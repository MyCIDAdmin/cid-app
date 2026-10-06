"""Export Excel de la Jahresbilanz (2026-10-06) — classeur multi-feuilles avec formules pour les
totaux/écarts, afin que le service financier puisse corriger/compléter les chiffres à la main."""

from openpyxl import Workbook
from openpyxl.styles import Font

from .exports import _ecrire_en_tete

SOURCES_LIBELLES = {
    "cotisations": "Cotisations / Beiträge",
    "dons": "Dons / Spenden",
    "projets": "Projets / Projekte",
    "adhesions": "Adhésions / Mitgliedschaften",
    "boutique": "Boutique / Shop",
    "evenements": "Événements / Veranstaltungen",
}
EUR = '#,##0.00 "€"'
GRAS = Font(bold=True)


def _f(valeur):
    return float(valeur)


def construire_classeur_bilan(bilan, depenses) -> Workbook:
    classeur = Workbook()

    # --- Bilan : recettes puis dépenses par catégorie, avec total/résultat en formules.
    feuille = classeur.active
    feuille.title = "Bilan"
    annee = bilan["annee"]
    _ecrire_en_tete(feuille, ["Poste", str(annee), str(annee - 1), "Écart"])
    ligne = 2
    feuille.cell(row=ligne, column=1, value="RECETTES / EINNAHMEN").font = GRAS
    debut = ligne + 1
    for r in bilan["recettes"]["lignes"]:
        ligne += 1
        feuille.cell(row=ligne, column=1, value=SOURCES_LIBELLES[r["cle"]])
        feuille.cell(row=ligne, column=2, value=_f(r["montant"]))
        feuille.cell(row=ligne, column=3, value=_f(r["montant_precedent"]))
        feuille.cell(row=ligne, column=4, value=f"=B{ligne}-C{ligne}")
    ligne += 1
    total_rec = ligne
    feuille.cell(row=ligne, column=1, value="Total recettes").font = GRAS
    for col in "BC":
        feuille[f"{col}{ligne}"] = f"=SUM({col}{debut}:{col}{ligne - 1})"
    feuille[f"D{ligne}"] = f"=B{ligne}-C{ligne}"
    ligne += 2
    feuille.cell(row=ligne, column=1, value="DÉPENSES / AUSGABEN").font = GRAS
    debut = ligne + 1
    for d in bilan["depenses"]["lignes"]:
        ligne += 1
        feuille.cell(row=ligne, column=1, value=d["nom"])
        feuille.cell(row=ligne, column=2, value=_f(d["montant"]))
        feuille.cell(row=ligne, column=3, value=_f(d["montant_precedent"]))
        feuille.cell(row=ligne, column=4, value=f"=B{ligne}-C{ligne}")
    ligne += 1
    total_dep = ligne
    feuille.cell(row=ligne, column=1, value="Total dépenses").font = GRAS
    for col in "BC":
        feuille[f"{col}{ligne}"] = f"=SUM({col}{debut}:{col}{ligne - 1})"
    feuille[f"D{ligne}"] = f"=B{ligne}-C{ligne}"
    ligne += 2
    feuille.cell(row=ligne, column=1, value="RÉSULTAT / ERGEBNIS").font = GRAS
    for col in "BC":
        feuille[f"{col}{ligne}"] = f"={col}{total_rec}-{col}{total_dep}"
    feuille[f"D{ligne}"] = f"=B{ligne}-C{ligne}"
    for row in feuille.iter_rows(min_row=2, min_col=2, max_col=4):
        for cellule in row:
            cellule.number_format = EUR
    feuille.column_dimensions["A"].width = 38

    # --- Budget vs. Ist
    feuille = classeur.create_sheet("Budget vs Ist")
    _ecrire_en_tete(feuille, ["Catégorie", "Budget", "Ist", "Écart", "% du budget", "Statut"])
    for i, d in enumerate(bilan["depenses"]["lignes"], start=2):
        feuille.cell(row=i, column=1, value=d["nom"])
        feuille.cell(row=i, column=2, value=_f(d["budget"])).number_format = EUR
        feuille.cell(row=i, column=3, value=_f(d["montant"])).number_format = EUR
        feuille.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = EUR
        feuille.cell(row=i, column=5, value=f'=IF(B{i}=0,"",C{i}/B{i})').number_format = "0.0%"
        feuille.cell(row=i, column=6, value=d["statut_budget"])
    feuille.column_dimensions["A"].width = 38

    # --- Mensuel
    feuille = classeur.create_sheet("Mensuel")
    _ecrire_en_tete(feuille, ["Mois", "Recettes", "Dépenses", "Résultat", "Cumul"])
    for i, m in enumerate(bilan["mensuel"], start=2):
        feuille.cell(row=i, column=1, value=m["mois"])
        feuille.cell(row=i, column=2, value=_f(m["recettes"])).number_format = EUR
        feuille.cell(row=i, column=3, value=_f(m["depenses"])).number_format = EUR
        feuille.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = EUR
        feuille.cell(row=i, column=5, value=f"=SUM(D$2:D{i})").number_format = EUR

    # --- Événements et projets
    for titre, cle in (
        ("Résultat événements", "resultats_evenements"),
        ("Résultat projets", "resultats_projets"),
    ):
        feuille = classeur.create_sheet(titre)
        _ecrire_en_tete(feuille, ["Titre", "Recettes", "Dépenses", "Résultat"])
        for i, e in enumerate(bilan[cle], start=2):
            feuille.cell(row=i, column=1, value=e["titre"])
            feuille.cell(row=i, column=2, value=_f(e["recettes"])).number_format = EUR
            feuille.cell(row=i, column=3, value=_f(e["depenses"])).number_format = EUR
            feuille.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = EUR
        feuille.column_dimensions["A"].width = 38

    # --- Détail des dépenses
    feuille = classeur.create_sheet("Dépenses")
    _ecrire_en_tete(
        feuille,
        [
            "Date",
            "Fournisseur",
            "Catégorie",
            "Montant",
            "Statut",
            "Événement",
            "Projet",
            "Saisi par",
            "Décidé par",
        ],
    )
    for i, d in enumerate(depenses, start=2):
        valeurs = [
            d.date_depense.strftime("%d/%m/%Y"),
            d.fournisseur,
            d.categorie.nom,
            _f(d.montant),
            d.statut,
            d.evenement.titre if d.evenement else "",
            d.projet.titre if d.projet else "",
            d.saisie_par.email if d.saisie_par else "",
            d.decide_par.email if d.decide_par else "",
        ]
        for col, v in enumerate(valeurs, start=1):
            feuille.cell(row=i, column=col, value=v)
        feuille.cell(row=i, column=4).number_format = EUR
    return classeur


def construire_csv_buchungen(zeilen) -> str:
    """CSV für den Steuerberater : Semikolon, Dezimalkomma, Datum TT.MM.JJJJ, UTF-8 mit BOM
    (öffnet in Excel/DATEV-Importen ohne Umlaut-Probleme). Ausgaben stehen negativ."""
    import csv
    import io

    puffer = io.StringIO()
    puffer.write("\ufeff")
    w = csv.writer(puffer, delimiter=";", lineterminator="\r\n")
    w.writerow(
        [
            "Datum",
            "Typ",
            "Kategorie",
            "Beschreibung",
            "Gegenpartei",
            "Betrag (EUR)",
            "Beleg",
            "Referenz",
        ]
    )
    for z in zeilen:
        w.writerow(
            [
                z["datum"].strftime("%d.%m.%Y"),
                z["typ"],
                z["kategorie"],
                z["beschreibung"],
                z["gegenpartei"],
                f"{z['betrag']:.2f}".replace(".", ","),
                z["beleg"],
                z["referenz"],
            ]
        )
    return puffer.getvalue()
