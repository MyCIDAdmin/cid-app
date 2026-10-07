"""
Excel-Export des Dashboards — app stats ("Statistiken & KPIs", Export als PDF/Excel-Dashboard).

Ein Arbeitsmappe, sechs Blätter (Sprache de/fr, Standard Deutsch — siehe apps.stats.i18n):
Kennzahlen, Finanzdaten, Mitglieder, Veranstaltungen, Projekte, Top-Beitragende. Alle Zahlen sind
dieselben wie auf dem Bildschirm (die Aufrufer berechnen/filtern, siehe views.StatsExportExcelView).
Gleicher Stil (roter Kopf #CC0000, fixierte Kopfzeile) wie die anderen Exporte der App.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.membres.utils_http import safe_cell

from . import i18n

_EN_TETE_STYLE = Font(bold=True, color="FFFFFF")
_REMPLISSAGE = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")

TEXTE = {
    "de": {
        "kennzahlen": "Kennzahlen",
        "finanzdaten": "Finanzdaten",
        "mitglieder": "Mitglieder",
        "veranstaltungen": "Veranstaltungen",
        "projekte": "Projekte",
        "top": "Top-Beitragende",
        "kennzahl": "Kennzahl",
        "wert": "Wert",
        "jahr": "Jahr",
        "saldo": "Saldo (€)",
        "einnahmen": "Einnahmen (€)",
        "ausgaben": "Ausgaben (€)",
        "einzug": "Einzugsquote (%)",
        "offene_beitraege": "Offene Beiträge (€)",
        "shop": "Einnahmen Shop (€)",
        "mitgliedschaften": "Einnahmen Mitgliedschaften (€)",
        "veranst_einnahmen": "Einnahmen Veranstaltungen (€)",
        "fin": "Finanzen",
        "mit": "Mitglieder",
        "ver": "Veranstaltungen",
        "gesamt": "Gesamt",
        "aktiv": "Aktiv",
        "inaktiv": "Inaktiv",
        "anzahl": "Anzahl",
        "fuellung": "Durchschnittliche Auslastung (%)",
        "anmeldungen": "Anmeldungen gesamt",
        "typ": "Typ",
        "datum": "Datum",
        "mitglied": "Mitglied / Partei",
        "beschreibung": "Beschreibung",
        "betrag": "Betrag (€)",
        "status": "Status",
        "stadt": "Stadt",
        "altersgruppe": "Altersgruppe",
        "veranstaltung": "Veranstaltung",
        "reserviert": "Reservierte Plätze",
        "max": "Maximale Plätze",
        "projekt": "Projekt",
        "aufgaben": "Aufgaben erledigt",
        "gesamt_aufgaben": "Aufgaben gesamt",
        "plan": "Plan (€)",
        "ist": "Ist (€)",
        "offen": "Offen (€)",
        "ergebnis": "Ergebnis (€)",
        "name": "Name",
        "total": "Summe (€)",
    },
    "fr": {
        "kennzahlen": "Indicateurs",
        "finanzdaten": "Données financières",
        "mitglieder": "Membres",
        "veranstaltungen": "Événements",
        "projekte": "Projets",
        "top": "Top contributeurs",
        "kennzahl": "Indicateur",
        "wert": "Valeur",
        "jahr": "Année",
        "saldo": "Solde (€)",
        "einnahmen": "Recettes (€)",
        "ausgaben": "Dépenses (€)",
        "einzug": "Taux de collecte (%)",
        "offene_beitraege": "Cotisations en attente (€)",
        "shop": "Revenus boutique (€)",
        "mitgliedschaften": "Revenus adhésions (€)",
        "veranst_einnahmen": "Revenus événements (€)",
        "fin": "Finances",
        "mit": "Membres",
        "ver": "Événements",
        "gesamt": "Total",
        "aktiv": "Actifs",
        "inaktiv": "Inactifs",
        "anzahl": "Nombre",
        "fuellung": "Taux de remplissage moyen (%)",
        "anmeldungen": "Inscriptions totales",
        "typ": "Type",
        "datum": "Date",
        "mitglied": "Membre / partie",
        "beschreibung": "Description",
        "betrag": "Montant (€)",
        "status": "Statut",
        "stadt": "Ville",
        "altersgruppe": "Tranche d'âge",
        "veranstaltung": "Événement",
        "reserviert": "Places réservées",
        "max": "Places maximum",
        "projekt": "Projet",
        "aufgaben": "Tâches terminées",
        "gesamt_aufgaben": "Tâches au total",
        "plan": "Prévu (€)",
        "ist": "Réel (€)",
        "offen": "Ouvert (€)",
        "ergebnis": "Résultat (€)",
        "name": "Nom",
        "total": "Total (€)",
    },
}


def _ecrire_en_tete(feuille, colonnes):
    for col_idx, libelle in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=1, column=col_idx, value=libelle)
        cellule.font = _EN_TETE_STYLE
        cellule.fill = _REMPLISSAGE
        feuille.column_dimensions[cellule.column_letter].width = max(len(libelle) + 2, 14)
    feuille.freeze_panes = "A2"


def _zeilen_schreiben(feuille, zeilen):
    for ligne_idx, valeurs in enumerate(zeilen, start=2):
        for col_idx, valeur in enumerate(valeurs, start=1):
            feuille.cell(row=ligne_idx, column=col_idx, value=safe_cell(valeur))


def _blatt(classeur, titel, spalten, zeilen, *, erstes=False, breiten=None):
    feuille = classeur.active if erstes else classeur.create_sheet()
    feuille.title = titel[:31]
    _ecrire_en_tete(feuille, spalten)
    _zeilen_schreiben(feuille, zeilen)
    for buchstabe, breite in (breiten or {}).items():
        feuille.column_dimensions[buchstabe].width = breite
    return feuille


def _kennzahlen(kf, km, ke, x):
    return [
        (f"{x['fin']} — {x['jahr']}", kf["annee"]),
        (f"{x['fin']} — {x['saldo']}", float(kf["solde"])),
        (f"{x['fin']} — {x['einnahmen']}", float(kf["recettes"])),
        (f"{x['fin']} — {x['ausgaben']}", float(kf["depenses"])),
        (f"{x['fin']} — {x['einzug']}", kf["taux_collecte"]),
        (f"{x['fin']} — {x['offene_beitraege']}", float(kf["cotisations_en_attente"])),
        (f"{x['fin']} — {x['shop']}", float(kf["revenus_boutique"])),
        (f"{x['fin']} — {x['mitgliedschaften']}", float(kf["revenus_adhesions"])),
        (f"{x['fin']} — {x['veranst_einnahmen']}", float(kf["revenus_evenements"])),
        (f"{x['mit']} — {x['gesamt']}", km["total"]),
        (f"{x['mit']} — {x['aktiv']}", km["actifs"]),
        (f"{x['mit']} — {x['inaktiv']}", km["inactifs"]),
        (f"{x['ver']} — {x['jahr']}", ke["annee"]),
        (f"{x['ver']} — {x['anzahl']}", ke["nombre_evenements"]),
        (f"{x['ver']} — {x['fuellung']}", ke["taux_remplissage_moyen"]),
        (f"{x['ver']} — {x['anmeldungen']}", ke["inscriptions_totales"]),
        (f"{x['ver']} — {x['einnahmen']}", float(ke["revenus"])),
    ]


def construire_classeur_dashboard(
    *, kpis_financier, kpis_membres, kpis_evenements, finances, kpis_projets=None, langue=None
) -> Workbook:
    """Baut (ohne zu speichern) die Excel-Arbeitsmappe des Dashboards ; alle Argumente sind vom
    Aufrufer bereits berechnet/gefiltert (siehe views.StatsExportExcelView)."""
    langue = langue if langue in TEXTE else i18n.STANDARD
    x = TEXTE[langue]
    classeur = Workbook()

    _blatt(
        classeur,
        x["kennzahlen"],
        [x["kennzahl"], x["wert"]],
        _kennzahlen(kpis_financier, kpis_membres, kpis_evenements, x),
        erstes=True,
        breiten={"A": 48, "B": 18},
    )
    _blatt(
        classeur,
        x["finanzdaten"],
        [x["typ"], x["datum"], x["mitglied"], x["beschreibung"], x["betrag"], x["status"]],
        [
            (
                i18n.typ(z["type"], langue),
                z["date"].strftime("%d.%m.%Y"),
                z["membre_nom"],
                z["description"],
                float(z["montant"]),
                i18n.status(z["statut"], langue),
            )
            for z in finances
        ],
        breiten={"D": 50},
    )
    mitglieder = [
        (x["stadt"] + ": " + z["ville_de"], z["nombre"]) for z in kpis_membres["par_ville"]
    ]
    mitglieder += [
        (x["altersgruppe"] + ": " + z["tranche"], z["nombre"])
        for z in kpis_membres["pyramide_ages"]
    ]
    _blatt(
        classeur,
        x["mitglieder"],
        [x["kennzahl"], x["anzahl"]],
        [
            (x["gesamt"], kpis_membres["total"]),
            (x["aktiv"], kpis_membres["actifs"]),
            (x["inaktiv"], kpis_membres["inactifs"]),
        ]
        + mitglieder,
        breiten={"A": 40},
    )
    _blatt(
        classeur,
        x["veranstaltungen"],
        [x["veranstaltung"], x["reserviert"], x["max"]],
        [
            (e["titre"], e["places_reservees"], e["places_max"])
            for e in kpis_evenements["participation_par_evenement"]
        ],
        breiten={"A": 50},
    )
    if kpis_projets is not None:
        _blatt(
            classeur,
            x["projekte"],
            [
                x["projekt"],
                x["status"],
                x["aufgaben"],
                x["gesamt_aufgaben"],
                x["plan"],
                x["ist"],
                x["offen"],
                x["einnahmen"],
                x["ergebnis"],
            ],
            [
                (
                    z["titre"],
                    i18n.status(z["statut"], langue),
                    z["aufgaben_erledigt"],
                    z["aufgaben_gesamt"],
                    float(z["plan"]),
                    float(z["ist"]),
                    float(z["offen"]),
                    float(z["einnahmen"]),
                    float(z["ergebnis"]),
                )
                for z in kpis_projets["projekte"]
            ],
            breiten={"A": 40},
        )
    _blatt(
        classeur,
        x["top"],
        [x["name"], x["total"]],
        [(z["nom"], float(z["total"])) for z in kpis_financier["top_contributeurs"]],
        breiten={"A": 40},
    )
    return classeur
