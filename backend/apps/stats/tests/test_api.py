"""Tests API — app stats (accès réservé Admin/DG/Bureau Admin — FDD écrans Release 1 §2.3)."""

import datetime
from decimal import Decimal

import pytest
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.cotisations.models import StatutCotisation, TypeArticle
from apps.cotisations.tests.factories import CotisationFactory
from apps.evenements.models import StatutEvenement
from apps.evenements.tests.factories import EvenementFactory, InscriptionFactory
from apps.membres.models import StatutMembre
from apps.membres.tests.factories import MembreFactory

pytestmark = pytest.mark.django_db


@pytest.fixture
def api_client():
    return APIClient()


def _user_avec_membre(role, email, **membre_kwargs):
    user = User.objects.create_user(email=email, password="Password123!", role=role, is_active=True)
    membre = MembreFactory(user=user, **membre_kwargs)
    return user, membre


def _auth(api_client, user):
    api_client.force_authenticate(user=user)
    return api_client


FINANCIER_URL = "stats:financier"
MEMBRES_URL = "stats:membres"
EVENEMENTS_URL = "stats:evenements"
FINANCES_URL = "stats:finances"
EXPORT_EXCEL_URL = "stats:export-excel"
EXPORT_PDF_URL = "stats:export-pdf"


def test_financier_non_authentifie_refuse(api_client):
    resp = api_client.get(reverse(FINANCIER_URL))
    assert resp.status_code == 401


def test_membre_normal_ne_peut_pas_voir_les_stats(api_client):
    user, _ = _user_avec_membre(Role.MEMBRE, "m1@example.de")
    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 403


def test_rh_ne_peut_pas_voir_les_stats(api_client):
    user, _ = _user_avec_membre(Role.RH, "rh1@example.de")
    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 403


def test_bureau_admin_peut_voir_les_stats_financieres(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "bureau1@example.de")
    CotisationFactory(membre=membre)

    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 200
    assert Decimal(resp.data["recettes"]) == Decimal("45.00")
    assert "top_contributeurs" in resp.data


def test_directeur_financier_peut_voir_les_stats_membres(api_client):
    user, _ = _user_avec_membre(Role.DIR_FINANCIER, "dg1@example.de")
    MembreFactory(statut=StatutMembre.ACTIF, ville_de="Berlin")
    MembreFactory(statut=StatutMembre.INACTIF, ville_de="Berlin")

    resp = _auth(api_client, user).get(reverse(MEMBRES_URL))
    assert resp.status_code == 200
    assert resp.data["total"] >= 2
    assert resp.data["actifs"] >= 1


def test_admin_peut_voir_les_stats_evenements(api_client):
    user, membre = _user_avec_membre(Role.SUPER_ADMIN, "admin1@example.de")
    evenement = EvenementFactory(statut=StatutEvenement.PUBLIE, places_max=10)
    InscriptionFactory(evenement=evenement, membre=membre, places=2)

    resp = _auth(api_client, user).get(
        reverse(EVENEMENTS_URL), {"annee": evenement.date_evenement.year}
    )
    assert resp.status_code == 200
    assert resp.data["nombre_evenements"] >= 1
    assert resp.data["inscriptions_totales"] >= 2


def test_annee_invalide_refusee(api_client):
    user, _ = _user_avec_membre(Role.SUPER_ADMIN, "admin2@example.de")
    resp = _auth(api_client, user).get(reverse(FINANCIER_URL), {"annee": "pas-une-annee"})
    assert resp.status_code == 400


def test_filtre_land_transmis_aux_stats_membres(api_client):
    # Ajouté le 2026-09-19 (demande utilisateur : "Bei ... Statistiken & KPIs füge mehr
    # Filtermöglichten hinzu z.B. Bundesland").
    user, _ = _user_avec_membre(Role.SUPER_ADMIN, "admin3@example.de")
    MembreFactory(land_de="BY")
    MembreFactory(land_de="BE")

    resp = _auth(api_client, user).get(reverse(MEMBRES_URL), {"land": "BY"})
    assert resp.status_code == 200
    assert resp.data["total"] == 1


def test_filtre_pays_et_dates_adhesion_transmis_aux_stats_membres(api_client):
    user, _ = _user_avec_membre(Role.SUPER_ADMIN, "admin4@example.de")
    MembreFactory(pays="TN", date_adhesion=datetime.date(2026, 6, 1))
    MembreFactory(pays="DE", date_adhesion=datetime.date(2020, 1, 1))

    resp = _auth(api_client, user).get(
        reverse(MEMBRES_URL),
        {"pays": "TN", "date_adhesion_apres": "2026-01-01"},
    )
    assert resp.status_code == 200
    assert resp.data["total"] == 1


# ---------------------------------------------------------------------------
# Onglet "Finanzdaten" + exports PDF/Excel (module "Statistiken & KPIs", ajouté le 2026-09-25 :
# "Tab für alle Finanzdaten (filterbar/sortierbar)" / "Export als PDF/Excel-Dashboard").
# ---------------------------------------------------------------------------


def test_finances_non_authentifie_refuse(api_client):
    resp = api_client.get(reverse(FINANCES_URL))
    assert resp.status_code == 401


def test_membre_normal_ne_peut_pas_voir_les_finances(api_client):
    user, _ = _user_avec_membre(Role.MEMBRE, "fin-membre@example.de")
    resp = _auth(api_client, user).get(reverse(FINANCES_URL))
    assert resp.status_code == 403


def test_bureau_admin_voit_les_finances(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "fin-bureau@example.de")
    CotisationFactory(
        membre=membre,
        type_article=TypeArticle.DON,
        statut=StatutCotisation.PAYEE,
        montant=Decimal("20.00"),
    )

    resp = _auth(api_client, user).get(reverse(FINANCES_URL))
    assert resp.status_code == 200
    assert any(ligne["type"] == "don" for ligne in resp.data["results"])


def test_finances_type_transaction_invalide_refuse(api_client):
    user, _ = _user_avec_membre(Role.SUPER_ADMIN, "fin-invalide@example.de")
    resp = _auth(api_client, user).get(reverse(FINANCES_URL), {"type_transaction": "inexistant"})
    assert resp.status_code == 400


def test_finances_filtre_type_transaction(api_client):
    user, membre = _user_avec_membre(Role.SUPER_ADMIN, "fin-filtre@example.de")
    CotisationFactory(
        membre=membre,
        type_article=TypeArticle.DON,
        statut=StatutCotisation.PAYEE,
        montant=Decimal("20.00"),
    )
    CotisationFactory(
        membre=membre,
        type_article=TypeArticle.COTISATION,
        statut=StatutCotisation.PAYEE,
        montant=Decimal("45.00"),
    )

    resp = _auth(api_client, user).get(reverse(FINANCES_URL), {"type_transaction": "don"})
    assert resp.status_code == 200
    assert all(ligne["type"] == "don" for ligne in resp.data["results"])


def test_export_excel_non_authentifie_refuse(api_client):
    resp = api_client.get(reverse(EXPORT_EXCEL_URL))
    assert resp.status_code == 401


def test_membre_normal_ne_peut_pas_exporter(api_client):
    user, _ = _user_avec_membre(Role.MEMBRE, "export-membre@example.de")
    resp = _auth(api_client, user).get(reverse(EXPORT_EXCEL_URL))
    assert resp.status_code == 403


def test_export_excel_retourne_un_classeur_xlsx(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "export-bureau@example.de")
    CotisationFactory(membre=membre, type_article=TypeArticle.COTISATION, montant=Decimal("45.00"))

    resp = _auth(api_client, user).get(reverse(EXPORT_EXCEL_URL))
    assert resp.status_code == 200
    assert (
        resp["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    import io

    from openpyxl import load_workbook

    classeur = load_workbook(io.BytesIO(resp.content))
    assert classeur.sheetnames == [
        "Kennzahlen",
        "Finanzdaten",
        "Mitglieder",
        "Veranstaltungen",
        "Projekte",
        "Top-Beitragende",
    ]
    assert classeur["Kennzahlen"]["A1"].value == "Kennzahl"


def test_export_excel_franzoesisch_mit_langue_param(api_client):
    user, _ = _user_avec_membre(Role.BUREAU_ADMIN, "export-fr@example.de")
    resp = _auth(api_client, user).get(reverse(EXPORT_EXCEL_URL), {"langue": "fr"})
    import io

    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(resp.content)).sheetnames[0] == "Indicateurs"


def test_export_pdf_non_authentifie_refuse(api_client):
    resp = api_client.get(reverse(EXPORT_PDF_URL))
    assert resp.status_code == 401


def test_membre_normal_ne_peut_pas_exporter_pdf(api_client):
    user, _ = _user_avec_membre(Role.MEMBRE, "export-pdf-membre@example.de")
    resp = _auth(api_client, user).get(reverse(EXPORT_PDF_URL))
    assert resp.status_code == 403


def test_export_pdf_retourne_un_pdf(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "export-pdf-bureau@example.de")
    CotisationFactory(membre=membre, type_article=TypeArticle.COTISATION, montant=Decimal("45.00"))

    resp = _auth(api_client, user).get(reverse(EXPORT_PDF_URL))
    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF-")


# ---------------------------------------------------------------------------
# Phase D (ajoutée le 2026-09-23) — page de gestion "Statistiken & KPIs" (page_stats) désormais
# pilotée par apps.rbac (real enforcement, y compris pour les rôles système eux-mêmes).
# ---------------------------------------------------------------------------


def _set_matrice_cellule(role_slug, module_slug, niveau_acces):
    from apps.rbac.models import RoleDefinition, RoleModulePermission

    role = RoleDefinition.objects.get(slug=role_slug, is_system=True)
    RoleModulePermission.objects.update_or_create(
        role=role, module=module_slug, defaults={"niveau_acces": niveau_acces}
    )


def test_phase_d_bureau_admin_perd_lacces_aux_stats_si_matrice_le_dit(api_client):
    from apps.rbac.models import NiveauAcces

    _set_matrice_cellule("bureau_admin", "page_stats", NiveauAcces.AUCUN)
    user, _ = _user_avec_membre(Role.BUREAU_ADMIN, "phased-stats-restrict@example.de")

    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 403


def test_phase_d_role_personnalise_peut_voir_les_stats_via_la_matrice(api_client):
    from apps.rbac.models import NiveauAcces
    from apps.rbac.tests.factories import (
        RoleDefinitionFactory,
        RoleModulePermissionFactory,
        UserRoleAssignmentFactory,
    )

    user, _ = _user_avec_membre(Role.MEMBRE, "phased-stats-grant@example.de")
    role_perso = RoleDefinitionFactory(slug="stats-viewer")
    RoleModulePermissionFactory(
        role=role_perso, module="page_stats", niveau_acces=NiveauAcces.LECTURE
    )
    UserRoleAssignmentFactory(user=user, role=role_perso)

    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 200


def test_phase_d_super_admin_voit_toujours_les_stats_meme_si_matrice_dit_aucun(api_client):
    from apps.rbac.models import NiveauAcces

    _set_matrice_cellule("super_admin", "page_stats", NiveauAcces.AUCUN)
    user, _ = _user_avec_membre(Role.SUPER_ADMIN, "phased-stats-super@example.de")

    resp = _auth(api_client, user).get(reverse(FINANCIER_URL))
    assert resp.status_code == 200


# --- Pivot (2026-10-06) ------------------------------------------------------------------------


def _beitrag(membre, jahr, montant="45.00", monat=3):
    return CotisationFactory(
        membre=membre,
        type_article=TypeArticle.COTISATION,
        montant=Decimal(montant),
        statut=StatutCotisation.PAYEE,
        date_paiement=datetime.datetime(jahr, monat, 5, 12, tzinfo=datetime.timezone.utc),
    )


def _ausgabe(jahr, montant="20.00", lieferant="Druckerei Meier"):
    from apps.finances.models import CategorieDepense, Depense, StatutDepense

    return Depense.objects.create(
        date_depense=datetime.date(jahr, 4, 2),
        montant=Decimal(montant),
        categorie=CategorieDepense.objects.filter(actif=True).first(),
        fournisseur=lieferant,
        statut=StatutDepense.APPROUVEE,
    )


def test_pivot_liefert_matrix_mit_summen(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "pivot@example.de")
    jahr = datetime.date.today().year
    _beitrag(membre, jahr)
    resp = _auth(api_client, user).get(
        reverse("stats:pivot"),
        {"zeilen": "kategorie", "spalten": "jahr", "kennzahlen": "einnahmen"},
    )
    assert resp.status_code == 200
    daten = resp.json()
    assert daten["gesamt"] == [45.0]
    assert [s["label"] for s in daten["spalten"]] == [str(jahr)]
    zeile = next(z for z in daten["zeilen"] if z["label"] == "Beiträge")
    assert zeile["werte"] == [[45.0]] and zeile["summe"] == [45.0]


def test_pivot_trennt_einnahmen_und_ausgaben(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-ea@example.de")
    jahr = datetime.date.today().year
    _beitrag(membre, jahr, "100.00")
    _ausgabe(jahr, "30.00")
    resp = _auth(api_client, user).get(
        reverse("stats:pivot"),
        {"zeilen": "jahr", "kennzahlen": "einnahmen,ausgaben,saldo,anzahl"},
    )
    assert resp.status_code == 200
    daten = resp.json()
    assert daten["kennzahlen"] == ["einnahmen", "ausgaben", "saldo", "anzahl"]
    # Ausgaben erscheinen positiv und getrennt, der Saldo ist die Differenz.
    assert daten["gesamt"] == [100.0, 30.0, 70.0, 2]
    assert daten["zeilen"][0]["summe"] == [100.0, 30.0, 70.0, 2]


def test_pivot_mehrere_dimensionen_je_achse(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-md@example.de")
    jahr = datetime.date.today().year
    _beitrag(membre, jahr, "100.00", monat=1)
    _ausgabe(jahr, "30.00")
    resp = _auth(api_client, user).get(
        reverse("stats:pivot"),
        {"zeilen": "jahr,typ", "spalten": "kategorie", "kennzahlen": "saldo,anzahl"},
    )
    assert resp.status_code == 200
    daten = resp.json()
    assert daten["zeilen_dims"] == ["jahr", "typ"]
    labels = {tuple(z["labels"]) for z in daten["zeilen"]}
    assert labels == {(str(jahr), "Einnahme"), (str(jahr), "Ausgabe")}
    assert all(len(w) == 2 for z in daten["zeilen"] for w in z["werte"])
    assert len(daten["spalten"]) == 2  # Beiträge + Ausgabenkategorie


def test_pivot_filter_typ_kategorie_und_gegenpartei(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-f@example.de")
    jahr = datetime.date.today().year
    _beitrag(membre, jahr, "100.00")
    _ausgabe(jahr, "30.00", "Druckerei Meier")
    _ausgabe(jahr, "10.00", "Catering Schmidt")
    client = _auth(api_client, user)
    basis = {"zeilen": "typ", "kennzahlen": "saldo"}

    nur_ausgaben = client.get(reverse("stats:pivot"), {**basis, "f_typ": "ausgabe"}).json()
    assert nur_ausgaben["gesamt"] == [-40.0]
    assert nur_ausgaben["anzahl_buchungen"] == 2

    lieferant = client.get(reverse("stats:pivot"), {**basis, "f_gegenpartei": "druckerei"}).json()
    assert lieferant["gesamt"] == [-30.0]

    kategorie = client.get(reverse("stats:pivot"), {**basis, "f_kategorie": "Beiträge"}).json()
    assert kategorie["gesamt"] == [100.0]
    assert kategorie["filter"] == {"kategorie": ["Beiträge"]}


def test_pivot_optionen_liefern_kategorien(api_client):
    user, membre = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-o@example.de")
    _beitrag(membre, datetime.date.today().year)
    resp = _auth(api_client, user).get(reverse("stats:pivot-optionen"))
    assert resp.status_code == 200
    assert "Beiträge" in resp.json()["kategorie"]
    assert resp.json()["typ"] == ["einnahme", "ausgabe"]


def test_pivot_dimension_doppelt_oder_zu_viele_400(api_client):
    user, _ = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-dup@example.de")
    client = _auth(api_client, user)
    assert (
        client.get(reverse("stats:pivot"), {"zeilen": "jahr", "spalten": "jahr"}).status_code == 400
    )
    assert (
        client.get(reverse("stats:pivot"), {"zeilen": "jahr,monat,typ,kategorie"}).status_code
        == 400
    )
    assert client.get(reverse("stats:pivot"), {"kennzahlen": "umsatz"}).status_code == 400


def test_pivot_ungueltige_dimension_400(api_client):
    user, _ = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-bad@example.de")
    resp = _auth(api_client, user).get(reverse("stats:pivot"), {"zeilen": "unsinn"})
    assert resp.status_code == 400


def test_pivot_membre_normal_verboten(api_client):
    user, _ = _user_avec_membre(Role.MEMBRE, "pivot-m@example.de")
    assert _auth(api_client, user).get(reverse("stats:pivot")).status_code == 403


def test_pivot_export_csv_und_excel(api_client):
    user, _ = _user_avec_membre(Role.BUREAU_ADMIN, "pivot-exp@example.de")
    client = _auth(api_client, user)
    csv_resp = client.get(reverse("stats:export-pivot"), {"datei": "csv", "zeilen": "typ"})
    assert csv_resp.status_code == 200
    assert csv_resp["Content-Type"].startswith("text/csv")
    xlsx = client.get(
        reverse("stats:export-pivot"),
        {"zeilen": "typ,kategorie", "spalten": "monat", "kennzahlen": "einnahmen,ausgaben"},
    )
    assert xlsx.status_code == 200
    assert xlsx["Content-Disposition"].endswith('.xlsx"')
