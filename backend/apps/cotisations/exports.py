"""
Export Excel des paiements — app cotisations (demande utilisateur du 2026-09-25, module
"Ausstehende Zahlungen" (renommé "Zahlungen" — voir CotisationsEnAttentePage.tsx côté frontend) :
"Excel-Export der Zahlungen").

Le queryset est filtré/scopé par l'appelant (CotisationViewSet.export — mêmes filtres que la
liste, voir CotisationFilter, et même scope IDOR que get_queryset) : ce module ne fait que
construire le classeur, même séparation des responsabilités que
apps.boutique.exports/apps.membres.exports. Même style visuel (en-tête rouge #CC0000) que ces
deux modules, pour une seule identité visuelle d'export dans toute l'application.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

_EN_TETE_STYLE = Font(bold=True, color="FFFFFF")
_REMPLISSAGE = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")

# Colonnes disponibles à l'export, dans leur ordre d'affichage — pas de sélection de colonnes ici
# (même choix que apps.boutique.exports : la demande utilisateur ne porte que sur l'export
# lui-même, sans mention d'un besoin de personnalisation des colonnes).
_COLONNES = [
    "Membre",
    "N° membre",
    "Email",
    "Type d'article",
    "Libellé",
    "Montant (€)",
    "Mode de paiement",
    "Statut",
    "Référence transaction",
    "Date de création",
    "Date de paiement",
    "Saisi par (staff)",
]

# Colonnes forcées au format Texte Excel (number_format '@') — un numéro de membre ou une
# référence de transaction ne doit jamais être réinterprété comme un nombre à l'ouverture du
# fichier (même raison que apps.boutique.exports._COLONNES_TEXTE).
_COLONNES_TEXTE = {2, 9}  # N° membre, Référence transaction


def _formater_date(valeur) -> str:
    return valeur.strftime("%d/%m/%Y %H:%M") if valeur else ""


def construire_classeur_cotisations(queryset) -> Workbook:
    """
    Construit (sans l'enregistrer) le classeur .xlsx d'export des paiements — `queryset` doit
    déjà être filtré/scopé par l'appelant (voir CotisationViewSet.export). Un export n'est pas
    paginé : il contient toute la sélection filtrée, comme pour
    apps.boutique.exports.construire_classeur_commandes.
    """
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Zahlungen"

    for col_idx, libelle in enumerate(_COLONNES, start=1):
        cellule = feuille.cell(row=1, column=col_idx, value=libelle)
        cellule.font = _EN_TETE_STYLE
        cellule.fill = _REMPLISSAGE

    ligne_idx = 2
    for cotisation in queryset.select_related("membre", "saisie_par"):
        valeurs = [
            f"{cotisation.membre.prenom} {cotisation.membre.nom}" if cotisation.membre else "",
            cotisation.membre.numero_membre if cotisation.membre else "",
            cotisation.membre.email if cotisation.membre else "",
            cotisation.get_type_article_display(),
            cotisation.libelle,
            float(cotisation.montant),
            cotisation.get_mode_paiement_display() if cotisation.mode_paiement else "",
            cotisation.get_statut_display(),
            cotisation.reference_transaction or "",
            _formater_date(cotisation.created_at),
            _formater_date(cotisation.date_paiement),
            (
                f"{cotisation.saisie_par.prenom} {cotisation.saisie_par.nom}"
                if cotisation.saisie_par
                else ""
            ),
        ]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne_idx, column=col_idx, value=valeur)
            if col_idx in _COLONNES_TEXTE:
                cellule.number_format = "@"
        ligne_idx += 1

    for col_idx, libelle in enumerate(_COLONNES, start=1):
        feuille.column_dimensions[feuille.cell(row=1, column=col_idx).column_letter].width = max(
            len(libelle) + 2, 12
        )
    feuille.freeze_panes = "A2"

    return classeur
