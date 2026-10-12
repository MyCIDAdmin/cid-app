"""Massenimport von Zahlungen per Excel (Release 3, 2026-10-12) — nutzt die zweistufige Mechanik
der Mitglieder-Importe (apps.membres.import_gemeinsam : Prüfen -> Bestätigen -> Bericht).

Entscheidung des Nutzers (AskUserQuestion 2026-10-12): ALLE Zahlungen werden übernommen, auch ohne
Zuordnung. Spalten : datum, betrag (Pflicht) · zahlungsart, referenz, verwendungszweck, email,
cin, artikel, jahr (optional).
  - Mitglied (email und/oder cin) gefunden UND Artikel angegeben -> sofort gebucht (bezahlte
    Cotisation, keine Benachrichtigung).
  - sonst -> „offen“ (nicht zugeordnet / unvollständig) ; später in der Oberfläche ergänzen,
    korrigieren, zuordnen und buchen.
  - Referenz schon vorhanden (oder gleiche Zahlung desselben Mitglieds am selben Tag) -> Dublette,
    wird übersprungen (nie doppelt gebucht).
Prüfen schreibt nichts."""

import logging
from decimal import Decimal

from django.db import transaction

from apps.membres.import_gemeinsam import (
    ERG_FEHLER,
    ERG_IMPORTIERT,
    ERG_UEBERSPRUNGEN,
    STATUS_DUBLETTE,
    STATUS_FEHLER,
    STATUS_NEU,
    Analyse,
    Ergebnis,
    ImportDateiFehler,
    ZeileAnalyse,
    ZeilenErgebnis,
    erzeuge_bericht,
    lies_kopfzeile_und_zeilen,
    oeffne_arbeitsmappe,
)
from apps.membres.imports_beitraege import _parse_betrag, _parse_datum, _parse_modus
from apps.membres.imports_historique import (
    ANNEE_MAX,
    ANNEE_MIN,
    _lade_caches,
    _map_identite,
    _normalize,
    _resoudre_membre,
)

from .models import (
    Cotisation,
    ImportListe,
    ImportListenWert,
    StatutCotisation,
    StatutZahlungseingang,
    TypeArticle,
    Zahlungseingang,
)
from .zahlungseingang_services import ARTIKEL_ERLAUBT, buchen

logger = logging.getLogger(__name__)

_SPALTEN = {
    "datum": ("datum", "zahlungsdatum", "date", "date_paiement", "buchungsdatum"),
    "betrag": ("betrag", "montant", "amount", "summe"),
    "modus": ("zahlungsart", "mode", "mode_paiement", "zahlungsmodus", "payment_mode"),
    "referenz": ("referenz", "reference", "référence", "transaktion", "ref"),
    "zweck": ("verwendungszweck", "bemerkung", "notiz", "libelle", "libellé", "note", "zweck"),
    "artikel": ("artikel", "article", "typ", "type", "art"),
    "jahr": ("jahr", "annee", "année", "year"),
}
_ARTIKEL_ALIASE = {
    TypeArticle.COTISATION: ("cotisation", "mitgliedsbeitrag", "beitrag", "jahresbeitrag"),
    TypeArticle.ADHESION: ("adhesion", "adhésion", "aufnahme", "aufnahmegebühr", "mitgliedschaft"),
    TypeArticle.DON: ("don", "spende"),
    TypeArticle.AUTRE_LIBRE: ("autre", "sonstiges", "sonstige", "anders", "autre_libre"),
}
_ARTIKEL_LOOKUP = {a: t.value for t, alias in _ARTIKEL_ALIASE.items() for a in alias}
_ARTIKEL_TEXT = {
    TypeArticle.COTISATION: "Mitgliedsbeitrag",
    TypeArticle.ADHESION: "Aufnahmegebühr",
    TypeArticle.DON: "Spende",
    TypeArticle.AUTRE_LIBRE: "Sonstiges",
}


def _spalten(kopf) -> dict:
    norm = [_normalize(c) for c in kopf]
    mapping = {}
    for champ, aliase in _SPALTEN.items():
        for idx, header in enumerate(norm):
            if header in aliase:
                mapping[champ] = idx
                break
    return mapping


def _zelle(row, mapping, champ):
    idx = mapping.get(champ)
    return row[idx] if idx is not None and idx < len(row) else None


_DROPDOWN_ZEILEN = 1000  # Zeilen 2..1000 der Vorlage erhalten die Auswahlliste


def _dropdowns_anbringen(classeur, feuille, kopf):
    """Verstecktes Blatt „Listen“ + Datenüberprüfung für zahlungsart (streng) und
    verwendungszweck (Warnung: freier Text bleibt möglich)."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    listen = classeur.create_sheet("Listen")
    spalten = [
        ("zahlungsart", ImportListe.ZAHLUNGSART, "stop"),
        ("verwendungszweck", ImportListe.VERWENDUNGSZWECK, "warning"),
    ]
    for lcol, (feld, liste, stil) in enumerate(spalten, start=1):
        werte = _listenwerte(liste)
        listen.cell(row=1, column=lcol, value=feld)
        for zeile, wert in enumerate(werte, start=2):
            listen.cell(row=zeile, column=lcol, value=wert)
        if not werte:
            continue
        buchstabe = get_column_letter(lcol)
        validierung = DataValidation(
            type="list",
            formula1=f"=Listen!${buchstabe}$2:${buchstabe}${len(werte) + 1}",
            allow_blank=True,
            errorStyle=stil,
            showErrorMessage=True,
            errorTitle=feld,
            error="Bitte einen Wert aus der Liste wählen.",
        )
        ziel = get_column_letter(kopf.index(feld) + 1)
        validierung.add(f"{ziel}2:{ziel}{_DROPDOWN_ZEILEN}")
        feuille.add_data_validation(validierung)
    listen.sheet_state = "hidden"


def construire_classeur_template_zahlungen():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Zahlungen"
    kopf = [
        "datum", "betrag", "zahlungsart", "referenz", "verwendungszweck",
        "email", "cin", "artikel", "jahr",
    ]  # fmt: skip
    for col, champ in enumerate(kopf, start=1):
        zelle = feuille.cell(row=1, column=col, value=champ)
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")
        feuille.column_dimensions[zelle.column_letter].width = max(len(champ) + 2, 14)
    beispiele = [
        ["15.03.2026", 45, "Überweisung", "SEPA-2026-0042", "Beitrag 2026",
         "amira@example.de", "", "Mitgliedsbeitrag", 2026],
        ["16.03.2026", 20, "Bar", "", "Spende Mitgliederversammlung", "", "", "", ""],
    ]  # fmt: skip
    for zeile, werte in enumerate(beispiele, start=2):
        for col, wert in enumerate(werte, start=1):
            feuille.cell(row=zeile, column=col, value=wert)
    _dropdowns_anbringen(classeur, feuille, kopf)
    notes = classeur.create_sheet("Notes")
    hinweise = [
        "Pflicht: datum (TT.MM.JJJJ) und betrag. Alles andere ist optional.",
        "Mit email und/oder cin sowie artikel (Mitgliedsbeitrag / Aufnahmegebühr / Spende / "
        "Sonstiges) wird die Zahlung sofort gebucht.",
        "Fehlt das Mitglied (oder ist es unbekannt) oder der Artikel, bleibt die Zahlung „offen“ "
        "und kann später in der Oberfläche ergänzt, korrigiert, zugeordnet und gebucht werden.",
        "referenz verhindert doppelte Importe: dieselbe Referenz wird nie ein zweites Mal "
        "übernommen. jahr gilt für Mitgliedsbeiträge (Standard: Jahr des Datums).",
        "zahlungsart und verwendungszweck sind Auswahllisten (Blatt „Listen“). Die Werte werden in "
        "der Oberfläche unter Zahlungen > Importieren > „Listen verwalten“ gepflegt ; bei "
        "verwendungszweck ist auch freier Text erlaubt.",
        "Importierte Zahlungen lösen keine Benachrichtigung aus und ändern keinen Mitgliedsstatus.",
    ]
    for i, text in enumerate(hinweise, start=1):
        notes.cell(row=i, column=1, value=text)
    notes.column_dimensions["A"].width = 120
    return classeur


def _modus_aus_liste(roh, listen_modus: dict) -> str:
    """Zahlungsart: zuerst die verwalteten Listenwerte (auch inaktive), dann die festen Aliase."""
    text = _normalize(roh)
    if text in listen_modus:
        return listen_modus[text]
    return _parse_modus(roh)


def _listenwerte(liste: str) -> list[str]:
    return list(
        ImportListenWert.objects.filter(liste=liste, aktiv=True).values_list(
            "bezeichnung", flat=True
        )
    )


def _fmt(betrag) -> str:
    return f"{Decimal(betrag):.2f} €"


def _bezeichnung(membre) -> str:
    return f"{membre.numero_membre} ({membre.prenom} {membre.nom})"


def analyser_zahlungen(fichier, dateiname: str = "", importeur=None) -> Analyse:
    """PRÜFPHASE — schreibt nichts."""
    classeur = oeffne_arbeitsmappe(fichier)
    kopf, daten = lies_kopfzeile_und_zeilen(classeur)
    spalten = _spalten(kopf)
    fehlend = [c for c in ("datum", "betrag") if c not in spalten]
    if fehlend:
        raise ImportDateiFehler("Pflichtspalten fehlen: " + ", ".join(fehlend))
    identite = _map_identite(kopf)

    cache_email, cache_cin = _lade_caches()
    listen_modus = {
        _normalize(w.bezeichnung): w.modus
        for w in ImportListenWert.objects.filter(liste=ImportListe.ZAHLUNGSART).exclude(modus="")
    }
    referenzen_db = set(
        Zahlungseingang.objects.exclude(status=StatutZahlungseingang.VERWORFEN)
        .exclude(referenz="")
        .values_list("referenz", flat=True)
    ) | set(
        Cotisation.objects.exclude(reference_transaction__isnull=True).values_list(
            "reference_transaction", flat=True
        )
    )
    gesehen_ref = {}
    zeilen = []

    for nummer, row in daten:
        email = (
            row[identite["email"]] if "email" in identite and identite["email"] < len(row) else None
        )
        cin = row[identite["cin"]] if "cin" in identite and identite["cin"] < len(row) else None
        anzeige = {"prenom": "", "nom": "", "email": str(email or "").strip()}

        def fehler(grund):
            zeilen.append(
                ZeileAnalyse(ligne=nummer, status=STATUS_FEHLER, grund=grund, anzeige=anzeige)
            )

        try:
            datum = _parse_datum(_zelle(row, spalten, "datum"))
            if datum is None:
                raise ValueError("Datum fehlt")
            betrag = _parse_betrag(_zelle(row, spalten, "betrag"))
            if betrag is None or betrag <= 0:
                raise ValueError("Betrag fehlt oder ist nicht größer als 0")
            modus = _modus_aus_liste(_zelle(row, spalten, "modus"), listen_modus)
            artikel_roh = _normalize(_zelle(row, spalten, "artikel"))
            artikel = ""
            if artikel_roh:
                artikel = _ARTIKEL_LOOKUP.get(artikel_roh, "")
                if not artikel:
                    raise ValueError(f"Unbekannter Artikel {artikel_roh!r}")
            jahr_roh = _normalize(_zelle(row, spalten, "jahr"))
            jahr = None
            if jahr_roh:
                jahr_text = jahr_roh[:-2] if jahr_roh.endswith(".0") else jahr_roh
                if not jahr_text.isdigit() or not ANNEE_MIN <= int(jahr_text) <= ANNEE_MAX:
                    raise ValueError(f"Ungültiges Jahr {jahr_roh!r}")
                jahr = int(jahr_text)
        except ValueError as exc:
            fehler(str(exc))
            continue
        referenz = str(_zelle(row, spalten, "referenz") or "").strip()
        if len(referenz) > 64:
            fehler("Referenz ist länger als 64 Zeichen")
            continue
        zweck = str(_zelle(row, spalten, "zweck") or "").strip()[:255]

        membre, hinweis_mitglied = None, ""
        if str(email or "").strip() or str(cin or "").strip():
            try:
                membre = _resoudre_membre(email, cin, cache_email, cache_cin)
                anzeige.update(prenom=membre.prenom, nom=membre.nom)
            except ValueError as exc:
                hinweis_mitglied = str(exc)

        if referenz and referenz in gesehen_ref:
            fehler(
                "Referenz kommt in der Datei mehrfach vor "
                f"(bereits in Zeile {gesehen_ref[referenz]})"
            )
            continue
        if referenz:
            gesehen_ref[referenz] = nummer

        doppelt = referenz in referenzen_db if referenz else False
        if not referenz and membre is not None:
            doppelt = (
                Zahlungseingang.objects.filter(membre=membre, datum=datum, betrag=betrag)
                .exclude(status=StatutZahlungseingang.VERWORFEN)
                .exists()
                or Cotisation.objects.filter(
                    membre=membre,
                    montant=betrag,
                    date_paiement__date=datum,
                    statut=StatutCotisation.PAYEE,
                ).exists()
            )
        if doppelt:
            zeilen.append(
                ZeileAnalyse(
                    ligne=nummer,
                    status=STATUS_DUBLETTE,
                    grund=(
                        "Zahlung bereits vorhanden (Referenz bzw. gleiches Mitglied/Datum/Betrag)"
                        " — wird übersprungen"
                    ),
                    anzeige=anzeige,
                    existant=membre,
                )
            )
            continue

        fehlt = []
        if membre is None:
            fehlt.append("Mitglied")
        if artikel not in ARTIKEL_ERLAUBT:
            fehlt.append("Artikel")
        ergebnis_text = (
            "wird gebucht" if not fehlt else "bleibt offen — es fehlt: " + ", ".join(fehlt)
        )
        aenderungen = [
            {"champ": "betrag", "label": "Betrag", "alt": "", "neu": _fmt(betrag), "art": "neu"},
            {
                "champ": "datum",
                "label": "Datum",
                "alt": "",
                "neu": datum.strftime("%d.%m.%Y"),
                "art": "neu",
            },
            {
                "champ": "mitglied",
                "label": "Mitglied",
                "alt": "",
                "neu": _bezeichnung(membre) if membre else "nicht zugeordnet",
                "art": "neu",
            },
            {
                "champ": "artikel",
                "label": "Artikel",
                "alt": "",
                "neu": _ARTIKEL_TEXT.get(artikel, "nicht angegeben"),
                "art": "neu",
            },
            {
                "champ": "ergebnis",
                "label": "Ergebnis",
                "alt": "",
                "neu": ergebnis_text,
                "art": "neu",
            },
        ]
        grund = ergebnis_text[0].upper() + ergebnis_text[1:]
        if hinweis_mitglied:
            grund += f" ({hinweis_mitglied})"
        zeilen.append(
            ZeileAnalyse(
                ligne=nummer,
                status=STATUS_NEU,
                grund=grund,
                anzeige=anzeige,
                existant=membre,
                aenderungen=aenderungen,
                daten={
                    "datum": datum,
                    "betrag": betrag,
                    "modus": modus,
                    "referenz": referenz,
                    "zweck": zweck,
                    "artikel": artikel,
                    "jahr": jahr,
                    "hinweis": hinweis_mitglied,
                    "buchbar": not fehlt,
                },
            )
        )
    return Analyse(zeilen=zeilen, classeur=classeur, dateiname=dateiname)


def _zeile_schreiben(zeile: ZeileAnalyse, quelle: str, importeur) -> ZeilenErgebnis:
    d = zeile.daten
    try:
        with transaction.atomic():
            eingang = Zahlungseingang.objects.create(
                datum=d["datum"],
                betrag=d["betrag"],
                mode_paiement=d["modus"],
                referenz=d["referenz"],
                verwendungszweck=d["zweck"],
                membre=zeile.existant,
                type_article=d["artikel"],
                annee=d["jahr"],
                libelle=d["zweck"] or "",
                hinweis=d["hinweis"] or ("" if d["buchbar"] else "Mitglied oder Artikel fehlt"),
                quelle_datei=quelle[:200],
                importiert_von=importeur,
            )
            if d["buchbar"]:
                buchen(eingang, importeur)
    except Exception:  # noqa: BLE001 — eine fehlerhafte Zeile darf den Rest nicht stoppen
        logger.exception("import zahlungen: Zeile %s nicht gespeichert", zeile.ligne)
        return ZeilenErgebnis(zeile.ligne, ERG_FEHLER, "Technischer Fehler beim Speichern")
    if d["buchbar"]:
        return ZeilenErgebnis(
            zeile.ligne,
            ERG_IMPORTIERT,
            f"Gebucht: {_fmt(d['betrag'])} für {_bezeichnung(zeile.existant)}",
        )
    return ZeilenErgebnis(
        zeile.ligne,
        ERG_IMPORTIERT,
        f"Als offen importiert ({_fmt(d['betrag'])}) — später zuordnen/ergänzen und buchen",
    )


def ausfuehren_zahlungen(fichier, ueberschreiben, dateiname: str = "", importeur=None) -> Ergebnis:
    """BESTÄTIGUNGSPHASE — erneute Auswertung, dann schreiben. `ueberschreiben` wird ignoriert
    (Zahlungen werden nie überschrieben ; Dubletten werden übersprungen)."""
    analyse = analyser_zahlungen(fichier, dateiname)
    ergebnisse = []
    for zeile in analyse.zeilen:
        if zeile.status == STATUS_FEHLER:
            ergebnisse.append(ZeilenErgebnis(zeile.ligne, ERG_FEHLER, zeile.grund))
        elif zeile.status == STATUS_DUBLETTE:
            ergebnisse.append(ZeilenErgebnis(zeile.ligne, ERG_UEBERSPRUNGEN, zeile.grund))
        else:
            ergebnisse.append(_zeile_schreiben(zeile, dateiname, importeur))
    name, inhalt = erzeuge_bericht(analyse.classeur, ergebnisse, dateiname)
    logger.info("import zahlungen: %s Zeilen verarbeitet", len(ergebnisse))
    return Ergebnis(zeilen=ergebnisse, bericht_name=name, bericht_bytes=inhalt)
