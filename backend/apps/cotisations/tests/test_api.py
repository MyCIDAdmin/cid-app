"""
Tests API — app cotisations (TDD §2.4, SCD §2.3 A01 : IDOR sur les endpoints financiers).
"""

import pytest
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.cotisations.models import ArticleCatalogue, Cotisation, StatutCotisation, TypeArticle
from apps.cotisations.tests.factories import ArticleCatalogueFactory, CotisationFactory
from apps.membres.tests.factories import MembreFactory
from apps.notifications.models import Notification, TypeNotification

pytestmark = pytest.mark.django_db


@pytest.fixture
def api_client():
    return APIClient()


def _user_avec_membre(role, email):
    user = User.objects.create_user(email=email, password="Password123!", role=role, is_active=True)
    membre = MembreFactory(user=user)
    return user, membre


def _auth(api_client, user):
    api_client.force_authenticate(user=user)
    return api_client


LIST_URL = "cotisations:cotisation-list"


def _detail_url(cotisation):
    return reverse("cotisations:cotisation-detail", args=[cotisation.id])


# --- Authentification ---


def test_list_non_authentifie_refuse(api_client):
    resp = api_client.get(reverse(LIST_URL))
    assert resp.status_code == 401


# --- Libre-service (F-004 stepper) ---


def test_membre_peut_payer_sa_propre_cotisation(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {"type_article": TypeArticle.COTISATION, "mode_paiement": "carte", "statut": "payee"},
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["membre"]) == str(membre.id)
    assert resp.data["saisie_par"] is None
    assert str(resp.data["montant"]) == "45.00"  # tarif catalogue, pas de montant du client
    assert resp.data["libelle"] == f"Cotisation annuelle {resp.data['annee']}"
    # AHM-53 : jamais fait confiance au statut envoyé par le client en libre-service — reste
    # en_attente tant que le Directeur Financier/Admin ne l'a pas confirmé (marquer_payee).
    assert resp.data["statut"] == StatutCotisation.EN_ATTENTE
    assert resp.data["reference_transaction"] is None
    assert resp.data["date_paiement"] is None


def test_paiement_libre_service_en_attente_notifie_le_directeur_financier(api_client):
    """Ajouté le 2026-09-16 (retour utilisateur : couverture "allen Admin Modulen") — voir
    notifications.notifier_nouveau_paiement_attente_staff : tout Directeur Financier+ est
    notifié dès qu'un paiement libre-service reste en_attente de confirmation manuelle
    (AHM-53), jamais RH/Bureau Admin (niveau insuffisant)."""
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    df = User.objects.create_user(
        email="df@example.de", password="Password123!", role=Role.DIR_FINANCIER, is_active=True
    )
    rh = User.objects.create_user(
        email="rh@example.de", password="Password123!", role=Role.RH, is_active=True
    )
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {"type_article": TypeArticle.COTISATION, "mode_paiement": "carte", "statut": "payee"},
    )

    assert resp.status_code == 201, resp.data
    notification = Notification.objects.get(destinataire=df)
    assert notification.type_notification == TypeNotification.COTISATION_PAIEMENT_ATTENTE
    assert notification.lien == "/cotisations/en-attente"
    assert not Notification.objects.filter(destinataire=rh).exists()


def test_df_qui_saisit_pour_autrui_nest_pas_soi_meme_notifie(api_client):
    """La branche "saisie DF pour un autre membre" n'utilise pas le statut en_attente
    libre-service — pas de notification staff à générer ici (voir views.perform_create)."""
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "df@example.de")
    _autre_user, autre_membre = _user_avec_membre(Role.MEMBRE, "autre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.COTISATION,
            "mode_paiement": "carte",
            "statut": "payee",
            "membre": str(autre_membre.id),
        },
    )

    assert resp.status_code == 201, resp.data
    assert not Notification.objects.filter(
        type_notification=TypeNotification.COTISATION_PAIEMENT_ATTENTE
    ).exists()


def test_statut_libre_service_toujours_en_attente_quel_que_soit_le_mode_de_paiement(api_client):
    # AHM-53 : "je reçois une quittance immédiate pour un virement SEPA" — plus vrai pour aucun
    # mode de paiement en libre-service, aucune passerelle réelle ne pouvant le vérifier.
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    _auth(api_client, user)

    for mode in ("carte", "virement_sepa", "paypal"):
        resp = api_client.post(
            reverse(LIST_URL),
            {
                "type_article": TypeArticle.ADHESION,
                "mode_paiement": mode,
                "statut": "payee",  # ignoré côté serveur
            },
        )
        assert resp.status_code == 201, resp.data
        assert resp.data["statut"] == StatutCotisation.EN_ATTENTE, mode
        assert resp.data["mode_paiement"] == mode
        assert resp.data["reference_transaction"] is None


def test_montant_catalogue_impose_meme_si_client_en_envoie_un_autre(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.COTISATION,
            "montant": "1.00",  # doit être ignoré — tarif catalogue imposé côté serveur
            "mode_paiement": "carte",
            "statut": "payee",
        },
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["montant"]) == "45.00"


# --- Tarifs cotisation/adhésion pilotés par ArticleCatalogue.type_fixe (retour utilisateur du
# 2026-09-17 : "die bestehende [Cotisation annuelle/Frais d'adhésion] müssen auch verwaltbar
# sein") ---


def test_montant_cotisation_suit_le_tarif_configure_par_ladmin(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.COTISATION).update(montant="50.00")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.COTISATION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["montant"]) == "50.00"


def test_montant_adhesion_suit_le_tarif_configure_par_ladmin(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.ADHESION).update(montant="20.00")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.ADHESION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["montant"]) == "20.00"


def test_paiement_cotisation_refuse_si_type_desactive_par_ladmin(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.COTISATION).update(actif=False)
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.COTISATION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 400
    assert "type_article" in resp.data["details"]


def test_paiement_adhesion_refuse_si_type_desactive_par_ladmin(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.ADHESION).update(actif=False)
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.ADHESION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 400
    assert "type_article" in resp.data["details"]


def test_paiement_cotisation_refuse_si_ligne_type_fixe_supprimee(api_client):
    # Corrigé le 2026-09-17 (retour utilisateur répété : "Die Artikel müssen komplett gelöscht
    # werden") — article_catalogue_fixe_actif() est désormais fail-CLOSED : une ligne type_fixe
    # supprimée (ex. via Django Admin, en dehors de l'API) doit rendre le type indisponible, tout
    # comme une désactivation explicite — pas retomber sur l'ancien tarif MONTANTS_CATALOGUE. Voir
    # docstring de la fonction dans models.py pour le raisonnement complet.
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.COTISATION).delete()
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.COTISATION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 400
    assert "type_article" in resp.data["details"]


def test_paiement_adhesion_refuse_si_ligne_type_fixe_supprimee(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    ArticleCatalogue.objects.filter(type_fixe=TypeArticle.ADHESION).delete()
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.ADHESION, "mode_paiement": "carte"}
    )

    assert resp.status_code == 400
    assert "type_article" in resp.data["details"]


def test_don_libre_conserve_le_montant_transmis(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "donateur@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.DON,
            "libelle": "Don libre — soutien projet",
            "montant": "20.00",
            "mode_paiement": "paypal",
            "statut": "payee",
        },
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["montant"]) == "20.00"


def test_don_sans_libelle_refuse(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "donateur@example.de")
    _auth(api_client, user)

    resp = api_client.post(reverse(LIST_URL), {"type_article": TypeArticle.DON, "montant": "20.00"})

    assert resp.status_code == 400
    assert "libelle" in resp.data["details"]


def test_evenement_sans_montant_refuse(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {"type_article": TypeArticle.EVENEMENT, "libelle": "Sortie Dortmund"},
    )

    assert resp.status_code == 400
    assert "montant" in resp.data["details"]


# --- Type "autre_libre" (ajouté le 2026-09-21, retour utilisateur : "Füge noch einen Artikeltyp
# 'anders' mit einem Freitextfeld hinzu") — même mécanique que "don" (libellé/montant libres,
# voir CotisationSerializer.validate), mais réservé à la saisie pour autrui par le Directeur
# Financier/Admin (voir TYPES_ARTICLE_ESPECES côté frontend), jamais au stepper libre-service.


def test_df_peut_saisir_un_paiement_autre_libre_pour_un_autre_membre(api_client):
    user, membre_df = _user_avec_membre(Role.DIR_FINANCIER, "df-libre@example.de")
    _autre_user, autre_membre = _user_avec_membre(Role.MEMBRE, "autre-libre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.AUTRE_LIBRE,
            "libelle": "Remboursement frais essence",
            "montant": "12.50",
            "mode_paiement": "especes",
            "statut": "payee",
            "membre": str(autre_membre.id),
        },
    )

    assert resp.status_code == 201, resp.data
    assert resp.data["libelle"] == "Remboursement frais essence"
    assert str(resp.data["montant"]) == "12.50"
    assert str(resp.data["saisie_par"]) == str(membre_df.id)


def test_autre_libre_sans_libelle_refuse(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre-libre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.AUTRE_LIBRE, "montant": "12.50"}
    )

    assert resp.status_code == 400
    assert "libelle" in resp.data["details"]


# --- Article catalogue, type_article=autre (retour utilisateur du 2026-09-17) ---


def test_paiement_article_catalogue_recalcule_libelle_et_montant_cote_serveur(api_client):
    """CLAUDE.md §8 : comme pour cotisation/adhésion ci-dessus, le libellé/montant d'un article du
    catalogue est toujours recalculé côté serveur à partir de l'ArticleCatalogue référencé, jamais
    fait confiance au client — voir CotisationSerializer.validate, branche TypeArticle.AUTRE."""
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    article = ArticleCatalogueFactory(libelle="T-shirt du club", montant="20.00")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.AUTRE,
            "article_catalogue": str(article.id),
            "libelle": "Faux libellé",  # doit être ignoré
            "montant": "1.00",  # doit être ignoré
            "mode_paiement": "carte",
            "statut": "payee",
        },
    )

    assert resp.status_code == 201, resp.data
    assert resp.data["libelle"] == "T-shirt du club"
    assert str(resp.data["montant"]) == "20.00"
    assert resp.data["statut"] == StatutCotisation.EN_ATTENTE


def test_article_catalogue_sans_id_refuse(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.AUTRE, "mode_paiement": "carte"}
    )

    assert resp.status_code == 400
    assert "article_catalogue" in resp.data["details"]


def test_article_catalogue_desactive_refuse(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    article = ArticleCatalogueFactory(actif=False)
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.AUTRE,
            "article_catalogue": str(article.id),
            "mode_paiement": "carte",
        },
    )

    assert resp.status_code == 400
    assert "article_catalogue" in resp.data["details"]


def test_creation_sans_fiche_membre_liee_refusee(api_client):
    user = User.objects.create_user(
        email="sans-fiche@example.de", password="Password123!", role=Role.MEMBRE, is_active=True
    )
    _auth(api_client, user)

    resp = api_client.post(reverse(LIST_URL), {"type_article": TypeArticle.COTISATION})
    assert resp.status_code == 400


# --- Saisie pour autrui (F-015) ---


def test_membre_ne_peut_pas_saisir_pour_un_autre_membre(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "membre@example.de")
    autre = MembreFactory()
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL), {"type_article": TypeArticle.COTISATION, "membre": str(autre.id)}
    )
    assert resp.status_code == 403


def test_directeur_financier_peut_saisir_pour_un_autre_membre(api_client):
    user, _membre_dg = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    autre = MembreFactory()
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.ADHESION,
            "mode_paiement": "virement_sepa",
            "statut": "payee",
            "membre": str(autre.id),
        },
    )

    assert resp.status_code == 201, resp.data
    assert str(resp.data["membre"]) == str(autre.id)
    assert str(resp.data["saisie_par"]) == str(_membre_dg.id)


# --- Enregistrer un paiement en espèces (retour utilisateur du 2026-09-21, "Es soll möglich sein
# eine Zahlung als Barzahlung einzutragen") ---


def test_df_peut_enregistrer_un_paiement_en_especes_deja_payee_pour_un_autre_membre(api_client):
    user, membre_df = _user_avec_membre(Role.DIR_FINANCIER, "df@example.de")
    autre_user, autre_membre = _user_avec_membre(Role.MEMBRE, "autre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.ADHESION,
            "mode_paiement": "especes",
            "statut": "payee",
            "membre": str(autre_membre.id),
        },
    )

    assert resp.status_code == 201, resp.data
    assert resp.data["mode_paiement"] == "especes"
    assert resp.data["statut"] == StatutCotisation.PAYEE
    assert resp.data["reference_transaction"] is not None
    assert str(resp.data["saisie_par"]) == str(membre_df.id)
    # Effet de bord attendu, même principe que marquer_payee (voir
    # views.CotisationViewSet.perform_create) : le membre concerné reçoit la notification
    # "paiement confirmé", pas seulement le DF qui a saisi la transaction.
    notification = Notification.objects.get(destinataire=autre_user)
    assert notification.type_notification == TypeNotification.PAIEMENT_CONFIRME


def test_saisie_especes_en_attente_ne_notifie_pas_de_paiement_confirme(api_client):
    """Une transaction saisie en espèces mais laissée en_attente/echouee (cas rare mais possible,
    ex. correction manuelle ultérieure prévue) ne doit pas déclencher la notification "paiement
    confirmé" — seul un statut payee le doit (voir perform_create)."""
    user, _membre_df = _user_avec_membre(Role.DIR_FINANCIER, "df@example.de")
    autre_user, autre_membre = _user_avec_membre(Role.MEMBRE, "autre@example.de")
    _auth(api_client, user)

    resp = api_client.post(
        reverse(LIST_URL),
        {
            "type_article": TypeArticle.ADHESION,
            "mode_paiement": "especes",
            "statut": "en_attente",
            "membre": str(autre_membre.id),
        },
    )

    assert resp.status_code == 201, resp.data
    assert not Notification.objects.filter(
        destinataire=autre_user, type_notification=TypeNotification.PAIEMENT_CONFIRME
    ).exists()


# --- Scope liste / IDOR (SCD §2.3 A01) ---


def test_membre_ne_voit_que_ses_propres_cotisations(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    CotisationFactory(membre=membre)
    CotisationFactory()  # une autre fiche, non liée à ce compte

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL))

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert str(resp.data["results"][0]["membre"]) == str(membre.id)


def test_role_personnalise_eleve_voit_toutes_les_cotisations(api_client):
    """apps.rbac Phase B : un rôle personnalisé avec au moins la lecture sur "cotisations" voit
    TOUTES les cotisations, comme RH/Admin — voir is_elevated_for_module."""
    from apps.rbac.models import NiveauAcces
    from apps.rbac.tests.factories import (
        RoleDefinitionFactory,
        RoleModulePermissionFactory,
        UserRoleAssignmentFactory,
    )

    user, membre = _user_avec_membre(Role.MEMBRE, "vertrieb-idor@example.de")
    role = RoleDefinitionFactory(slug="vertrieb-cotisations-idor")
    RoleModulePermissionFactory(role=role, module="cotisations", niveau_acces=NiveauAcces.LECTURE)
    UserRoleAssignmentFactory(user=user, role=role)
    CotisationFactory(membre=membre)
    CotisationFactory()  # une autre fiche, sans lien avec `user`

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL))

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 2


def test_membre_sans_role_eleve_ne_voit_toujours_que_ses_propres_cotisations(api_client):
    """Régression explicite après le câblage Phase B : sans UserRoleAssignment
    supplémentaire, le comportement historique (SCD §2.3 A01) n'a pas bougé."""
    user, membre = _user_avec_membre(Role.MEMBRE, "membre-seul-idor@example.de")
    CotisationFactory(membre=membre)
    CotisationFactory()

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL))

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert str(resp.data["results"][0]["membre"]) == str(membre.id)


def test_membre_peut_recuperer_sa_propre_cotisation(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation = CotisationFactory(membre=membre)

    _auth(api_client, user)
    resp = api_client.get(_detail_url(cotisation))

    assert resp.status_code == 200
    assert str(resp.data["id"]) == str(cotisation.id)


def test_membre_ne_peut_pas_recuperer_la_cotisation_dun_autre(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation_autrui = CotisationFactory()

    _auth(api_client, user)
    resp = api_client.get(_detail_url(cotisation_autrui))

    assert resp.status_code == 404


def test_rh_peut_recuperer_le_detail_dune_cotisation_dautrui(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    cotisation = CotisationFactory()

    _auth(api_client, user)
    resp = api_client.get(_detail_url(cotisation))

    assert resp.status_code == 200
    assert str(resp.data["id"]) == str(cotisation.id)


def test_rh_liste_toutes_les_cotisations(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    CotisationFactory.create_batch(3)

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL))

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 3


# --- Registre append-only ---


def test_update_non_autorise(api_client):
    user, membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(membre=membre)

    _auth(api_client, user)
    resp = api_client.patch(_detail_url(cotisation), {"statut": StatutCotisation.ANNULEE})

    assert resp.status_code == 405


def test_delete_non_autorise(api_client):
    user, membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(membre=membre)

    _auth(api_client, user)
    resp = api_client.delete(_detail_url(cotisation))

    assert resp.status_code == 405


# --- Filtres ---


def test_filtre_par_statut(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    CotisationFactory(statut=StatutCotisation.PAYEE)
    CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL), {"statut": StatutCotisation.EN_ATTENTE})

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert resp.data["results"][0]["statut"] == StatutCotisation.EN_ATTENTE


# --- Filtres ajoutés le 2026-09-21 (page "Ausstehende Zahlungen", retour utilisateur : "Filter
# Möglichkeiten hinzufügen") ---


def test_filtre_par_mode_paiement(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    CotisationFactory(mode_paiement="especes", statut=StatutCotisation.PAYEE)
    CotisationFactory(mode_paiement="carte", statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL), {"mode_paiement": "especes"})

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert resp.data["results"][0]["mode_paiement"] == "especes"


def test_filtre_q_recherche_par_nom_de_membre(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    cible = MembreFactory(nom="Trabelsi", prenom="Sami")
    CotisationFactory(membre=cible)
    CotisationFactory()  # un autre membre, ne doit pas remonter

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL), {"q": "trabelsi"})

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert str(resp.data["results"][0]["membre"]) == str(cible.id)


def test_filtre_q_recherche_par_libelle(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    CotisationFactory(libelle="Don libre exceptionnel")
    CotisationFactory(libelle="Frais d'adhésion")

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL), {"q": "don libre"})

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert resp.data["results"][0]["libelle"] == "Don libre exceptionnel"


def test_filtre_par_date_de_creation(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    ancienne = CotisationFactory()
    recente = CotisationFactory()

    Cotisation.objects.filter(pk=ancienne.pk).update(created_at="2020-01-01T10:00:00Z")
    Cotisation.objects.filter(pk=recente.pk).update(created_at="2026-06-15T10:00:00Z")

    _auth(api_client, user)
    resp = api_client.get(reverse(LIST_URL), {"date_creation_apres": "2025-01-01"})

    assert resp.status_code == 200
    assert len(resp.data["results"]) == 1
    assert str(resp.data["results"][0]["id"]) == str(recente.id)


# --- Reçu PDF (AHM-17, RICEFW R-010) ---


def _receipt_url(cotisation):
    return reverse("cotisations:cotisation-receipt", args=[cotisation.id])


def test_receipt_non_authentifie_refuse(api_client):
    cotisation = CotisationFactory()
    resp = api_client.get(_receipt_url(cotisation))
    assert resp.status_code == 401


def test_receipt_disponible_pour_le_proprietaire(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation = CotisationFactory(membre=membre, statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.get(_receipt_url(cotisation))

    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/pdf"
    assert resp["Content-Disposition"] == (
        f'attachment; filename="recu-{cotisation.reference_transaction}.pdf"'
    )
    assert resp.content.startswith(b"%PDF-")


def test_receipt_refuse_pour_la_cotisation_dun_autre_membre(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation_autrui = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.get(_receipt_url(cotisation_autrui))

    assert resp.status_code == 404  # IDOR : ne révèle même pas l'existence de la ressource


def test_rh_peut_telecharger_le_recu_dun_autre_membre(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.get(_receipt_url(cotisation))

    assert resp.status_code == 200
    assert resp.content.startswith(b"%PDF-")


def test_receipt_refuse_si_cotisation_pas_encore_payee(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation = CotisationFactory(
        membre=membre, statut=StatutCotisation.EN_ATTENTE, reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.get(_receipt_url(cotisation))

    assert resp.status_code == 400


# --- Confirmation manuelle de paiement (AHM-53) ---


def _marquer_payee_url(cotisation):
    return reverse("cotisations:cotisation-marquer-payee", args=[cotisation.id])


def test_marquer_payee_non_authentifie_refuse(api_client):
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)
    resp = api_client.post(_marquer_payee_url(cotisation))
    assert resp.status_code == 401


def test_directeur_financier_peut_marquer_payee(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE, mode_paiement="", reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})

    assert resp.status_code == 200, resp.data
    assert resp.data["statut"] == StatutCotisation.PAYEE
    assert resp.data["mode_paiement"] == "virement_sepa"
    assert resp.data["reference_transaction"].startswith("TXN-")
    assert resp.data["date_paiement"] is not None


def test_marquer_payee_cree_une_notification_in_app_pour_le_membre(api_client):
    from apps.notifications.models import Notification, TypeNotification

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg2@example.de")
    user_paye = User.objects.create_user(
        email="paye@example.de", password="Password123!", is_active=True
    )
    membre_paye = MembreFactory(user=user_paye)
    cotisation = CotisationFactory(
        membre=membre_paye,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    notification = Notification.objects.get(destinataire=user_paye)
    assert notification.type_notification == TypeNotification.PAIEMENT_CONFIRME
    assert notification.lien == "/cotisation"


# Statut associatif automatique (demande utilisateur du 2026-09-19) : confirmer une cotisation
# ANNUELLE réactive automatiquement le membre — voir
# apps.membres.services.enregistrer_statut_annuel, déclenché depuis
# apps.cotisations.notifications.notifier_paiement_confirme.
def test_marquer_payee_dune_cotisation_annuelle_reactive_le_membre(api_client):
    from apps.membres.models import HistoriqueStatutMembre, RaisonChangementStatut, StatutMembre

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg3@example.de")
    user_paye = User.objects.create_user(
        email="paye2@example.de", password="Password123!", is_active=True
    )
    membre_paye = MembreFactory(user=user_paye, statut=StatutMembre.INACTIF)
    cotisation = CotisationFactory(
        membre=membre_paye,
        type_article=TypeArticle.COTISATION,
        annee=2027,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    membre_paye.refresh_from_db()
    assert membre_paye.statut == StatutMembre.ACTIF
    historique = HistoriqueStatutMembre.objects.get(membre=membre_paye, annee=2027)
    assert historique.statut == StatutMembre.ACTIF
    assert historique.raison == RaisonChangementStatut.PAIEMENT_CONFIRME


def test_marquer_payee_dune_adhesion_active_le_membre(api_client):
    # Demande utilisateur du 2026-10-05 (point 4) : adhésion approuvée et payée -> le compte
    # devient immédiatement membre actif (remplace l'ancien comportement "statut inchangé").
    from apps.membres.models import HistoriqueStatutMembre, StatutMembre

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg4@example.de")
    user_paye = User.objects.create_user(
        email="paye3@example.de", password="Password123!", is_active=True
    )
    membre_paye = MembreFactory(user=user_paye, statut=StatutMembre.INACTIF)
    cotisation = CotisationFactory(
        membre=membre_paye,
        type_article=TypeArticle.ADHESION,
        annee=2027,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    membre_paye.refresh_from_db()
    assert membre_paye.statut == StatutMembre.ACTIF
    assert HistoriqueStatutMembre.objects.filter(
        membre=membre_paye, statut=StatutMembre.ACTIF
    ).exists()


def test_marquer_payee_dune_souscription_dadhesion_marque_la_souscription_payee(api_client):
    """
    Ajouté le 2026-09-19 (retour utilisateur : "Die Zahlung taucht nicht im Modul Ausstehende
    Zahlungen") — cascade Cotisation payee -> Souscription payee, voir
    apps.cotisations.notifications.notifier_paiement_confirme et
    apps.adhesions.services.synchroniser_cotisation (qui crée la Cotisation liée).
    """
    from apps.adhesions.models import StatutSouscription
    from apps.adhesions.tests.factories import SouscriptionFactory

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg5@example.de")
    membre_paye = MembreFactory()
    cotisation = CotisationFactory(
        membre=membre_paye,
        type_article=TypeArticle.ADHESION,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )
    souscription = SouscriptionFactory(
        membre=membre_paye,
        cotisation=cotisation,
        statut=StatutSouscription.EN_ATTENTE_PAIEMENT,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    souscription.refresh_from_db()
    assert souscription.statut == StatutSouscription.PAYEE


def test_marquer_payee_dune_inscription_evenement_marque_linscription_confirmee(api_client):
    """
    Ajouté le 2026-09-20 (retour utilisateur : "Confirmer et payer" doit sauter directement au
    paiement) — cascade Cotisation payee -> Inscription confirmee, voir
    apps.cotisations.notifications.notifier_paiement_confirme et
    apps.evenements.services.synchroniser_cotisation (qui crée la Cotisation liée).
    """
    from apps.evenements.models import StatutInscription
    from apps.evenements.tests.factories import InscriptionFactory

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg6@example.de")
    membre_paye = MembreFactory()
    cotisation = CotisationFactory(
        membre=membre_paye,
        type_article=TypeArticle.EVENEMENT,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )
    inscription = InscriptionFactory(
        membre=membre_paye,
        cotisation=cotisation,
        statut=StatutInscription.EN_ATTENTE_PAIEMENT,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    inscription.refresh_from_db()
    assert inscription.statut == StatutInscription.CONFIRMEE


def test_admin_peut_marquer_payee(api_client):
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin@example.de")
    cotisation = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE, mode_paiement="", reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})

    assert resp.status_code == 200, resp.data
    assert resp.data["statut"] == StatutCotisation.PAYEE


def test_marquer_payee_conserve_le_mode_de_paiement_existant_si_non_precise(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="virement_sepa",
        reference_transaction=None,
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation))

    assert resp.status_code == 200, resp.data
    assert resp.data["mode_paiement"] == "virement_sepa"


def test_marquer_payee_refuse_si_mode_de_paiement_absent(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE, mode_paiement="", reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation))

    assert resp.status_code == 400
    assert "mode_paiement" in resp.data["details"]


def test_rh_ne_peut_pas_marquer_payee(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})

    assert resp.status_code == 403


def test_membre_ne_peut_pas_marquer_sa_propre_cotisation_payee(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation = CotisationFactory(
        membre=membre, statut=StatutCotisation.EN_ATTENTE, reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})

    assert resp.status_code == 403


def test_marquer_payee_refuse_pour_la_cotisation_dun_autre_membre_hors_scope(api_client):
    # IDOR : un Membre normal (rôle < RH) ne peut même pas voir la ressource d'un autre membre —
    # get_object() renvoie 404 avant que la vérification de rôle DF/Admin ne soit atteinte.
    user, _membre = _user_avec_membre(Role.MEMBRE, "a@example.de")
    cotisation_autrui = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE, reference_transaction=None
    )

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation_autrui), {"mode_paiement": "carte"})

    assert resp.status_code == 404


def test_marquer_payee_refuse_si_deja_payee(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "carte"})

    assert resp.status_code == 400


def test_marquer_payee_refuse_si_annulee(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.ANNULEE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "carte"})

    assert resp.status_code == 400


def test_marquer_payee_refuse_mode_de_paiement_invalide(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "bitcoin"})

    assert resp.status_code == 400


def test_marquer_payee_journalise_lhistorique(api_client):
    """Ajouté le 2026-09-19 — marquer_payee alimente aussi HistoriqueStatutCotisation, même
    point d'écriture que changer_statut ci-dessous, mais sans motif (transition standard)."""
    from apps.cotisations.models import HistoriqueStatutCotisation

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-hist@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data

    entree = HistoriqueStatutCotisation.objects.get(cotisation=cotisation)
    assert entree.ancien_statut == StatutCotisation.EN_ATTENTE
    assert entree.nouveau_statut == StatutCotisation.PAYEE
    assert entree.motif == ""


# --- date_paiement (transaction backdatée, demande utilisateur du 2026-09-29 : "Bei
# Zahlungsbestätigung Im Modul 'Zahlungen' [...] das Transaktionsdatum bei der Bestätigung
# hinzufügen") --------------------------------------------------------------------------


def test_marquer_payee_avec_date_paiement_backdate_la_transaction(api_client):
    from datetime import date

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-date1@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(
        _marquer_payee_url(cotisation),
        {"mode_paiement": "virement_sepa", "date_paiement": "2026-08-15"},
    )

    assert resp.status_code == 200, resp.data
    cotisation.refresh_from_db()
    # localtime() : date_paiement est stocké en UTC (USE_TZ=True) — _debut_jour_aware
    # construit minuit dans le fuseau LOCAL (Europe/Berlin), donc un .date() direct sur la
    # valeur UTC peut retomber sur la veille selon le fuseau.
    from django.utils import timezone

    assert timezone.localtime(cotisation.date_paiement).date() == date(2026, 8, 15)


def test_marquer_payee_sans_date_paiement_utilise_aujourdhui(api_client):
    """Comportement inchangé quand `date_paiement` est absent — même résultat qu'avant
    l'ajout de ce champ (voir Cotisation.save())."""
    from django.utils import timezone

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-date2@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})

    assert resp.status_code == 200, resp.data
    cotisation.refresh_from_db()
    assert timezone.localtime(cotisation.date_paiement).date() == timezone.localdate()


def test_marquer_payee_avec_date_du_jour_garde_lheure_reelle(api_client):
    """Das vorbelegte Datum "heute" darf nicht auf 00:00 zurückfallen (Reihenfolge der Belege)."""
    from django.utils import timezone

    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-date3@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(
        _marquer_payee_url(cotisation),
        {"mode_paiement": "virement_sepa", "date_paiement": timezone.localdate().isoformat()},
    )

    assert resp.status_code == 200, resp.data
    cotisation.refresh_from_db()
    assert timezone.localtime(cotisation.date_paiement).date() == timezone.localdate()
    assert (timezone.now() - cotisation.date_paiement).total_seconds() < 120


def test_marquer_payee_refuse_date_paiement_future(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-date3@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(
        _marquer_payee_url(cotisation),
        {"mode_paiement": "virement_sepa", "date_paiement": "2099-01-01"},
    )

    assert resp.status_code == 400
    assert "date_paiement" in resp.data["details"]


# --- changer-statut (ajouté le 2026-09-19, correction rétroactive + historique complet) ---


def _changer_statut_url(cotisation):
    return reverse("cotisations:cotisation-changer-statut", args=[cotisation.id])


def _historique_statuts_url(cotisation):
    return reverse("cotisations:cotisation-historique-statuts", args=[cotisation.id])


def test_changer_statut_non_authentifie_refuse(api_client):
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "remboursee"})
    assert resp.status_code == 401


def test_admin_peut_revenir_dun_statut_paye_a_un_autre_statut(api_client):
    """Cœur de la demande utilisateur : contrairement à marquer_payee (restreint
    EN_ATTENTE/ECHOUEE -> PAYEE), changer_statut autorise à revenir en arrière depuis PAYEE."""
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(
        _changer_statut_url(cotisation),
        {"statut": "remboursee", "motif": "Paiement finalement rejeté par la banque"},
    )

    assert resp.status_code == 200, resp.data
    assert resp.data["statut"] == StatutCotisation.REMBOURSEE

    cotisation.refresh_from_db()
    assert cotisation.statut == StatutCotisation.REMBOURSEE


def test_directeur_financier_peut_changer_statut(api_client):
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "dg-cs@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "annulee"})

    assert resp.status_code == 200, resp.data
    assert resp.data["statut"] == StatutCotisation.ANNULEE


def test_rh_ne_peut_pas_changer_statut(api_client):
    user, _membre = _user_avec_membre(Role.RH, "rh-cs@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "annulee"})

    assert resp.status_code == 403


def test_membre_ne_peut_pas_changer_le_statut_de_sa_propre_cotisation(api_client):
    user, membre = _user_avec_membre(Role.MEMBRE, "membre-cs@example.de")
    cotisation = CotisationFactory(membre=membre, statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "annulee"})

    assert resp.status_code == 403


def test_changer_statut_refuse_statut_invalide(api_client):
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs2@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "inexistant"})

    assert resp.status_code == 400


def test_changer_statut_refuse_si_meme_statut(api_client):
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs3@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "en_attente"})

    assert resp.status_code == 400


def test_changer_statut_vers_payee_notifie_comme_marquer_payee(api_client):
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs4@example.de")
    user_paye = User.objects.create_user(
        email="paye-cs@example.de", password="Password123!", is_active=True
    )
    membre_paye = MembreFactory(user=user_paye)
    cotisation = CotisationFactory(membre=membre_paye, statut=StatutCotisation.EN_ATTENTE)

    _auth(api_client, user)
    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "payee"})
    assert resp.status_code == 200, resp.data

    notification = Notification.objects.get(destinataire=user_paye)
    assert notification.type_notification == TypeNotification.PAIEMENT_CONFIRME


def test_changer_statut_journalise_le_motif_et_lauteur(api_client):
    from apps.cotisations.models import HistoriqueStatutCotisation

    user, membre_admin = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs5@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(
        _changer_statut_url(cotisation), {"statut": "annulee", "motif": "Erreur de saisie"}
    )
    assert resp.status_code == 200, resp.data

    entree = HistoriqueStatutCotisation.objects.get(cotisation=cotisation)
    assert entree.ancien_statut == StatutCotisation.PAYEE
    assert entree.nouveau_statut == StatutCotisation.ANNULEE
    assert entree.motif == "Erreur de saisie"
    assert entree.modifie_par_id == membre_admin.id


def test_changer_statut_vers_payee_avec_date_paiement_backdate_la_transaction(api_client):
    from datetime import date

    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs-date1@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(
        _changer_statut_url(cotisation), {"statut": "payee", "date_paiement": "2026-07-01"}
    )
    assert resp.status_code == 200, resp.data

    from django.utils import timezone

    cotisation.refresh_from_db()
    assert timezone.localtime(cotisation.date_paiement).date() == date(2026, 7, 1)


def test_changer_statut_ignore_date_paiement_hors_transition_payee(api_client):
    """`date_paiement` n'a de sens que pour statut=payee (voir docstring views.py) —
    silencieusement ignoré pour toute autre transition, aucune erreur levée."""
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs-date2@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.post(
        _changer_statut_url(cotisation),
        {"statut": "annulee", "date_paiement": "2026-07-01"},
    )
    assert resp.status_code == 200, resp.data


def test_changer_statut_refuse_date_paiement_future(api_client):
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "admin-cs-date3@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE, reference_transaction=None)

    _auth(api_client, user)
    resp = api_client.post(
        _changer_statut_url(cotisation), {"statut": "payee", "date_paiement": "2099-01-01"}
    )
    assert resp.status_code == 400
    assert "date_paiement" in resp.data["details"]


def test_historique_statuts_visible_par_le_proprietaire(api_client):
    from apps.cotisations.models import HistoriqueStatutCotisation

    user, membre = _user_avec_membre(Role.MEMBRE, "hist-owner@example.de")
    cotisation = CotisationFactory(membre=membre, statut=StatutCotisation.PAYEE)
    HistoriqueStatutCotisation.objects.create(
        cotisation=cotisation,
        ancien_statut=StatutCotisation.EN_ATTENTE,
        nouveau_statut=StatutCotisation.PAYEE,
    )

    _auth(api_client, user)
    resp = api_client.get(_historique_statuts_url(cotisation))

    assert resp.status_code == 200
    assert len(resp.data) == 1
    assert resp.data[0]["nouveau_statut"] == StatutCotisation.PAYEE


def test_historique_statuts_refuse_pour_la_cotisation_dun_autre_membre(api_client):
    user, _membre = _user_avec_membre(Role.MEMBRE, "hist-autrui@example.de")
    cotisation_autrui = CotisationFactory(statut=StatutCotisation.PAYEE)

    _auth(api_client, user)
    resp = api_client.get(_historique_statuts_url(cotisation_autrui))

    assert resp.status_code in (403, 404)


# ---------------------------------------------------------------------------
# Phase D (ajoutée le 2026-09-23) — page de gestion "Ausstehende Zahlungen"
# (page_cotisations_attente) désormais pilotée par apps.rbac (real enforcement, y compris pour
# les rôles système eux-mêmes) — couvre marquer_payee ET changer_statut, la même page.
# ---------------------------------------------------------------------------


def _set_matrice_cellule(role_slug, module_slug, niveau_acces):
    from apps.rbac.models import RoleDefinition, RoleModulePermission

    role = RoleDefinition.objects.get(slug=role_slug, is_system=True)
    RoleModulePermission.objects.update_or_create(
        role=role, module=module_slug, defaults={"niveau_acces": niveau_acces}
    )


def test_phase_d_directeur_financier_perd_lacces_a_marquer_payee_si_matrice_le_dit(api_client):
    from apps.rbac.models import NiveauAcces

    _set_matrice_cellule("dir_financier", "page_cotisations_attente", NiveauAcces.AUCUN)
    user, _membre = _user_avec_membre(Role.DIR_FINANCIER, "phased-attente-restrict@example.de")
    cotisation = CotisationFactory(statut=StatutCotisation.EN_ATTENTE)
    _auth(api_client, user)

    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 403


def test_phase_d_directeur_financier_perd_lacces_a_changer_statut_si_matrice_le_dit(api_client):
    from apps.rbac.models import NiveauAcces

    _set_matrice_cellule("dir_financier", "page_cotisations_attente", NiveauAcces.AUCUN)
    user, _membre = _user_avec_membre(
        Role.DIR_FINANCIER, "phased-attente-statut-restrict@example.de"
    )
    cotisation = CotisationFactory(statut=StatutCotisation.PAYEE)
    _auth(api_client, user)

    resp = api_client.post(_changer_statut_url(cotisation), {"statut": "annulee"})
    assert resp.status_code == 403


def test_phase_d_role_personnalise_peut_marquer_payee_via_la_matrice(api_client):
    """Vérifie le gate de page apporté par has_admin_page_access("page_cotisations_attente").
    La visibilité au niveau du queryset (get_queryset de CotisationViewSet) reste pilotée par
    is_elevated_for_module(user, "cotisations") — le module DATA "cotisations", pas la page
    admin "page_cotisations_attente" — donc explicitement hors scope de cette phase (voir plan).
    Comme pour les justificatifs, la cotisation ciblée appartient donc à la propre fiche membre
    du user pour rester dans le queryset filtré "mes cotisations"."""
    from apps.rbac.models import NiveauAcces
    from apps.rbac.tests.factories import (
        RoleDefinitionFactory,
        RoleModulePermissionFactory,
        UserRoleAssignmentFactory,
    )

    user, membre = _user_avec_membre(Role.MEMBRE, "phased-attente-grant@example.de")
    role_perso = RoleDefinitionFactory(slug="cotisations-attente-manager")
    RoleModulePermissionFactory(
        role=role_perso,
        module="page_cotisations_attente",
        niveau_acces=NiveauAcces.LECTURE_ECRITURE,
    )
    UserRoleAssignmentFactory(user=user, role=role_perso)
    cotisation = CotisationFactory(
        membre=membre,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )
    _auth(api_client, user)

    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data


def test_phase_d_super_admin_marque_toujours_payee_meme_si_matrice_dit_aucun(api_client):
    from apps.rbac.models import NiveauAcces

    _set_matrice_cellule("super_admin", "page_cotisations_attente", NiveauAcces.AUCUN)
    user, _membre = _user_avec_membre(Role.SUPER_ADMIN, "phased-attente-super@example.de")
    cotisation = CotisationFactory(
        statut=StatutCotisation.EN_ATTENTE, mode_paiement="", reference_transaction=None
    )
    _auth(api_client, user)

    resp = api_client.post(_marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"})
    assert resp.status_code == 200, resp.data


# ---------------------------------------------------------------------------
# Lecture vs écriture (ajouté le 2026-09-24, task #214, retour utilisateur sur Quiz-Verwaltung —
# voir apps.rbac.services.has_admin_page_access) : marquer_payee/changer_statut n'ont pas
# d'équivalent "lecture" gaté par page_cotisations_attente (list/retrieve de CotisationViewSet
# restent gouvernés par CotisationPermission + le module DATA "cotisations", pas cette page —
# voir docstring de test_phase_d_role_personnalise_peut_marquer_payee_via_la_matrice ci-dessus).
# Ces tests vérifient donc uniquement qu'une cellule `lecture` seule sur page_cotisations_attente
# ne débloque plus les deux actions d'écriture de cette page.
# ---------------------------------------------------------------------------


def test_role_lecture_seule_ne_peut_pas_marquer_payee_ni_changer_statut(api_client):
    from apps.rbac.models import NiveauAcces
    from apps.rbac.tests.factories import (
        RoleDefinitionFactory,
        RoleModulePermissionFactory,
        UserRoleAssignmentFactory,
    )

    user, membre = _user_avec_membre(Role.MEMBRE, "readonly-attente@example.de")
    role_perso = RoleDefinitionFactory(slug="attente-lecteur")
    RoleModulePermissionFactory(
        role=role_perso, module="page_cotisations_attente", niveau_acces=NiveauAcces.LECTURE
    )
    UserRoleAssignmentFactory(user=user, role=role_perso)
    cotisation = CotisationFactory(
        membre=membre,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )
    _auth(api_client, user)

    resp_marquer = api_client.post(
        _marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"}
    )
    assert resp_marquer.status_code == 403

    resp_changer = api_client.post(_changer_statut_url(cotisation), {"statut": "annulee"})
    assert resp_changer.status_code == 403


def test_role_lecture_ecriture_peut_marquer_payee_et_changer_statut(api_client):
    from apps.rbac.models import NiveauAcces
    from apps.rbac.tests.factories import (
        RoleDefinitionFactory,
        RoleModulePermissionFactory,
        UserRoleAssignmentFactory,
    )

    user, membre = _user_avec_membre(Role.MEMBRE, "readwrite-attente@example.de")
    role_perso = RoleDefinitionFactory(slug="attente-editeur")
    RoleModulePermissionFactory(
        role=role_perso,
        module="page_cotisations_attente",
        niveau_acces=NiveauAcces.LECTURE_ECRITURE,
    )
    UserRoleAssignmentFactory(user=user, role=role_perso)
    cotisation = CotisationFactory(
        membre=membre,
        statut=StatutCotisation.EN_ATTENTE,
        mode_paiement="",
        reference_transaction=None,
    )
    _auth(api_client, user)

    resp_marquer = api_client.post(
        _marquer_payee_url(cotisation), {"mode_paiement": "virement_sepa"}
    )
    assert resp_marquer.status_code == 200, resp_marquer.data

    autre_cotisation = CotisationFactory(membre=membre, statut=StatutCotisation.PAYEE)
    resp_changer = api_client.post(_changer_statut_url(autre_cotisation), {"statut": "annulee"})
    assert resp_changer.status_code == 200, resp_changer.data


# ---------------------------------------------------------------------------
# Export Excel des paiements (module "Ausstehende Zahlungen"/"Zahlungen", ajouté le 2026-09-25,
# demande utilisateur : "Excel-Export der Zahlungen") — même principe que
# apps.boutique.tests.test_api::test_export_commandes_* (même scope IDOR/filtres que la liste).
# ---------------------------------------------------------------------------


def _export_url():
    return reverse("cotisations:cotisation-export")


def test_export_non_authentifie_refuse(api_client):
    resp = api_client.get(_export_url())
    assert resp.status_code == 401


def test_export_retourne_un_classeur_xlsx(api_client):
    import io

    from openpyxl import load_workbook

    admin, _ = _user_avec_membre(Role.BUREAU_ADMIN, "export-cot-admin1@example.de")
    _, membre1 = _user_avec_membre(Role.MEMBRE, "export-cot-m1@example.de")
    _, membre2 = _user_avec_membre(Role.MEMBRE, "export-cot-m2@example.de")
    CotisationFactory(membre=membre1, libelle="Export Un")
    CotisationFactory(membre=membre2, libelle="Export Deux")

    resp = _auth(api_client, admin).get(_export_url())
    assert resp.status_code == 200
    assert resp["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    classeur = load_workbook(io.BytesIO(resp.content))
    assert classeur.sheetnames == ["Zahlungen"]
    feuille = classeur.active
    entetes = [c.value for c in feuille[1]]
    assert entetes[0] == "Membre"
    libelles = {feuille.cell(row=r, column=5).value for r in (2, 3)}
    assert libelles == {"Export Un", "Export Deux"}


def test_export_respecte_le_filtre_statut(api_client):
    import io

    from openpyxl import load_workbook

    admin, _ = _user_avec_membre(Role.BUREAU_ADMIN, "export-cot-admin2@example.de")
    _, membre = _user_avec_membre(Role.MEMBRE, "export-cot-m3@example.de")
    CotisationFactory(
        membre=membre,
        statut=StatutCotisation.EN_ATTENTE,
        libelle="Attente",
        reference_transaction=None,
    )
    CotisationFactory(membre=membre, statut=StatutCotisation.PAYEE, libelle="Payee")

    resp = _auth(api_client, admin).get(_export_url(), {"statut": StatutCotisation.EN_ATTENTE})
    assert resp.status_code == 200

    classeur = load_workbook(io.BytesIO(resp.content))
    feuille = classeur.active
    assert feuille.max_row == 2  # en-tête + 1 seule cotisation
    assert feuille.cell(row=2, column=5).value == "Attente"


def test_export_scope_membre_ne_voit_que_les_siennes(api_client):
    import io

    from openpyxl import load_workbook

    user, membre1 = _user_avec_membre(Role.MEMBRE, "export-cot-m4@example.de")
    _, membre2 = _user_avec_membre(Role.MEMBRE, "export-cot-m5@example.de")
    CotisationFactory(membre=membre1, libelle="Mine")
    CotisationFactory(membre=membre2, libelle="PasMoi")

    resp = _auth(api_client, user).get(_export_url())
    assert resp.status_code == 200

    classeur = load_workbook(io.BytesIO(resp.content))
    feuille = classeur.active
    assert feuille.max_row == 2
    assert feuille.cell(row=2, column=5).value == "Mine"
