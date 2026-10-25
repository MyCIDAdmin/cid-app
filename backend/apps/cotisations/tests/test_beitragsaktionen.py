"""Tests — Beitragsaktionen (Wunsch Ghazi, 2026-10-25) : Beschreibung, Dringend, Frist,
Zielbetrag, Fortschritt, Export, öffentliche Seite, Erinnerungen."""

from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.cotisations.aktions_erinnerungen import erinnere_beitragsaktionen, faellige_regeln
from apps.cotisations.models import (
    AktionsErinnerung,
    AktionsRegel,
    ArticleCatalogue,
    StatutCotisation,
    TypeArticle,
)
from apps.cotisations.tests.factories import ArticleCatalogueFactory, CotisationFactory
from apps.membres.models import StatutMembre
from apps.membres.tests.factories import MembreFactory
from apps.notifications.models import Notification, TypeNotification

pytestmark = pytest.mark.django_db

LIST_URL = "cotisations:article-catalogue-list"


def _user(role, email, langue="de"):
    return User.objects.create_user(
        email=email, password="Password123!", role=role, is_active=True, langue_preferee=langue
    )


def _client(user):
    c = APIClient()
    c.force_authenticate(user=user)
    return c


_zaehler = {"n": 0}


def _admin():
    _zaehler["n"] += 1
    return _client(_user(Role.SUPER_ADMIN, f"admin{_zaehler['n']}@example.de"))


def _mitglied(email="m@example.de", langue="de", statut=StatutMembre.ACTIF):
    user = _user(Role.MEMBRE, email, langue)
    return MembreFactory(user=user, statut=statut)


def _aktion(**kw):
    kw.setdefault("libelle", "Busfahrt Auswärtsspiel")
    kw.setdefault("montant", "25.00")
    return ArticleCatalogueFactory(**kw)


def _zahlung(artikel, membre, betrag="25.00", statut=StatutCotisation.PAYEE):
    return CotisationFactory(
        membre=membre,
        type_article=TypeArticle.AUTRE,
        article_catalogue=artikel,
        libelle=artikel.libelle,
        montant=betrag,
        statut=statut,
        annee=None,
    )


# --- Verwaltung : Felder, Bereinigung, Prüfungen ---


def test_admin_legt_aktion_mit_allen_feldern_an():
    client = _admin()
    frist = (timezone.localdate() + timedelta(days=10)).isoformat()
    resp = client.post(
        reverse(LIST_URL),
        {
            "libelle": "Spendenaufruf Fahne",
            "montant": "10.00",
            "beschreibung": "<p>Wir brauchen <strong>dich</strong></p>",
            "dringend": True,
            "frist": frist,
            "zielbetrag": "500.00",
            "oeffentlich": True,
            "erinnerungen": True,
        },
        format="json",
    )
    assert resp.status_code == 201, resp.content
    assert resp.data["dringend"] is True
    assert resp.data["fortschritt"] == 0
    assert resp.data["abgelaufen"] is False
    assert resp.data["gesammelt"] == "0.00"


def test_beschreibung_wird_serverseitig_bereinigt():
    client = _admin()
    resp = client.post(
        reverse(LIST_URL),
        {
            "libelle": "X",
            "montant": "5.00",
            "beschreibung": '<p onclick="x()">Hallo</p><script>alert(1)</script>'
            '<a href="javascript:evil()">Link</a>',
        },
        format="json",
    )
    assert resp.status_code == 201
    html = resp.data["beschreibung"]
    assert "script" not in html and "onclick" not in html and "javascript:" not in html
    assert "Hallo" in html


def test_frist_in_der_vergangenheit_wird_abgelehnt():
    client = _admin()
    gestern = (timezone.localdate() - timedelta(days=1)).isoformat()
    resp = client.post(
        reverse(LIST_URL),
        {"libelle": "Alt", "montant": "5.00", "frist": gestern},
        format="json",
    )
    assert resp.status_code == 400
    assert "frist" in resp.data["details"]


def test_unveraenderte_alte_frist_blockiert_andere_aenderungen_nicht():
    artikel = _aktion(frist=timezone.localdate() - timedelta(days=3))
    client = _admin()
    resp = client.patch(
        reverse("cotisations:article-catalogue-detail", args=[artikel.id]),
        {"dringend": True, "frist": artikel.frist.isoformat()},
        format="json",
    )
    assert resp.status_code == 200


def test_erinnerungen_brauchen_frist_oder_zielbetrag():
    client = _admin()
    resp = client.post(
        reverse(LIST_URL),
        {"libelle": "Ohne", "montant": "5.00", "erinnerungen": True},
        format="json",
    )
    assert resp.status_code == 400
    assert "erinnerungen" in resp.data["details"]


def test_aktionsfelder_der_technischen_zeilen_werden_ignoriert():
    zeile = ArticleCatalogue.objects.get(type_fixe=TypeArticle.COTISATION)
    client = _admin()
    resp = client.patch(
        reverse("cotisations:article-catalogue-detail", args=[zeile.id]),
        {"dringend": True, "zielbetrag": "100.00", "beschreibung": "<p>x</p>"},
        format="json",
    )
    assert resp.status_code == 200
    zeile.refresh_from_db()
    assert zeile.dringend is False and zeile.zielbetrag is None and zeile.beschreibung == ""


def test_normales_mitglied_kann_nicht_schreiben():
    mitglied = _mitglied()
    resp = _client(mitglied.user).post(
        reverse(LIST_URL), {"libelle": "X", "montant": "1.00", "dringend": True}, format="json"
    )
    assert resp.status_code == 403


# --- Fortschritt zählt nur bezahlte Beiträge ---


def test_fortschritt_zaehlt_nur_bezahlte_beitraege():
    artikel = _aktion(zielbetrag=Decimal("100.00"))
    m1, m2, m3 = _mitglied("a@x.de"), _mitglied("b@x.de"), _mitglied("c@x.de")
    _zahlung(artikel, m1, "25.00")
    _zahlung(artikel, m1, "25.00")  # gleiche Person zweimal
    _zahlung(artikel, m2, "10.00")
    _zahlung(artikel, m3, "99.00", statut=StatutCotisation.EN_ATTENTE)  # offen : zählt nicht
    resp = _admin().get(reverse("cotisations:article-catalogue-detail", args=[artikel.id]))
    assert resp.status_code == 200
    assert resp.data["gesammelt"] == "60.00"
    assert resp.data["beitragende"] == 2
    assert resp.data["fortschritt"] == 60


def test_fortschritt_ist_bei_100_gedeckelt():
    artikel = _aktion(zielbetrag=Decimal("20.00"))
    _zahlung(artikel, _mitglied(), "50.00")
    resp = _admin().get(reverse("cotisations:article-catalogue-detail", args=[artikel.id]))
    assert resp.data["fortschritt"] == 100


def test_ohne_zielbetrag_kein_fortschritt():
    artikel = _aktion()
    resp = _admin().get(reverse("cotisations:article-catalogue-detail", args=[artikel.id]))
    assert resp.data["fortschritt"] is None


def test_dringende_aktionen_stehen_in_der_liste_zuerst():
    _aktion(libelle="AAA normal")
    _aktion(libelle="ZZZ dringend", dringend=True)
    mitglied = _mitglied()
    resp = _client(mitglied.user).get(reverse(LIST_URL))
    namen = [e["libelle"] for e in resp.data["results"] if e["type_fixe"] is None]
    assert namen == ["ZZZ dringend", "AAA normal"]


# --- Frist ---


def _bezahlen(client, artikel):
    return client.post(
        reverse("cotisations:cotisation-list"),
        {
            "type_article": TypeArticle.AUTRE,
            "article_catalogue": str(artikel.id),
            "mode_paiement": "virement_sepa",
        },
        format="json",
    )


def test_nach_fristende_keine_neue_zahlung():
    artikel = _aktion(frist=timezone.localdate() - timedelta(days=1))
    mitglied = _mitglied()
    resp = _bezahlen(_client(mitglied.user), artikel)
    assert resp.status_code == 400
    assert "article_catalogue" in resp.data["details"]


def test_am_fristtag_ist_zahlung_noch_moeglich():
    artikel = _aktion(frist=timezone.localdate())
    mitglied = _mitglied()
    resp = _bezahlen(_client(mitglied.user), artikel)
    assert resp.status_code == 201, resp.content


def test_nach_erreichtem_ziel_sind_weitere_beitraege_moeglich():
    artikel = _aktion(zielbetrag=Decimal("10.00"))
    _zahlung(artikel, _mitglied("a@x.de"), "50.00")
    mitglied = _mitglied("b@x.de")
    resp = _bezahlen(_client(mitglied.user), artikel)
    assert resp.status_code == 201


# --- Export ---


def _export_url(artikel):
    return reverse("cotisations:article-catalogue-teilnehmer-export", args=[artikel.id])


def test_export_nur_fuer_app_admin():
    artikel = _aktion()
    mitglied = _mitglied()
    assert _client(mitglied.user).get(_export_url(artikel)).status_code == 403
    assert _admin().get(_export_url(artikel)).status_code == 200


def test_export_enthaelt_nur_bezahlte_standardmaessig_und_die_summe():
    from io import BytesIO

    from openpyxl import load_workbook

    artikel = _aktion()
    bezahlt = _mitglied("p@x.de")
    offen = _mitglied("o@x.de")
    _zahlung(artikel, bezahlt, "25.00")
    _zahlung(artikel, offen, "25.00", statut=StatutCotisation.EN_ATTENTE)

    resp = _admin().get(_export_url(artikel))
    assert resp["Content-Type"].startswith("application/vnd.openxmlformats")
    blatt = load_workbook(BytesIO(resp.content)).active
    zeilen = [[z.value for z in r] for r in blatt.iter_rows(min_row=4)]
    namen = [z[0] for z in zeilen if z[0]]
    assert len(namen) == 1
    assert any(z[2] == "Summe bezahlt" and z[3] == 25.0 for z in zeilen)

    alle = _admin().get(_export_url(artikel), {"status": "alle"})
    blatt = load_workbook(BytesIO(alle.content)).active
    assert len([r for r in blatt.iter_rows(min_row=4) if r[0].value]) == 2


def test_export_ungueltiger_status_wird_abgelehnt():
    artikel = _aktion()
    assert _admin().get(_export_url(artikel), {"status": "x"}).status_code == 400


# --- Öffentliche Seite und Vorschau ---


def test_oeffentliche_seite_nur_wenn_freigegeben():
    versteckt = _aktion()
    frei = _aktion(libelle="Frei", oeffentlich=True, beschreibung="<p>Hilf mit</p>")
    anon = APIClient()
    assert anon.get(reverse("aktionen:oeffentlich", args=[versteckt.id])).status_code == 404
    assert anon.get(reverse("aktionen:teilen", args=[versteckt.id])).status_code == 404
    resp = anon.get(reverse("aktionen:oeffentlich", args=[frei.id]))
    assert resp.status_code == 200
    assert resp.data["titel"] == "Frei"


def test_inaktive_aktion_ist_nicht_oeffentlich():
    artikel = _aktion(oeffentlich=True, actif=False)
    assert APIClient().get(reverse("aktionen:oeffentlich", args=[artikel.id])).status_code == 404


def test_technische_zeile_ist_nie_oeffentlich():
    zeile = ArticleCatalogue.objects.get(type_fixe=TypeArticle.COTISATION)
    ArticleCatalogue.objects.filter(pk=zeile.pk).update(oeffentlich=True)
    assert APIClient().get(reverse("aktionen:oeffentlich", args=[zeile.id])).status_code == 404


def test_oeffentliche_daten_enthalten_keine_namen():
    artikel = _aktion(oeffentlich=True, zielbetrag=Decimal("100.00"))
    mitglied = _mitglied("geheim@x.de")
    _zahlung(artikel, mitglied, "25.00")
    resp = APIClient().get(reverse("aktionen:oeffentlich", args=[artikel.id]))
    text = str(resp.data)
    assert mitglied.nom not in text and mitglied.prenom not in text and "geheim@x.de" not in text
    assert resp.data["gesammelt"] == "25.00"
    assert resp.data["beitragende"] == 1
    assert resp.data["fortschritt"] == 25


def test_teilen_seite_hat_open_graph_und_maskiert_html():
    artikel = _aktion(
        libelle='Fahne <b>"neu"</b>', oeffentlich=True, beschreibung="<p>Für <em>alle</em></p>"
    )
    resp = APIClient().get(reverse("aktionen:teilen", args=[artikel.id]))
    assert resp.status_code == 200
    html = resp.content.decode()
    assert 'property="og:title"' in html and 'property="og:image"' in html
    assert "<b>" not in html
    assert f"/aktion/{artikel.id}" in html
    assert "Für alle" in html


# --- Erinnerungen ---


HEUTE = date(2026, 11, 10)


def test_regeln_frist_und_ziel():
    artikel = _aktion(frist=HEUTE + timedelta(days=3), zielbetrag=Decimal("100.00"))
    assert faellige_regeln(artikel, HEUTE) == [AktionsRegel.FRIST_BALD]
    _zahlung(artikel, _mitglied("z@x.de"), "85.00")
    assert faellige_regeln(artikel, HEUTE) == [AktionsRegel.FRIST_BALD, AktionsRegel.ZIEL_FAST]


def test_keine_regel_ausserhalb_des_fensters_oder_nach_ziel_oder_ablauf():
    weit = _aktion(frist=HEUTE + timedelta(days=4))
    assert faellige_regeln(weit, HEUTE) == []
    erreicht = _aktion(zielbetrag=Decimal("10.00"))
    _zahlung(erreicht, _mitglied("e@x.de"), "10.00")
    assert faellige_regeln(erreicht, HEUTE) == []
    vorbei = _aktion(frist=HEUTE - timedelta(days=1), zielbetrag=Decimal("100.00"))
    _zahlung(vorbei, _mitglied("v@x.de"), "90.00")
    assert faellige_regeln(vorbei, HEUTE) == []


def test_erinnerung_geht_nur_an_aktive_ohne_beitrag_und_nur_einmal(mailoutbox):
    artikel = _aktion(erinnerungen=True, frist=HEUTE + timedelta(days=2))
    ohne = _mitglied("ohne@x.de")
    bezahlt = _mitglied("bezahlt@x.de")
    offen = _mitglied("offen@x.de")
    inaktiv = _mitglied("inaktiv@x.de", statut=StatutMembre.INACTIF)
    _zahlung(artikel, bezahlt)
    _zahlung(artikel, offen, statut=StatutCotisation.EN_ATTENTE)

    ergebnis = erinnere_beitragsaktionen(HEUTE)

    assert ergebnis["gesendet"] == 1
    assert [m.to for m in mailoutbox] == [[ohne.user.email]]
    assert AktionsErinnerung.objects.filter(membre=ohne, regel=AktionsRegel.FRIST_BALD).exists()
    assert not AktionsErinnerung.objects.filter(membre=inaktiv).exists()
    notiz = Notification.objects.get(destinataire=ohne.user)
    assert notiz.type_notification == TypeNotification.AKTION_FRIST
    assert notiz.lien == f"/cotisation?aktion={artikel.id}"

    assert erinnere_beitragsaktionen(HEUTE)["gesendet"] == 0  # kein Doppelversand
    assert len(mailoutbox) == 1


def test_ohne_eingeschaltete_erinnerungen_wird_nichts_gesendet(mailoutbox):
    _aktion(erinnerungen=False, frist=HEUTE + timedelta(days=1))
    _mitglied()
    assert erinnere_beitragsaktionen(HEUTE)["gesendet"] == 0
    assert len(mailoutbox) == 0


def test_erinnerung_in_der_sprache_des_mitglieds(mailoutbox):
    _aktion(erinnerungen=True, frist=HEUTE + timedelta(days=1))
    _mitglied("fr@x.de", langue="fr")
    erinnere_beitragsaktionen(HEUTE)
    assert "contribution" in mailoutbox[0].subject.lower()


def test_ziel_fast_erreicht_nennt_den_fehlbetrag(mailoutbox):
    artikel = _aktion(erinnerungen=True, zielbetrag=Decimal("100.00"))
    _zahlung(artikel, _mitglied("spender@x.de"), "90.00")
    _mitglied("neu@x.de")
    erinnere_beitragsaktionen(HEUTE)
    mail = next(m for m in mailoutbox if m.to == ["neu@x.de"])
    assert "10,00" in mail.body


def test_aktion_mit_abgeschalteten_mails_erzeugt_trotzdem_hinweis(mailoutbox):
    from apps.notifications.models import ParametresNotification

    ParametresNotification.objects.update_or_create(pk=1, defaults={"email_cotisations": False})
    _aktion(erinnerungen=True, frist=HEUTE + timedelta(days=1))
    m = _mitglied()
    erinnere_beitragsaktionen(HEUTE)
    assert len(mailoutbox) == 0
    assert Notification.objects.filter(destinataire=m.user).count() == 1


# --- Mindestbetrag (Betrag darf vom Mitglied erhöht werden) ---


def _bezahlen_mit_betrag(client, artikel, betrag):
    return client.post(
        reverse("cotisations:cotisation-list"),
        {
            "type_article": TypeArticle.AUTRE,
            "article_catalogue": str(artikel.id),
            "mode_paiement": "virement_sepa",
            "montant": betrag,
        },
        format="json",
    )


def test_mindestbetrag_darf_erhoeht_werden():
    artikel = _aktion(betrag_erhoehbar=True, montant="15.00")
    resp = _bezahlen_mit_betrag(_client(_mitglied().user), artikel, "40.50")
    assert resp.status_code == 201, resp.content
    assert resp.data["montant"] == "40.50"


def test_mindestbetrag_genau_erlaubt_und_ohne_angabe_gilt_der_mindestbetrag():
    artikel = _aktion(betrag_erhoehbar=True, montant="15.00")
    assert (
        _bezahlen_mit_betrag(_client(_mitglied("a@x.de").user), artikel, "15.00").status_code == 201
    )
    resp = _bezahlen(_client(_mitglied("b@x.de").user), artikel)
    assert resp.status_code == 201
    assert resp.data["montant"] == "15.00"


def test_unter_dem_mindestbetrag_wird_abgelehnt():
    artikel = _aktion(betrag_erhoehbar=True, montant="15.00")
    resp = _bezahlen_mit_betrag(_client(_mitglied().user), artikel, "14.99")
    assert resp.status_code == 400
    assert "montant" in resp.data["details"]


def test_fester_preis_ignoriert_einen_gesendeten_betrag():
    artikel = _aktion(montant="25.00")
    resp = _bezahlen_mit_betrag(_client(_mitglied().user), artikel, "99.00")
    assert resp.status_code == 201
    assert resp.data["montant"] == "25.00"


def test_betrag_erhoehbar_wird_verwaltet_und_oeffentlich_angezeigt():
    client = _admin()
    resp = client.post(
        reverse(LIST_URL),
        {
            "libelle": "Spendenaktion",
            "montant": "10.00",
            "betrag_erhoehbar": True,
            "oeffentlich": True,
        },
        format="json",
    )
    assert resp.status_code == 201, resp.content
    assert resp.data["betrag_erhoehbar"] is True
    seite = APIClient().get(f"/api/v1/aktionen/{resp.data['id']}/")
    assert seite.status_code == 200
    assert seite.data["betrag_erhoehbar"] is True


def test_erhoehter_betrag_zaehlt_im_fortschritt():
    artikel = _aktion(betrag_erhoehbar=True, montant="15.00", zielbetrag=Decimal("100.00"))
    _zahlung(artikel, _mitglied(), "60.00")
    gesammelt, anzahl = ArticleCatalogue.objects.get(pk=artikel.pk).fortschritt()
    assert gesammelt == Decimal("60.00")
    assert anzahl == 1
