"""
Génère un classeur .xlsx vide avec les en-têtes attendues par l'import
membres (RICEFW C-001/W-008, voir apps.membres.imports.REQUIRED_COLUMNS /
OPTIONAL_COLUMNS) — à fournir à l'association pour saisir les ~300 lignes
historiques avant de les uploader via POST /api/v1/membres/import/.

Usage : python manage.py generer_template_import_membres [--out chemin.xlsx]
"""

import openpyxl
from django.core.management.base import BaseCommand
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from apps.membres.imports import OPTIONAL_COLUMNS, REQUIRED_COLUMNS
from apps.membres.models import Bundesland

# En-tête "canonique" à écrire (le 1er alias déclaré dans imports.py) +
# exemple de valeur pour guider la saisie (ligne 2 du template).
_EXEMPLE = {
    "prenom": "Riadh",
    "nom": "Bchini",
    "date_naissance": "15/03/1985",
    "sexe": "Homme",
    "email": "riadh.bchini@example.de",
    "telephone": "+49 170 1234567",
    "cin": "12345678",
    "passeport": "",
    "adresse_de": "Musterstr. 1",
    "code_postal_de": "10115",
    "ville_de": "Berlin",
    "land_de": "Berlin",
    "ville_origine_tn": "Tunis",
    "gouvernorat_tn": "Tunis",
    "statut": "actif",
    "date_adhesion": "01/09/2019",
}


class Command(BaseCommand):
    help = "Génère le template Excel vide pour l'import des membres historiques (RICEFW C-001)."

    def add_arguments(self, parser):
        parser.add_argument("--out", default="template_import_membres.xlsx")

    def handle(self, *args, **options):
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.title = "Membres"

        champs = list(REQUIRED_COLUMNS) + list(OPTIONAL_COLUMNS)
        en_tete_style = Font(bold=True, color="FFFFFF")
        remplissage = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")

        for col_idx, champ in enumerate(champs, start=1):
            cellule = feuille.cell(row=1, column=col_idx, value=champ)
            cellule.font = en_tete_style
            cellule.fill = remplissage
            if champ not in REQUIRED_COLUMNS:
                cellule.comment = Comment("Optionnel", "CID")
            feuille.cell(row=2, column=col_idx, value=_EXEMPLE.get(champ, ""))
            feuille.column_dimensions[cellule.column_letter].width = max(len(champ) + 2, 14)

        notes = classeur.create_sheet("Notes")
        notes["A1"] = "Colonnes obligatoires : " + ", ".join(REQUIRED_COLUMNS)
        notes["A2"] = "Colonnes optionnelles : " + ", ".join(OPTIONAL_COLUMNS)
        notes["A3"] = "sexe : Homme / Femme (laisser vide si non renseigné)"
        notes["A4"] = "statut : actif / en_attente / inactif (défaut si vide : actif)"
        notes["A5"] = (
            "land_de : code à 2 lettres (BE, BY, ...) ou nom complet — Länder valides : "
            + ", ".join(f"{code} ({label})" for code, label in Bundesland.choices)
        )
        notes["A6"] = "Dates au format JJ/MM/AAAA."
        notes.column_dimensions["A"].width = 100

        classeur.save(options["out"])
        self.stdout.write(self.style.SUCCESS(f"Template généré : {options['out']}"))
