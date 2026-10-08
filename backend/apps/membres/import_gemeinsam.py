"""
Gemeinsame Bausteine der beiden zweistufigen Excel-Importe (Mitglieder-Stammdaten in
`apps.membres.imports`, Statushistorie in `apps.membres.imports_historique`).

Ablauf (Entscheidung des Nutzers, 2026-10-08) :
  1. PRÜFEN   — Datei wird komplett ausgewertet, aber NICHTS geschrieben. Ergebnis : eine
                detaillierte Zeilenliste (neu / Dublette / unverändert / Fehler).
  2. BESTÄTIGEN — dieselbe Datei wird erneut hochgeladen (zustandslos, kein Server-Cache), noch
                einmal ausgewertet (die Datenbank kann sich inzwischen geändert haben) und erst
                jetzt geschrieben. Dubletten werden nur für die ausdrücklich gewählten
                Zeilennummern überschrieben ; neue gültige Zeilen werden immer importiert.
  3. BERICHT  — die Originaldatei wird um die Spalten „Status" und „Grund" ergänzt und als
                Excel-Bericht zurückgegeben (Sprache : Deutsch, Entscheidung des Nutzers).
"""

import base64
import io
import os
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Font, PatternFill

from .utils_http import safe_cell

# --- Status einer Zeile in der Prüfphase -----------------------------------------------------
STATUS_NEU = "neu"
STATUS_DUBLETTE = "dublette"
STATUS_UNVERAENDERT = "unveraendert"
STATUS_FEHLER = "fehler"

# --- Ergebnis einer Zeile nach der Bestätigung (Schlüssel -> Text im Bericht) ----------------
ERG_IMPORTIERT = "importiert"
ERG_UEBERSCHRIEBEN = "ueberschrieben"
ERG_TEILWEISE = "teilweise"
ERG_UEBERSPRUNGEN = "uebersprungen"
ERG_UNVERAENDERT = "unveraendert"
ERG_FEHLER = "fehler"

ERGEBNIS_TEXT = {
    ERG_IMPORTIERT: "Importiert",
    ERG_UEBERSCHRIEBEN: "Überschrieben",
    ERG_TEILWEISE: "Teilweise importiert",
    ERG_UEBERSPRUNGEN: "Übersprungen",
    ERG_UNVERAENDERT: "Unverändert",
    ERG_FEHLER: "Fehler",
}

# Hintergrundfarbe der Status-Zelle im Bericht (hell, damit der Text lesbar bleibt).
_ERGEBNIS_FARBE = {
    ERG_IMPORTIERT: "C6EFCE",
    ERG_UEBERSCHRIEBEN: "BDD7EE",
    ERG_TEILWEISE: "FFEB9C",
    ERG_UEBERSPRUNGEN: "FFEB9C",
    ERG_UNVERAENDERT: "EDEDED",
    ERG_FEHLER: "FFC7CE",
}

SPALTE_STATUS = "Status"
SPALTE_GRUND = "Grund"


class ImportDateiFehler(Exception):
    """Datei unlesbar, leer oder ohne die nötigen Spalten — vor jeder Zeilenauswertung."""


@dataclass
class ZeileAnalyse:
    """Prüfergebnis einer Datenzeile. `ligne` = Excel-Zeilennummer (Kopfzeile = 1)."""

    ligne: int
    status: str
    grund: str = ""
    anzeige: dict = field(default_factory=dict)  # prenom / nom / email für die Prüfliste
    existant: object = None  # bestehende Membre (nur bei Dublette/unverändert)
    aenderungen: list = field(default_factory=list)  # [{champ,label,alt,neu,art}]
    daten: object = None  # interne Nutzlast der Ausführungsphase (nie serialisiert)

    @property
    def ueberschreibbar(self) -> bool:
        return self.status == STATUS_DUBLETTE and bool(self.aenderungen)

    def as_dict(self) -> dict:
        existant = None
        if self.existant is not None:
            existant = {
                "id": str(self.existant.id),
                "numero_membre": self.existant.numero_membre,
                "prenom": self.existant.prenom,
                "nom": self.existant.nom,
                "email": self.existant.email,
                "a_un_compte": self.existant.user_id is not None,
            }
        return {
            "ligne": self.ligne,
            "status": self.status,
            "grund": self.grund,
            "ueberschreibbar": self.ueberschreibbar,
            "anzeige": self.anzeige,
            "existant": existant,
            "aenderungen": self.aenderungen,
        }


@dataclass
class Analyse:
    zeilen: list
    classeur: object  # openpyxl.Workbook — für den Bericht wiederverwendet
    dateiname: str = ""

    def zaehler(self) -> dict:
        zaehler = {
            STATUS_NEU: 0,
            STATUS_DUBLETTE: 0,
            STATUS_UNVERAENDERT: 0,
            STATUS_FEHLER: 0,
        }
        for zeile in self.zeilen:
            zaehler[zeile.status] += 1
        return zaehler

    def as_dict(self) -> dict:
        return {
            "total": len(self.zeilen),
            "zaehler": self.zaehler(),
            "zeilen": [z.as_dict() for z in self.zeilen],
        }


@dataclass
class ZeilenErgebnis:
    ligne: int
    ergebnis: str  # ERG_*
    grund: str = ""

    def as_dict(self) -> dict:
        return {
            "ligne": self.ligne,
            "ergebnis": self.ergebnis,
            "status": ERGEBNIS_TEXT[self.ergebnis],
            "grund": self.grund,
        }


@dataclass
class Ergebnis:
    zeilen: list  # list[ZeilenErgebnis]
    bericht_name: str
    bericht_bytes: bytes

    def zaehler(self) -> dict:
        zaehler = {schluessel: 0 for schluessel in ERGEBNIS_TEXT}
        for zeile in self.zeilen:
            zaehler[zeile.ergebnis] += 1
        return zaehler

    def as_dict(self) -> dict:
        return {
            "total": len(self.zeilen),
            "zaehler": self.zaehler(),
            "zeilen": [z.as_dict() for z in self.zeilen],
            "bericht": {
                "dateiname": self.bericht_name,
                "inhalt_base64": base64.b64encode(self.bericht_bytes).decode("ascii"),
            },
        }


def oeffne_arbeitsmappe(fichier):
    """Lädt die erste Tabelle des .xlsx (Werte statt Formeln). Wirft ImportDateiFehler."""
    try:
        return openpyxl.load_workbook(fichier, data_only=True)
    except Exception as exc:  # openpyxl wirft je nach Problem unterschiedliche Typen
        raise ImportDateiFehler(f"Excel-Datei nicht lesbar: {exc}") from exc


def lies_kopfzeile_und_zeilen(classeur):
    """(Kopfzeile, [(Excel-Zeilennummer, Zellwerte)...]) der ersten Tabelle ; leere Zeilen
    werden übersprungen (kein Fehler)."""
    feuille = classeur.worksheets[0]
    zeilen = feuille.iter_rows(values_only=True)
    try:
        kopf = next(zeilen)
    except StopIteration:
        raise ImportDateiFehler("Die Datei ist leer.") from None
    daten = []
    for nummer, zeile in enumerate(zeilen, start=2):
        if zeile is None or all(c is None or c == "" for c in zeile):
            continue
        daten.append((nummer, zeile))
    return kopf, daten


def bericht_dateiname(originalname: str) -> str:
    """`liste.xlsx` -> `liste_bericht.xlsx` (nur sichere Zeichen, nie ein Pfad)."""
    basis = os.path.splitext(os.path.basename(originalname or ""))[0]
    basis = re.sub(r"[^A-Za-z0-9._-]+", "_", basis).strip("._") or "import"
    return f"{basis[:80]}_bericht.xlsx"


def erzeuge_bericht(classeur, ergebnisse: list, originalname: str):
    """
    Ergänzt die ERSTE Tabelle der Originaldatei um die Spalten „Status" und „Grund" (ganz
    rechts) und gibt (Dateiname, Bytes) zurück. Zeilen ohne Ergebnis (leere Zeilen) bleiben
    leer. Nutzerdaten gehen durch `safe_cell` (Formel-Injektion).
    """
    feuille = classeur.worksheets[0]
    spalte_status = feuille.max_column + 1
    spalte_grund = spalte_status + 1

    kopf_stil = Font(bold=True, color="FFFFFF")
    kopf_fuellung = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")
    for spalte, titel in ((spalte_status, SPALTE_STATUS), (spalte_grund, SPALTE_GRUND)):
        zelle = feuille.cell(row=1, column=spalte, value=titel)
        zelle.font = kopf_stil
        zelle.fill = kopf_fuellung

    for ergebnis in ergebnisse:
        zelle_status = feuille.cell(
            row=ergebnis.ligne, column=spalte_status, value=ERGEBNIS_TEXT[ergebnis.ergebnis]
        )
        zelle_status.fill = PatternFill(
            start_color=_ERGEBNIS_FARBE[ergebnis.ergebnis],
            end_color=_ERGEBNIS_FARBE[ergebnis.ergebnis],
            fill_type="solid",
        )
        feuille.cell(row=ergebnis.ligne, column=spalte_grund, value=safe_cell(ergebnis.grund))

    feuille.column_dimensions[feuille.cell(row=1, column=spalte_status).column_letter].width = 22
    feuille.column_dimensions[feuille.cell(row=1, column=spalte_grund).column_letter].width = 80

    puffer = io.BytesIO()
    classeur.save(puffer)
    return bericht_dateiname(originalname), puffer.getvalue()
