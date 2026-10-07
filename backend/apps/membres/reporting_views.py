"""API des Mitglieder-Reportings (siehe reporting.py) — nur lesend, ab Rolle RH (wie Export/
Import der Mitglieder)."""

from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.permissions import IsRHOrAbove

from .reporting import AKTIVITAETS_TYPEN, TYPEN, aktivitaeten, mitglieder_liste
from .utils_http import xlsx_response

MAX_SEITENGROESSE = 200


def _str(wert):
    return f"{wert:.2f}" if isinstance(wert, Decimal) else wert


def _seite(params, standard):
    try:
        seite = max(int(params.get("page", 1)), 1)
        groesse = min(max(int(params.get("page_size", standard)), 1), MAX_SEITENGROESSE)
    except ValueError:
        seite, groesse = 1, standard
    return seite, groesse


class MitgliederReportingView(APIView):
    """GET /membres/reporting/mitglieder/ — Mitgliederliste mit Statushistorie und Aktivitäts-
    zahlen. Filter: die der Mitgliederliste (statut, ville, land, pays, nom, q, date_adhesion_
    apres/avant), `historie_statut`, `historie_jahr`, `typ` (Aktivitätsarten, kommagetrennt),
    `von`/`bis`, `aktivitaet_status`, `min_betrag`/`max_betrag`, `min_aktivitaeten`,
    `ohne_aktivitaet=1`, `sortierung`, `page`, `page_size`."""

    permission_classes = [IsRHOrAbove]

    def get(self, request):
        zeilen = mitglieder_liste(request.query_params)
        seite, groesse = _seite(request.query_params, 25)
        start = (seite - 1) * groesse
        return Response(
            {
                "count": len(zeilen),
                "page": seite,
                "page_size": groesse,
                "summen": {
                    "mitglieder": len(zeilen),
                    "aktivitaeten": sum(z["aktivitaeten_gesamt"] for z in zeilen),
                    "betrag": _str(sum((z["betrag_gesamt"] for z in zeilen), Decimal("0.00"))),
                },
                "ergebnisse": [
                    {**z, "betrag_gesamt": _str(z["betrag_gesamt"])}
                    for z in zeilen[start : start + groesse]
                ],
            }
        )


class AktivitaetenReportingView(APIView):
    """GET /membres/reporting/aktivitaeten/ — alle Aktivitäten der gefilterten Mitglieder,
    neueste zuerst. Zusätzlich zu den Filtern oben: `membre` (einzelnes Mitglied)."""

    permission_classes = [IsRHOrAbove]

    def get(self, request):
        seite, groesse = _seite(request.query_params, 50)
        zeilen, anzahl, summen = aktivitaeten(request.query_params, seite, groesse)
        return Response(
            {
                "count": anzahl,
                "page": seite,
                "page_size": groesse,
                "typen": TYPEN,
                "summen": {
                    typ: {"anzahl": s["anzahl"], "betrag": _str(s["betrag"])}
                    for typ, s in summen.items()
                },
                "ergebnisse": [{**z, "betrag": _str(z["betrag"])} for z in zeilen],
            }
        )


def _kopf(blatt, spalten):
    blatt.append(spalten)
    for zelle in blatt[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill("solid", fgColor="C8102E")
    blatt.freeze_panes = "A2"


def _breiten(blatt):
    for spalte in blatt.columns:
        blatt.column_dimensions[spalte[0].column_letter].width = 22


class MitgliederReportingExportView(APIView):
    """GET /membres/reporting/export/?ansicht=mitglieder|aktivitaeten — gleiche Filter wie die
    Ansichten, alle Zeilen (ungeseitet) als Excel. Kennungen (Mitgliedsnummer) als Text."""

    permission_classes = [IsRHOrAbove]

    def get(self, request):
        arbeitsmappe = Workbook()
        blatt = arbeitsmappe.active
        if request.query_params.get("ansicht") == "aktivitaeten":
            blatt.title = "Aktivitäten"
            _kopf(
                blatt,
                ["Datum", "Art", "Mitglied", "Mitgliedsnr.", "Beschreibung", "Betrag", "Status"],
            )
            zeilen, _anzahl, _summen = aktivitaeten(request.query_params)
            for z in zeilen:
                blatt.append(
                    [
                        z["datum"][:16].replace("T", " "),
                        z["typ"],
                        z["membre_name"],
                        z["numero_membre"],
                        z["titel"],
                        float(z["betrag"]) if z["betrag"] is not None else None,
                        z["statut"],
                    ]
                )
            name = "aktivitaeten"
        else:
            blatt.title = "Mitglieder"
            _kopf(
                blatt,
                [
                    "Mitgliedsnr.",
                    "Nachname",
                    "Vorname",
                    "E-Mail",
                    "Land",
                    "Ort",
                    "Status",
                    "Mitglied seit",
                ]
                + AKTIVITAETS_TYPEN
                + ["Aktivitäten gesamt", "Betrag gesamt"],
            )
            historie = arbeitsmappe.create_sheet("Historie")
            _kopf(
                historie,
                ["Mitgliedsnr.", "Nachname", "Vorname", "Jahr", "Status", "Grund", "Wirksam ab"],
            )
            for z in mitglieder_liste(request.query_params):
                blatt.append(
                    [
                        z["numero_membre"],
                        z["nom"],
                        z["prenom"],
                        z["email"],
                        z["pays"],
                        z["ville"],
                        z["statut"],
                        z["date_adhesion"],
                    ]
                    + [z["aktivitaeten"][t] for t in AKTIVITAETS_TYPEN]
                    + [z["aktivitaeten_gesamt"], float(z["betrag_gesamt"])]
                )
                blatt.cell(row=blatt.max_row, column=1).number_format = "@"
                for h in z["historie"]:
                    historie.append(
                        [
                            z["numero_membre"],
                            z["nom"],
                            z["prenom"],
                            h["annee"],
                            h["statut"],
                            h["raison"],
                            h["date_effet"],
                        ]
                    )
            _breiten(historie)
            name = "mitglieder"
        _breiten(blatt)
        return xlsx_response(
            arbeitsmappe, f"reporting_{name}_{timezone.localtime():%Y%m%d_%H%M}.xlsx"
        )
