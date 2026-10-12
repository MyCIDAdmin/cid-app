"""
Import Excel « Historie MIT Angebot und Beitrag » (Roadmap Release 2, 2026-10-12) — Erweiterung
des bestehenden Historie-Imports (apps.membres.imports_historique), KEIN eigener Import : dieselben
Endpunkte (pruefen/ + bestaetigen/), dieselbe Prüfliste, derselbe Bericht, dieselbe Dublettenlogik.

Erkennung des Formats am Kopf der Datei :
  - Spalte `jahr` (auch annee/year) vorhanden -> LANGES Format (dieses Modul) : eine Zeile =
    Mitglied + Jahr + Angebot + Beitrag.
  - sonst -> BREITES Format (imports_historique) : eine Spalte pro Jahr mit actif/inactif.

Spalten (Kopfzeile, Groß-/Kleinschreibung egal) :
  email (Pflicht) · cin (optional) · jahr (Pflicht) · angebot (Pflicht, z. B. „Gold") ·
  kampagne (optional) · betrag (optional ; Pflicht nur, wenn das Angebot neu angelegt werden
  muss) · zahlungsdatum (optional) · zahlungsart (optional : karte / überweisung / paypal / bar).

Pro Zeile wird (Entscheidung des Nutzers, AskUserQuestion 2026-10-12) :
  1. die Kampagne des Jahres gesucht (bei Mehrdeutigkeit hilft die Spalte `kampagne`) — fehlt
     sie, wird sie als ARCHIVIERTE (clôturée) Kampagne angelegt ;
  2. das Angebot in dieser Kampagne gesucht — fehlt es, wird es unsichtbar angelegt ;
  3. eine bezahlte Souscription + eine bezahlte Cotisation (Zahlungsregister) angelegt/verknüpft ;
  4. der Jahresstatus des Mitglieds auf „aktiv" gesetzt (HistoriqueStatutMembre).
Nie eine Benachrichtigung/E-Mail an das Mitglied (Massenimport historischer Daten).

Prüfphase schreibt NICHTS (auch keine Kampagnen/Angebote) ; die Ausführung ist je Zeile atomar.
"""

import datetime
import logging
import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.adhesions.models import (
    CampagneAdhesion,
    OffreAdhesion,
    Souscription,
    StatutCampagne,
    StatutSouscription,
)
from apps.cotisations.models import Cotisation, ModePaiement, StatutCotisation, TypeArticle

from .import_gemeinsam import (
    ERG_FEHLER,
    ERG_IMPORTIERT,
    ERG_TEILWEISE,
    ERG_UEBERSCHRIEBEN,
    ERG_UEBERSPRUNGEN,
    STATUS_DUBLETTE,
    STATUS_FEHLER,
    STATUS_NEU,
    STATUS_UNVERAENDERT,
    Analyse,
    ImportDateiFehler,
    ZeileAnalyse,
    ZeilenErgebnis,
)
from .imports_historique import (
    ANNEE_MAX,
    ANNEE_MIN,
    ART_ABWEICHEND,
    ART_GLEICH,
    ART_NEU,
    _lade_caches,
    _map_identite,
    _normalize,
    _resoudre_membre,
)
from .models import HistoriqueStatutMembre, Membre, RaisonChangementStatut, StatutMembre
from .services import enregistrer_statut_annuel

logger = logging.getLogger(__name__)

_SPALTEN_ALIASE = {
    "jahr": ("jahr", "annee", "année", "year"),
    "kampagne": ("kampagne", "campagne", "campaign"),
    "angebot": ("angebot", "offre", "offer", "tarif"),
    "betrag": ("betrag", "beitrag", "montant", "amount"),
    "datum": ("zahlungsdatum", "datum", "date_paiement", "payment_date", "date"),
    "modus": ("zahlungsart", "zahlungsmodus", "mode_paiement", "mode", "payment_mode"),
}

_MODUS_ALIASE = {
    ModePaiement.CARTE: ("carte", "karte", "kartenzahlung", "card", "cb"),
    ModePaiement.VIREMENT_SEPA: (
        "virement",
        "virement_sepa",
        "überweisung",
        "ueberweisung",
        "sepa",
    ),
    ModePaiement.PAYPAL: ("paypal",),
    ModePaiement.ESPECES: ("especes", "espèces", "bar", "barzahlung", "bargeld", "cash"),
}
_MODUS_LOOKUP = {alias: modus for modus, aliase in _MODUS_ALIASE.items() for alias in aliase}
_MODUS_TEXT = {
    ModePaiement.CARTE: "Karte",
    ModePaiement.VIREMENT_SEPA: "Überweisung",
    ModePaiement.PAYPAL: "PayPal",
    ModePaiement.ESPECES: "Bar",
}

_PRIORITAET = {StatutCampagne.PUBLIEE: 0, StatutCampagne.CLOTUREE: 1, StatutCampagne.BROUILLON: 2}


@dataclass
class BeitragsDaten:
    """Nutzlast einer Zeile des langen Formats (nie serialisiert)."""

    annee: int
    kampagne_name: str
    angebot: str
    betrag: Decimal
    datum: datetime.date
    modus: str
    sous_art: str
    hist_art: str


def ist_langes_format(kopf) -> bool:
    return _spalten(kopf).get("jahr") is not None


def _spalten(kopf) -> dict:
    normalisiert = [_normalize(c) for c in kopf]
    mapping = {}
    for champ, aliase in _SPALTEN_ALIASE.items():
        for idx, header in enumerate(normalisiert):
            if header in aliase:
                mapping[champ] = idx
                break
    return mapping


def _zelle(row, mapping, champ):
    idx = mapping.get(champ)
    return row[idx] if idx is not None and idx < len(row) else None


def _parse_jahr(roh) -> int:
    text = _normalize(roh)
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit() or not ANNEE_MIN <= int(text) <= ANNEE_MAX:
        raise ValueError(f"Ungültiges Jahr {roh!r}")
    return int(text)


def _parse_betrag(roh):
    if roh is None or str(roh).strip() == "":
        return None
    if isinstance(roh, (int, float, Decimal)):
        wert = Decimal(str(roh))
    else:
        text = re.sub(r"[€\s]|eur", "", str(roh).strip().lower())
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
        try:
            wert = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"Ungültiger Betrag {roh!r}") from None
    if wert < 0 or wert >= Decimal("1000000"):
        raise ValueError(f"Ungültiger Betrag {roh!r}")
    return wert.quantize(Decimal("0.01"))


def _parse_datum(roh):
    if roh is None or str(roh).strip() == "":
        return None
    if isinstance(roh, datetime.datetime):
        return roh.date()
    if isinstance(roh, datetime.date):
        return roh
    text = str(roh).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Ungültiges Zahlungsdatum {roh!r}")


def _parse_modus(roh) -> str:
    text = _normalize(roh)
    if not text:
        return ""
    modus = _MODUS_LOOKUP.get(text)
    if modus is None:
        raise ValueError(f"Unbekannte Zahlungsart {roh!r}")
    return modus.value


def _finde_kampagne(annee: int, name: str):
    """Kampagne des Jahres oder None. Ohne Namen : bevorzugt veröffentlicht, dann archiviert,
    dann Entwurf ; mehrere gleichrangige -> ValueError (nie raten)."""
    kandidaten = CampagneAdhesion.objects.filter(annee=annee)
    if name:
        kandidaten = kandidaten.filter(nom__iexact=name)
    kandidaten = list(kandidaten)
    if not kandidaten:
        return None
    beste = min(_PRIORITAET[k.statut] for k in kandidaten)
    erste_gruppe = [k for k in kandidaten if _PRIORITAET[k.statut] == beste]
    if len(erste_gruppe) > 1:
        raise ValueError(f"Mehrere Kampagnen im Jahr {annee} — bitte Spalte „kampagne“ ausfüllen")
    return erste_gruppe[0]


def _finde_angebot(kampagne, name: str):
    if kampagne is None:
        return None
    return kampagne.offres.filter(nom__iexact=name).first()


def _fmt_datum(datum: datetime.date) -> str:
    return datum.strftime("%d.%m.%Y")


def _fmt_betrag(betrag) -> str:
    return f"{Decimal(betrag):.2f} €"


def construire_classeur_template_beitraege():
    """Vorlage des langen Formats (eine Zeile = Mitglied + Jahr + Angebot + Beitrag)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Historie mit Beitrag"
    kopf = ["email", "cin", "jahr", "kampagne", "angebot", "betrag", "zahlungsdatum", "zahlungsart"]
    for col, champ in enumerate(kopf, start=1):
        zelle = feuille.cell(row=1, column=col, value=champ)
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")
        feuille.column_dimensions[zelle.column_letter].width = max(len(champ) + 2, 14)
    beispiele = [
        ["ahmed@example.de", "12345678", 2023, "", "Gold", 120, "15.03.2023", "Überweisung"],
        ["ahmed@example.de", "12345678", 2024, "", "Gold", 120, "20.02.2024", "PayPal"],
    ]
    for zeile, werte in enumerate(beispiele, start=2):
        for col, wert in enumerate(werte, start=1):
            feuille.cell(row=zeile, column=col, value=wert)

    notes = classeur.create_sheet("Notes")
    hinweise = [
        "Eine Zeile = ein Mitglied + ein Jahr + ein Angebot + ein Beitrag. "
        "Mehrere Jahre = mehrere Zeilen.",
        "Pflicht : email, jahr, angebot. cin optional (hilft, wenn sich die E-Mail geändert hat).",
        "Das Mitglied muss bereits existieren. Fehlende Kampagnen und Angebote werden automatisch "
        "angelegt (Kampagne archiviert, Angebot unsichtbar).",
        "betrag : Pflicht nur, wenn das Angebot neu angelegt wird ; sonst gilt der Angebotspreis.",
        "zahlungsdatum : optional (TT.MM.JJJJ). zahlungsart : Karte / Überweisung / PayPal / Bar.",
        "Pro Zeile entstehen: bezahlte Mitgliedschaft + Zahlung + Jahresstatus „aktiv“. "
        "Das Mitglied wird nicht benachrichtigt.",
    ]
    for i, text in enumerate(hinweise, start=1):
        notes.cell(row=i, column=1, value=text)
    notes.column_dimensions["A"].width = 120
    return classeur


def analyser_beitraege(classeur, kopf, daten, dateiname: str = "") -> Analyse:
    """PRÜFPHASE des langen Formats — schreibt nichts."""
    identite = _map_identite(kopf)
    spalten = _spalten(kopf)
    if "email" not in identite:
        raise ImportDateiFehler("Pflichtspalten fehlen: email")
    if "angebot" not in spalten:
        raise ImportDateiFehler("Pflichtspalten fehlen: angebot")

    cache_email, cache_cin = _lade_caches()
    historie = {
        (mid, jahr): statut
        for mid, jahr, statut in HistoriqueStatutMembre.objects.values_list(
            "membre_id", "annee", "statut"
        )
    }
    gesehen = {}  # (membre_id, annee) -> erste Zeilennummer
    zeilen = []

    for nummer, row in daten:
        email = row[identite["email"]] if identite["email"] < len(row) else None
        cin = row[identite["cin"]] if "cin" in identite and identite["cin"] < len(row) else None
        anzeige = {"prenom": "", "nom": "", "email": str(email or "").strip()}

        def fehler(grund, membre=None):
            zeilen.append(
                ZeileAnalyse(
                    ligne=nummer,
                    status=STATUS_FEHLER,
                    grund=grund,
                    anzeige=anzeige,
                    existant=membre,
                )
            )

        try:
            membre = _resoudre_membre(email, cin, cache_email, cache_cin)
        except ValueError as exc:
            fehler(str(exc))
            continue
        anzeige.update(prenom=membre.prenom, nom=membre.nom)

        try:
            annee = _parse_jahr(_zelle(row, spalten, "jahr"))
            angebot = str(_zelle(row, spalten, "angebot") or "").strip()
            if not angebot:
                raise ValueError("Angebot fehlt")
            kampagne_name = str(_zelle(row, spalten, "kampagne") or "").strip()
            betrag_roh = _parse_betrag(_zelle(row, spalten, "betrag"))
            datum_roh = _parse_datum(_zelle(row, spalten, "datum"))
            modus = _parse_modus(_zelle(row, spalten, "modus"))
            kampagne = _finde_kampagne(annee, kampagne_name)
        except ValueError as exc:
            fehler(str(exc), membre)
            continue

        if (membre.id, annee) in gesehen:
            fehler(
                f"Mitglied + Jahr {annee} kommt mehrfach vor (bereits in Zeile "
                f"{gesehen[(membre.id, annee)]})",
                membre,
            )
            continue
        gesehen[(membre.id, annee)] = nummer

        offre = _finde_angebot(kampagne, angebot)
        if betrag_roh is None:
            if offre is None:
                fehler("Betrag erforderlich, da das Angebot neu angelegt werden muss", membre)
                continue
            betrag = offre.prix_plein
        else:
            betrag = betrag_roh
        datum = datum_roh or (kampagne.date_debut if kampagne else datetime.date(annee, 1, 1))

        bestehende = (
            Souscription.objects.select_related("offre", "cotisation")
            .filter(membre=membre, campagne=kampagne)
            .first()
            if kampagne
            else None
        )
        if bestehende is None:
            sous_art = ART_NEU
        elif (
            bestehende.statut == StatutSouscription.PAYEE
            and bestehende.offre.nom.lower() == angebot.lower()
            and bestehende.prix_paye == betrag
        ):
            sous_art = ART_GLEICH
        else:
            sous_art = ART_ABWEICHEND

        alt_hist = historie.get((membre.id, annee))
        if alt_hist is None:
            hist_art = ART_NEU
        elif alt_hist == StatutMembre.ACTIF:
            hist_art = ART_GLEICH
        else:
            hist_art = ART_ABWEICHEND

        aenderungen = []
        if sous_art != ART_GLEICH:
            if kampagne is None:
                aenderungen.append(
                    {
                        "champ": "kampagne",
                        "label": "Kampagne",
                        "alt": "",
                        "neu": f"{kampagne_name or f'Mitgliedschaft {annee}'} (neu, archiviert)",
                        "art": ART_NEU,
                    }
                )
            aenderungen.append(
                {
                    "champ": "angebot",
                    "label": "Angebot",
                    "alt": bestehende.offre.nom if bestehende else "",
                    "neu": angebot + (" (neu angelegt)" if offre is None else ""),
                    "art": sous_art,
                }
            )
            aenderungen.append(
                {
                    "champ": "betrag",
                    "label": "Beitrag",
                    "alt": _fmt_betrag(bestehende.prix_paye) if bestehende else "",
                    "neu": _fmt_betrag(betrag),
                    "art": sous_art,
                }
            )
            aenderungen.append(
                {
                    "champ": "datum",
                    "label": "Zahlungsdatum",
                    "alt": "",
                    "neu": _fmt_datum(datum) + (f" · {_MODUS_TEXT[modus]}" if modus else ""),
                    "art": sous_art,
                }
            )
        if hist_art != ART_GLEICH:
            aenderungen.append(
                {
                    "champ": "status",
                    "label": f"Status {annee}",
                    "alt": "inaktiv" if alt_hist else "",
                    "neu": "aktiv",
                    "art": hist_art,
                }
            )

        daten_zeile = BeitragsDaten(
            annee, kampagne_name, angebot, betrag, datum, modus, sous_art, hist_art
        )
        if not aenderungen:
            status, grund = STATUS_UNVERAENDERT, f"{annee}: bereits identisch vorhanden"
        elif ART_ABWEICHEND in (sous_art, hist_art):
            status, grund = STATUS_DUBLETTE, f"{annee}: bestehender Eintrag weicht ab"
        else:
            status, grund = STATUS_NEU, f"{annee}: Mitgliedschaft {angebot} · {_fmt_betrag(betrag)}"
        zeilen.append(
            ZeileAnalyse(
                ligne=nummer,
                status=status,
                grund=grund,
                anzeige=anzeige,
                existant=membre,
                aenderungen=aenderungen,
                daten=daten_zeile,
            )
        )

    return Analyse(zeilen=zeilen, classeur=classeur, dateiname=dateiname)


def _utc_mittag(datum: datetime.date) -> datetime.datetime:
    return datetime.datetime.combine(datum, datetime.time(12, 0), tzinfo=datetime.timezone.utc)


def _schreibe_souscription(membre: Membre, d: BeitragsDaten, kampagne, offre, importeur) -> None:
    souscription = (
        Souscription.objects.select_related("cotisation")
        .filter(membre=membre, campagne=kampagne)
        .first()
    ) or Souscription(membre=membre, campagne=kampagne, offre=offre)
    souscription.offre = offre
    souscription.rabais = None
    souscription.prix_paye = d.betrag
    souscription.statut = StatutSouscription.PAYEE
    souscription.date_souscription = _utc_mittag(d.datum)
    souscription.snapshot_avantages = offre.avantages
    souscription.save()

    libelle = f"Adhésion {kampagne.nom} — {offre.nom}"
    cotisation = souscription.cotisation
    if cotisation is None and d.betrag > 0:
        cotisation = Cotisation(membre=membre, type_article=TypeArticle.ADHESION)
    if cotisation is not None and d.betrag > 0:
        cotisation.libelle = libelle
        cotisation.montant = d.betrag
        cotisation.statut = StatutCotisation.PAYEE
        cotisation.mode_paiement = d.modus
        cotisation.annee = d.annee
        cotisation.date_paiement = _utc_mittag(d.datum)
        cotisation.saisie_par = importeur
        cotisation.save()
        if souscription.cotisation_id != cotisation.id:
            souscription.cotisation = cotisation
            souscription.save(update_fields=["cotisation", "updated_at"])


def schreibe_beitrag_zeile(zeile: ZeileAnalyse, ueberschreiben: bool, importeur) -> ZeilenErgebnis:
    d: BeitragsDaten = zeile.daten
    membre = zeile.existant
    sous_zu = d.sous_art == ART_NEU or (d.sous_art == ART_ABWEICHEND and ueberschreiben)
    hist_zu = d.hist_art == ART_NEU or (d.hist_art == ART_ABWEICHEND and ueberschreiben)

    if not (sous_zu or hist_zu):
        return ZeilenErgebnis(
            zeile.ligne,
            ERG_UEBERSPRUNGEN,
            f"Dublette — nicht zum Überschreiben ausgewählt (Jahr {d.annee})",
        )
    try:
        with transaction.atomic():
            if sous_zu:
                kampagne = _finde_kampagne(d.annee, d.kampagne_name)
                if kampagne is None:
                    kampagne = CampagneAdhesion.objects.create(
                        nom=d.kampagne_name or f"Mitgliedschaft {d.annee}",
                        annee=d.annee,
                        date_debut=datetime.date(d.annee, 1, 1),
                        date_fin=datetime.date(d.annee, 12, 31),
                        description="Importiert (Historie)",
                        statut=StatutCampagne.CLOTUREE,
                        created_by=importeur or membre,
                    )
                offre = _finde_angebot(kampagne, d.angebot) or OffreAdhesion.objects.create(
                    campagne=kampagne,
                    nom=d.angebot,
                    prix_plein=d.betrag,
                    visible=False,
                    ordre=99,
                )
                _schreibe_souscription(membre, d, kampagne, offre, importeur)
            if hist_zu:
                enregistrer_statut_annuel(
                    membre,
                    d.annee,
                    StatutMembre.ACTIF,
                    RaisonChangementStatut.PAIEMENT_CONFIRME,
                    date_effet=datetime.datetime(d.annee, 1, 1, tzinfo=datetime.timezone.utc),
                    notifier_membre=False,
                )
    except Exception:  # noqa: BLE001 — eine fehlerhafte Zeile darf den Rest nicht stoppen
        logger.exception("import beitraege: Zeile %s nicht gespeichert", zeile.ligne)
        return ZeilenErgebnis(zeile.ligne, ERG_FEHLER, "Technischer Fehler beim Speichern")

    abweichend_offen = (d.sous_art == ART_ABWEICHEND and not sous_zu) or (
        d.hist_art == ART_ABWEICHEND and not hist_zu
    )
    if abweichend_offen:
        return ZeilenErgebnis(
            zeile.ligne, ERG_TEILWEISE, f"Jahr {d.annee}: nur neue Teile importiert (Rest Dublette)"
        )
    if ueberschreiben and ART_ABWEICHEND in (d.sous_art, d.hist_art):
        return ZeilenErgebnis(zeile.ligne, ERG_UEBERSCHRIEBEN, f"Jahr {d.annee} überschrieben")
    return ZeilenErgebnis(
        zeile.ligne, ERG_IMPORTIERT, f"Jahr {d.annee}: {d.angebot} · {_fmt_betrag(d.betrag)}"
    )
