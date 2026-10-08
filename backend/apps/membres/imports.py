"""
Import Excel des membres historiques — RICEFW C-001 (Import initial membres,
P0) / W-008 (workflow "Import Excel membres").

Portée volontairement limitée à ce que RICEFW décrit : lecture openpyxl,
validation de schéma + doublons email/CIN, chiffrement (délégué à
Membre.save() via EncryptedCharField), et un journal ligne par ligne
OK/KO (RICEFW R-012 "Journal import Master Data").

Décision importante : l'import est volontairement "best effort" ligne par
ligne (pas une transaction unique tout-ou-rien) — une ligne invalide est
consignée en erreur et n'empêche pas l'import des ~299 autres. C'est ce que
décrit W-008 ("générer journal OK/KO/erreurs"), qui suppose un import
partiel possible.

Limite connue (documentée, pas un bug) : cin/passeport sont chiffrés en
AES-256-GCM (nonce aléatoire par valeur — deux chiffrements du même CIN en
clair produisent des ciphertexts différents). Il est donc impossible de
poser un index unique BDD dessus. La détection de doublon CIN se fait donc
en clair, en mémoire, en déchiffrant les CIN déjà en base — acceptable pour
un import ponctuel de quelques centaines de lignes (C-001 : ~300), mais ne
pas réutiliser cette approche pour une détection de doublon "à l'échelle".

Depuis le 2026-10-08 l'import se fait en DEUX temps (voir import_gemeinsam) : `analyser_membres`
évalue le fichier sans rien écrire (liste détaillée neu / Dublette / unverändert / Fehler),
`ausfuehren_membres` réévalue puis écrit — les doublons ne sont écrasés que pour les lignes
explicitement choisies (uniquement les cellules renseignées, jamais le numéro de membre ; l'email
d'une fiche liée à un compte reste inchangé) — et produit le rapport (colonnes Status/Grund).
"""

import datetime
import logging

from django.db import transaction

from .import_gemeinsam import (
    ERG_FEHLER,
    ERG_IMPORTIERT,
    ERG_UEBERSCHRIEBEN,
    ERG_UEBERSPRUNGEN,
    ERG_UNVERAENDERT,
    STATUS_DUBLETTE,
    STATUS_FEHLER,
    STATUS_NEU,
    STATUS_UNVERAENDERT,
    Analyse,
    Ergebnis,
    ImportDateiFehler,
    ZeileAnalyse,
    ZeilenErgebnis,
    erzeuge_bericht,
    lies_kopfzeile_und_zeilen,
    oeffne_arbeitsmappe,
)
from .models import CIN_PLATZHALTER, Bundesland, Membre, Sexe, StatutMembre

logger = logging.getLogger(__name__)

# Colonnes attendues dans la feuille (1ère ligne = en-têtes). Plusieurs
# libellés tolérés par colonne (casse et accents ignorés à la comparaison).
REQUIRED_COLUMNS = {
    "prenom": ("prenom", "prénom"),
    "nom": ("nom",),
    "date_naissance": ("date_naissance", "date de naissance"),
    "email": ("email", "e-mail", "mail"),
    "adresse_de": ("adresse_de", "adresse (allemagne)", "adresse"),
    "ville_de": ("ville_de", "ville (allemagne)", "ville"),
    "date_adhesion": ("date_adhesion", "date d'adhesion", "date d'adhésion"),
}
# telephone et cin sont facultatifs depuis le 2026-10-07 (décision utilisateur : beaucoup
# d'anciens membres n'ont ni l'un ni l'autre dans les registres historiques).
OPTIONAL_COLUMNS = {
    "sexe": ("sexe",),
    "telephone": ("telephone", "téléphone", "tel", "tél"),
    "cin": ("cin",),
    "passeport": ("passeport",),
    "code_postal_de": ("code_postal_de", "code postal"),
    "land_de": ("land_de", "land", "bundesland", "region"),
    "ville_origine_tn": ("ville_origine_tn", "ville d'origine (tunisie)", "ville d'origine"),
    "gouvernorat_tn": ("gouvernorat_tn", "gouvernorat (tunisie)", "gouvernorat"),
    "statut": ("statut",),
}
ALL_COLUMNS = {**REQUIRED_COLUMNS, **OPTIONAL_COLUMNS}

# En-tête "canonique" à écrire dans le template (1er alias déclaré par colonne) + exemple de
# valeur (ligne 2) pour guider la saisie — utilisé à la fois par la commande manage.py
# generer_template_import_membres et par MembreImportTemplateView (téléchargement HTTP,
# RICEFW W-008/F-019) : une seule source pour ne pas dupliquer les en-têtes/exemples.
_EXEMPLE_TEMPLATE = {
    "prenom": "Riadh",
    "nom": "Bchini",
    "date_naissance": "15/03/1985",
    "sexe": "Homme",
    "email": "riadh.bchini@example.de",
    "telephone": "+49 170 1234567",
    "cin": "12345678",
    "passeport": "",
    "adresse_de": "Musterstr. 1",
    "code_postal_de": "10115",
    "ville_de": "Berlin",
    "land_de": "Berlin",
    "ville_origine_tn": "Tunis",
    "gouvernorat_tn": "Tunis",
    "statut": "actif",
    "date_adhesion": "01/09/2019",
}


def construire_classeur_template():
    """
    Construit (sans l'enregistrer) le classeur .xlsx vide utilisé pour l'import — en-têtes
    REQUIRED_COLUMNS + OPTIONAL_COLUMNS avec exemple de saisie, feuille "Notes" avec la
    légende. Ne dépend pas du système de fichiers : appelable aussi bien depuis la commande
    manage.py (sauvegarde disque) que depuis une vue HTTP (réponse en mémoire).
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Membres"

    champs = list(REQUIRED_COLUMNS) + list(OPTIONAL_COLUMNS)
    en_tete_style = Font(bold=True, color="FFFFFF")
    remplissage = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")

    for col_idx, champ in enumerate(champs, start=1):
        cellule = feuille.cell(row=1, column=col_idx, value=champ)
        cellule.font = en_tete_style
        cellule.fill = remplissage
        if champ not in REQUIRED_COLUMNS:
            cellule.comment = Comment("Optionnel", "CID")
        feuille.cell(row=2, column=col_idx, value=_EXEMPLE_TEMPLATE.get(champ, ""))
        feuille.column_dimensions[cellule.column_letter].width = max(len(champ) + 2, 14)

    notes = classeur.create_sheet("Notes")
    notes["A1"] = "Colonnes obligatoires : " + ", ".join(REQUIRED_COLUMNS)
    notes["A2"] = "Colonnes optionnelles : " + ", ".join(OPTIONAL_COLUMNS)
    notes["A3"] = "sexe : Homme / Femme (laisser vide si non renseigné)"
    notes["A4"] = "statut : actif / en_attente / inactif (défaut si vide : actif)"
    notes["A5"] = (
        "land_de : code à 2 lettres (BE, BY, ...) ou nom complet — Länder valides : "
        + ", ".join(f"{code} ({label})" for code, label in Bundesland.choices)
    )
    notes["A6"] = "Dates au format JJ/MM/AAAA."
    notes["A7"] = f"cin : si vide, remplacé automatiquement par {CIN_PLATZHALTER} (valeur neutre)."
    notes.column_dimensions["A"].width = 100

    return classeur


_SEXE_ALIASES = {"homme": Sexe.HOMME, "h": Sexe.HOMME, "femme": Sexe.FEMME, "f": Sexe.FEMME}
_STATUT_ALIASES = {
    "actif": StatutMembre.ACTIF,
    "en_attente": StatutMembre.EN_ATTENTE,
    "en attente": StatutMembre.EN_ATTENTE,
    "inactif": StatutMembre.INACTIF,
}
# Land : accepte le code à 2 lettres (BW, BY, ...) ou le nom complet
# (Baden-Württemberg, ...), casse ignorée.
_LAND_BY_LABEL = {label.lower(): code for code, label in Bundesland.choices}

# Champs de la fiche importable, dans l'ordre d'affichage, avec le libellé (allemand) utilisé
# dans la liste de contrôle et le rapport.
CHAMPS_IMPORT = {
    "prenom": "Vorname",
    "nom": "Nachname",
    "date_naissance": "Geburtsdatum",
    "sexe": "Geschlecht",
    "email": "E-Mail",
    "telephone": "Telefon",
    "cin": "CIN",
    "passeport": "Reisepass",
    "adresse_de": "Adresse",
    "code_postal_de": "PLZ",
    "ville_de": "Stadt",
    "land_de": "Bundesland",
    "ville_origine_tn": "Herkunftsstadt (TN)",
    "gouvernorat_tn": "Gouvernorat (TN)",
    "statut": "Status",
    "date_adhesion": "Beitrittsdatum",
}
_CHAMPS_MASQUES = ("cin", "passeport")  # jamais en clair dans la liste de contrôle
_LIBELLES_CHOIX = {
    "sexe": dict(Sexe.choices),
    "statut": dict(StatutMembre.choices),
    "land_de": dict(Bundesland.choices),
}


def _normalize(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _texte_ou_vide(valeur) -> str:
    """Cellule -> texte. Excel stocke souvent CIN/téléphone comme nombres (12345678.0)."""
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return str(valeur).strip()


def _map_headers(header_row) -> dict:
    """Retourne {nom_champ: index_colonne} pour les colonnes reconnues."""
    normalized = [_normalize(cell) for cell in header_row]
    mapping = {}
    for field_name, aliases in ALL_COLUMNS.items():
        for idx, header in enumerate(normalized):
            if header in aliases:
                mapping[field_name] = idx
                break

    manquantes = [f for f in REQUIRED_COLUMNS if f not in mapping]
    if manquantes:
        raise ImportDateiFehler("Pflichtspalten fehlen: " + ", ".join(sorted(manquantes)))
    return mapping


def _parse_date(value, champ: str) -> datetime.date:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str) and value.strip():
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    raise ValueError(f"{champ} ungültig (Datum erwartet, erhalten: {value!r})")


def _parse_row(row, mapping: dict):
    """
    Construit (sans l'enregistrer) un Membre à partir d'une ligne et renvoie
    (membre, gefuellt) — `gefuellt` = ensemble des champs réellement renseignés dans le fichier
    (sert à n'écraser que les cellules remplies). Lève ValueError avec un message explicite si
    la ligne est invalide.
    """

    def cell(field_name, required=True):
        idx = mapping.get(field_name)
        value = row[idx] if idx is not None and idx < len(row) else None
        if isinstance(value, str):
            value = value.strip()
        if required and (value is None or value == ""):
            raise ValueError(f"Pflichtfeld fehlt: {field_name}")
        return value

    prenom = cell("prenom")
    nom = cell("nom")
    date_naissance = _parse_date(cell("date_naissance"), "date_naissance")
    email = cell("email")
    telephone = _texte_ou_vide(cell("telephone", required=False))
    # CIN leer -> Platzhalter "00000000" (kein echter Identifikator, siehe CIN_PLATZHALTER).
    cin = _texte_ou_vide(cell("cin", required=False)) or CIN_PLATZHALTER
    cin_reel = cin != CIN_PLATZHALTER
    adresse_de = cell("adresse_de")
    ville_de = cell("ville_de")
    date_adhesion = _parse_date(cell("date_adhesion"), "date_adhesion")

    if "@" not in email:
        raise ValueError(f"E-Mail ungültig: {email!r}")

    sexe_brut = _normalize(cell("sexe", required=False))
    sexe = _SEXE_ALIASES.get(sexe_brut, Sexe.NON_RENSEIGNE)

    statut_brut = _normalize(cell("statut", required=False))
    # Import historique = membres déjà actifs dans l'association (pas de
    # nouvelles inscriptions à valider) — défaut ACTIF, pas EN_ATTENTE.
    statut = _STATUT_ALIASES.get(statut_brut, StatutMembre.ACTIF)

    land_brut = cell("land_de", required=False)
    land_de = ""
    if land_brut:
        land_brut_norm = _normalize(land_brut)
        land_de = (
            land_brut_norm.upper()
            if land_brut_norm.upper() in Bundesland.values
            else _LAND_BY_LABEL.get(land_brut_norm, "")
        )
        if not land_de:
            raise ValueError(f"Bundesland nicht erkannt: {land_brut!r}")

    passeport = _texte_ou_vide(cell("passeport", required=False)) or None

    membre = Membre(
        prenom=prenom,
        nom=nom,
        date_naissance=date_naissance,
        sexe=sexe,
        email=email,
        telephone=telephone,
        cin=cin,
        passeport=passeport,
        adresse_de=adresse_de,
        code_postal_de=_texte_ou_vide(cell("code_postal_de", required=False)),
        ville_de=ville_de,
        land_de=land_de,
        ville_origine_tn=_texte_ou_vide(cell("ville_origine_tn", required=False)),
        gouvernorat_tn=_texte_ou_vide(cell("gouvernorat_tn", required=False)),
        statut=statut,
        date_adhesion=date_adhesion,
    )

    gefuellt = {champ for champ in CHAMPS_IMPORT if getattr(membre, champ)}
    if not cin_reel:
        gefuellt.discard("cin")
    if sexe == Sexe.NON_RENSEIGNE:
        gefuellt.discard("sexe")
    # Statut : nur „gefüllt", wenn die Zelle einen bekannten Wert enthält (sonst gilt der
    # Standardwert ACTIF, der bestehende Mitglieder nie überschreiben darf).
    if statut_brut not in _STATUT_ALIASES:
        gefuellt.discard("statut")
    return membre, gefuellt


def _anzeige(champ, wert) -> str:
    """Darstellungstext eines Feldwerts in der Prüfliste (CIN/Reisepass maskiert)."""
    if wert in (None, ""):
        return ""
    if isinstance(wert, datetime.date):
        return wert.strftime("%d.%m.%Y")
    if champ in _CHAMPS_MASQUES:
        text = str(wert)
        return "•" * max(len(text) - 2, 2) + text[-2:]
    if champ in _LIBELLES_CHOIX:
        return str(_LIBELLES_CHOIX[champ].get(wert, wert))
    return str(wert)


def _gleich(champ, alt, neu) -> bool:
    if isinstance(neu, datetime.date) or isinstance(alt, datetime.date):
        return alt == neu
    a, n = str(alt or "").strip(), str(neu or "").strip()
    return a.lower() == n.lower() if champ == "email" else a == n


def berechne_aenderungen(importiert: Membre, existant: Membre, gefuellt: set) -> list:
    """
    Unterschiede zwischen Importzeile und bestehender Fiche — NUR gefüllte Zellen
    (Entscheidung 2026-10-08). Die E-Mail einer Fiche mit Benutzerkonto bleibt unverändert (sie
    ist an das Login gebunden) ; die Mitgliedsnummer ist nicht importierbar und bleibt immer.
    """
    aenderungen = []
    for champ, label in CHAMPS_IMPORT.items():
        if champ not in gefuellt:
            continue
        if champ == "email" and existant.user_id is not None:
            continue
        alt, neu = getattr(existant, champ), getattr(importiert, champ)
        if _gleich(champ, alt, neu):
            continue
        aenderungen.append(
            {
                "champ": champ,
                "label": label,
                "alt": _anzeige(champ, alt),
                "neu": _anzeige(champ, neu),
                "art": "abweichend",
            }
        )
    return aenderungen


def _vorhandene_fichen() -> tuple:
    """
    ({email_klein: [Membre]}, {cin: [Membre]}) der bereits gespeicherten Fiches. CIN werden
    im Speicher entschlüsselt (siehe Modul-Docstring : kein DB-Index auf AES-GCM-Feldern
    möglich) ; der Platzhalter 00000000 zählt nie als Treffer.
    """
    par_email, par_cin = {}, {}
    for membre in Membre.objects.select_related("user"):
        par_email.setdefault(membre.email.strip().lower(), []).append(membre)
        if membre.cin and membre.cin != CIN_PLATZHALTER:
            par_cin.setdefault(membre.cin, []).append(membre)
    return par_email, par_cin


def _bezeichnung(membre: Membre) -> str:
    return f"{membre.numero_membre} ({membre.prenom} {membre.nom})"


def analyser_membres(fichier, dateiname: str = "") -> Analyse:
    """
    PRÜFPHASE — wertet die Datei vollständig aus, schreibt nichts. `fichier` : .xlsx-Objekt
    (z. B. InMemoryUploadedFile), 1. Tabelle, 1. Zeile = Kopfzeilen. Wirft ImportDateiFehler.
    """
    classeur = oeffne_arbeitsmappe(fichier)
    kopf, daten = lies_kopfzeile_und_zeilen(classeur)
    mapping = _map_headers(kopf)

    par_email, par_cin = _vorhandene_fichen()
    emails_vus, cins_vus = {}, {}  # Wert -> Zeilennummer der ersten Fundstelle

    zeilen = []
    for nummer, row in daten:
        try:
            membre, gefuellt = _parse_row(row, mapping)
        except ValueError as exc:
            zeilen.append(ZeileAnalyse(ligne=nummer, status=STATUS_FEHLER, grund=str(exc)))
            continue

        anzeige = {"prenom": membre.prenom, "nom": membre.nom, "email": membre.email}
        email_norm = membre.email.strip().lower()
        cin_reel = "cin" in gefuellt

        erste_zeile = emails_vus.get(email_norm) or (cin_reel and cins_vus.get(membre.cin))
        if erste_zeile:
            zeilen.append(
                ZeileAnalyse(
                    ligne=nummer,
                    status=STATUS_FEHLER,
                    grund=(
                        f"Doppelt in der Datei (E-Mail oder CIN bereits in Zeile {erste_zeile})"
                    ),
                    anzeige=anzeige,
                )
            )
            continue
        emails_vus[email_norm] = nummer
        if cin_reel:
            cins_vus[membre.cin] = nummer

        treffer = {m.id: m for m in par_email.get(email_norm, [])}
        treffer_cin = {m.id: m for m in par_cin.get(membre.cin, [])} if cin_reel else {}
        treffer.update(treffer_cin)

        if len(treffer) > 1:
            nummern = ", ".join(sorted(_bezeichnung(m) for m in treffer.values()))
            zeilen.append(
                ZeileAnalyse(
                    ligne=nummer,
                    status=STATUS_FEHLER,
                    grund=(
                        "E-Mail/CIN passen zu mehreren verschiedenen Mitgliedern "
                        f"({nummern}) — bitte manuell klären"
                    ),
                    anzeige=anzeige,
                )
            )
            continue

        if not treffer:
            zeilen.append(
                ZeileAnalyse(
                    ligne=nummer,
                    status=STATUS_NEU,
                    grund="Neues Mitglied",
                    anzeige=anzeige,
                    daten=membre,
                )
            )
            continue

        existant = next(iter(treffer.values()))
        kriterien = []
        if existant.id in {m.id for m in par_email.get(email_norm, [])}:
            kriterien.append("E-Mail")
        if existant.id in treffer_cin:
            kriterien.append("CIN")
        aenderungen = berechne_aenderungen(membre, existant, gefuellt)
        grund = f"Dublette von {_bezeichnung(existant)} (gleiche {' und '.join(kriterien)})"
        if aenderungen:
            status = STATUS_DUBLETTE
        else:
            status = STATUS_UNVERAENDERT
            grund += " — keine Abweichungen"
        zeilen.append(
            ZeileAnalyse(
                ligne=nummer,
                status=status,
                grund=grund,
                anzeige=anzeige,
                existant=existant,
                aenderungen=aenderungen,
                daten=membre,
            )
        )

    return Analyse(zeilen=zeilen, classeur=classeur, dateiname=dateiname)


def ausfuehren_membres(fichier, ueberschreiben, dateiname: str = "") -> Ergebnis:
    """
    BESTÄTIGUNGSPHASE — wertet die Datei erneut aus (die DB kann sich seit der Prüfung geändert
    haben), importiert alle neuen gültigen Zeilen, überschreibt Dubletten NUR für die Zeilen-
    nummern in `ueberschreiben` und liefert den Bericht (Originaldatei + Status/Grund).
    Jede Zeile läuft in einer eigenen Transaktion : ein Datenbankfehler betrifft nur sie.
    """
    ueberschreiben = set(ueberschreiben or ())
    analyse = analyser_membres(fichier, dateiname)

    ergebnisse = []
    for zeile in analyse.zeilen:
        if zeile.status == STATUS_FEHLER:
            ergebnisse.append(ZeilenErgebnis(zeile.ligne, ERG_FEHLER, zeile.grund))
        elif zeile.status == STATUS_UNVERAENDERT:
            ergebnisse.append(ZeilenErgebnis(zeile.ligne, ERG_UNVERAENDERT, zeile.grund))
        elif zeile.status == STATUS_NEU:
            ergebnisse.append(_neu_speichern(zeile))
        elif zeile.ligne in ueberschreiben:
            ergebnisse.append(_ueberschreiben(zeile))
        else:
            ergebnisse.append(
                ZeilenErgebnis(
                    zeile.ligne,
                    ERG_UEBERSPRUNGEN,
                    f"{zeile.grund} — nicht zum Überschreiben ausgewählt",
                )
            )

    name, inhalt = erzeuge_bericht(analyse.classeur, ergebnisse, dateiname)
    logger.info(
        "import membres: %s",
        {
            e: sum(1 for r in ergebnisse if r.ergebnis == e)
            for e in {r.ergebnis for r in ergebnisse}
        },
    )
    return Ergebnis(zeilen=ergebnisse, bericht_name=name, bericht_bytes=inhalt)


def _neu_speichern(zeile: ZeileAnalyse) -> ZeilenErgebnis:
    # Ligne par ligne (pas bulk_create) : Membre.save() génère numero_membre via une requête BDD
    # (CA-<année>-<compteur>) — voir apps.membres.models._generate_numero_membre.
    try:
        with transaction.atomic():
            zeile.daten.save()
    except Exception:  # noqa: BLE001 — eine fehlerhafte Zeile darf den Rest nicht stoppen
        logger.exception("import membres: Zeile %s nicht gespeichert", zeile.ligne)
        return ZeilenErgebnis(zeile.ligne, ERG_FEHLER, "Technischer Fehler beim Speichern")
    return ZeilenErgebnis(
        zeile.ligne, ERG_IMPORTIERT, f"Neues Mitglied {zeile.daten.numero_membre} angelegt"
    )


def _ueberschreiben(zeile: ZeileAnalyse) -> ZeilenErgebnis:
    existant, importiert = zeile.existant, zeile.daten
    felder = [a["champ"] for a in zeile.aenderungen]
    try:
        with transaction.atomic():
            for champ in felder:
                setattr(existant, champ, getattr(importiert, champ))
            existant.save(update_fields=[*felder, "updated_at"])
    except Exception:  # noqa: BLE001
        logger.exception("import membres: Zeile %s nicht überschrieben", zeile.ligne)
        return ZeilenErgebnis(zeile.ligne, ERG_FEHLER, "Technischer Fehler beim Speichern")
    liste = ", ".join(a["label"] for a in zeile.aenderungen)
    return ZeilenErgebnis(
        zeile.ligne,
        ERG_UEBERSCHRIEBEN,
        f"Stammdaten von {_bezeichnung(existant)} überschrieben: {liste}",
    )
