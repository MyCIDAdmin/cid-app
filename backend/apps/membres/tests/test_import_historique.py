"""
Tests import Excel de l'historique de statut associatif par année (demande utilisateur du
2026-09-19) — voir apps.membres.imports_historique. Zweistufig seit 2026-10-08 :
analyser_historique / ausfuehren_historique und die Endpunkte pruefen/ + bestaetigen/.
"""

import base64
import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.membres.import_gemeinsam import ImportDateiFehler
from apps.membres.imports_historique import analyser_historique, ausfuehren_historique
from apps.membres.models import HistoriqueStatutMembre, RaisonChangementStatut, StatutMembre
from apps.membres.tests.factories import MembreFactory

pytestmark = pytest.mark.django_db

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_file(headers, rows, filename="historique.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(filename, buf.getvalue(), content_type=XLSX_MIME)


def _eintrag(membre, annee, statut):
    return HistoriqueStatutMembre.objects.create(
        membre=membre,
        annee=annee,
        statut=statut,
        raison=RaisonChangementStatut.ECHEANCE_DEPASSEE,
        date_effet=f"{annee}-01-01T00:00:00Z",
    )


def _zeilen(analyse):
    return {z.ligne: z for z in analyse.zeilen}


def _ergebnisse(ergebnis):
    return {z.ligne: z for z in ergebnis.zeilen}


# --- Prüfphase ---


def test_pruefen_schreibt_nichts_und_klassifiziert_neu():
    membre = MembreFactory(email="riadh@example.de", cin="11112222")
    fichier = _xlsx_file(
        ["email", "cin", "2023", "2024", "2025"],
        [["riadh@example.de", "11112222", "inactif", "actif", "actif"]],
    )

    analyse = analyser_historique(fichier, "h.xlsx")

    zeile = _zeilen(analyse)[2]
    assert zeile.status == "neu"
    assert zeile.existant == membre
    assert [a["champ"] for a in zeile.aenderungen] == ["2023", "2024", "2025"]
    assert zeile.aenderungen[0]["neu"] == "inaktiv"
    assert HistoriqueStatutMembre.objects.count() == 0


def test_pruefen_abweichender_bestehender_eintrag_ist_dublette():
    membre = MembreFactory(email="ecrase@example.de", cin="55556666")
    _eintrag(membre, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(["email", "cin", "2024"], [["ecrase@example.de", "55556666", "actif"]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "dublette"
    assert zeile.ueberschreibbar
    assert zeile.aenderungen == [
        {"champ": "2024", "label": "2024", "alt": "inaktiv", "neu": "aktiv", "art": "abweichend"}
    ]


def test_pruefen_identische_eintraege_sind_unveraendert():
    membre = MembreFactory(email="gleich@example.de", cin="55556667")
    _eintrag(membre, 2024, StatutMembre.ACTIF)
    fichier = _xlsx_file(["email", "cin", "2024"], [["gleich@example.de", "55556667", "actif"]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "unveraendert"
    assert not zeile.ueberschreibbar


def test_pruefen_leere_zelle_wird_ignoriert():
    MembreFactory(email="vide@example.de", cin="77778888")
    fichier = _xlsx_file(
        ["email", "cin", "2024", "2025"], [["vide@example.de", "77778888", "actif", ""]]
    )

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert [a["champ"] for a in zeile.aenderungen] == ["2024"]


def test_pruefen_zeile_ohne_jahreswerte_ist_unveraendert():
    MembreFactory(email="leer@example.de", cin="77778889")
    fichier = _xlsx_file(["email", "cin", "2024"], [["leer@example.de", "77778889", ""]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "unveraendert"


def test_pruefen_ungueltiger_wert_macht_die_ganze_zeile_zum_fehler():
    MembreFactory(email="invalide@example.de", cin="12123434")
    fichier = _xlsx_file(
        ["email", "cin", "2023", "2024"],
        [["invalide@example.de", "12123434", "actif", "peut-etre"]],
    )

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "fehler"
    assert "2024" in zeile.grund


def test_pruefen_mitglied_unbekannt_ist_fehler():
    fichier = _xlsx_file(["email", "cin", "2024"], [["inconnu@example.de", "00000000", "actif"]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "fehler"
    assert "inconnu@example.de" in zeile.grund


def test_pruefen_email_und_cin_zweier_mitglieder_ist_mehrdeutig():
    MembreFactory(email="a@example.de", cin="11111111")
    MembreFactory(email="b@example.de", cin="22222222")
    fichier = _xlsx_file(["email", "cin", "2024"], [["a@example.de", "22222222", "actif"]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "fehler"
    assert "2 verschiedenen" in zeile.grund


def test_pruefen_nur_cin_passt_genuegt():
    membre = MembreFactory(email="ancien-mail@example.de", cin="33334444")
    fichier = _xlsx_file(["email", "cin", "2024"], [["autre@example.de", "33334444", "actif"]])

    zeile = _zeilen(analyser_historique(fichier))[2]

    assert zeile.status == "neu"
    assert zeile.existant == membre


def test_pruefen_mitglied_zweimal_in_der_datei_ist_fehler():
    MembreFactory(email="doppelt@example.de", cin="44445555")
    fichier = _xlsx_file(
        ["email", "2024", "2025"],
        [["doppelt@example.de", "actif", ""], ["DOPPELT@example.de", "", "actif"]],
    )

    zeilen = _zeilen(analyser_historique(fichier))

    assert zeilen[2].status == "neu"
    assert zeilen[3].status == "fehler"
    assert "Zeile 2" in zeilen[3].grund


def test_pruefen_leere_zeile_wird_ignoriert():
    fichier = _xlsx_file(["email", "cin", "2024"], [[None, None, None]])

    assert analyser_historique(fichier).zeilen == []


def test_pruefen_ohne_pflichtspalten_oder_jahresspalte_wirft():
    with pytest.raises(ImportDateiFehler):
        analyser_historique(_xlsx_file(["prenom", "2024"], [["Riadh", "actif"]]))
    with pytest.raises(ImportDateiFehler):
        analyser_historique(_xlsx_file(["email", "cin"], [["riadh@example.de", "11112222"]]))


def test_pruefen_ohne_cin_spalte_funktioniert_mit_email():
    membre = MembreFactory(email="sanscin@example.de", cin=None)
    fichier = _xlsx_file(["email", "2024"], [["sanscin@example.de", "actif"]])

    assert _zeilen(analyser_historique(fichier))[2].existant == membre


def test_pruefen_cin_platzhalter_identifiziert_kein_mitglied():
    MembreFactory(email="a@example.de", cin="00000000")
    MembreFactory(email="b@example.de", cin="00000000")
    fichier = _xlsx_file(["email", "cin", "2024"], [["inconnu@example.de", "00000000", "actif"]])

    assert _zeilen(analyser_historique(fichier))[2].status == "fehler"


# --- Bestätigungsphase ---


def test_ausfuehren_schreibt_pro_jahr_und_synchronisiert_das_aktuelle_statut():
    membre = MembreFactory(email="riadh@example.de", cin="11112222", statut=StatutMembre.EN_ATTENTE)
    fichier = _xlsx_file(
        ["email", "cin", "2023", "2024", "2025"],
        [["riadh@example.de", "11112222", "inactif", "actif", "actif"]],
    )

    ergebnis = ausfuehren_historique(fichier, set(), "h.xlsx")

    assert ergebnis.zaehler()["importiert"] == 1
    entrees = {h.annee: h.statut for h in HistoriqueStatutMembre.objects.filter(membre=membre)}
    assert entrees == {
        2023: StatutMembre.INACTIF,
        2024: StatutMembre.ACTIF,
        2025: StatutMembre.ACTIF,
    }
    assert all(
        h.raison == RaisonChangementStatut.MANUEL
        for h in HistoriqueStatutMembre.objects.filter(membre=membre)
    )
    membre.refresh_from_db()
    assert membre.statut == StatutMembre.ACTIF  # année la plus récente (2025)


def test_ausfuehren_benachrichtigt_nicht_und_sendet_keine_mail(mailoutbox):
    from apps.notifications.models import Notification

    user = User.objects.create_user(
        email="silencieux@example.de", password="Password123!", is_active=True
    )
    MembreFactory(
        user=user, email="silencieux@example.de", cin="99998888", statut=StatutMembre.INACTIF
    )
    fichier = _xlsx_file(["email", "cin", "2026"], [["silencieux@example.de", "99998888", "actif"]])

    ausfuehren_historique(fichier, set())

    assert Notification.objects.count() == 0
    assert len(mailoutbox) == 0


def test_ausfuehren_dublette_ohne_auswahl_wird_nicht_ueberschrieben():
    membre = MembreFactory(email="ecrase@example.de", cin="55556666")
    _eintrag(membre, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(["email", "cin", "2024"], [["ecrase@example.de", "55556666", "actif"]])

    ergebnis = ausfuehren_historique(fichier, set())

    assert _ergebnisse(ergebnis)[2].ergebnis == "uebersprungen"
    assert HistoriqueStatutMembre.objects.get(membre=membre, annee=2024).statut == "inactif"


def test_ausfuehren_dublette_mit_auswahl_wird_ueberschrieben():
    membre = MembreFactory(email="ecrase@example.de", cin="55556666")
    _eintrag(membre, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(["email", "cin", "2024"], [["ecrase@example.de", "55556666", "actif"]])

    ergebnis = ausfuehren_historique(fichier, {2})

    assert _ergebnisse(ergebnis)[2].ergebnis == "ueberschrieben"
    eintrag = HistoriqueStatutMembre.objects.get(membre=membre, annee=2024)
    assert eintrag.statut == StatutMembre.ACTIF
    assert eintrag.raison == RaisonChangementStatut.MANUEL


def test_ausfuehren_dublette_ohne_auswahl_importiert_trotzdem_die_neuen_jahre():
    membre = MembreFactory(email="gemischt@example.de", cin="55556668")
    _eintrag(membre, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(
        ["email", "cin", "2023", "2024"], [["gemischt@example.de", "55556668", "actif", "actif"]]
    )

    ergebnis = ausfuehren_historique(fichier, set())

    assert _ergebnisse(ergebnis)[2].ergebnis == "teilweise"
    assert HistoriqueStatutMembre.objects.get(membre=membre, annee=2023).statut == "actif"
    assert HistoriqueStatutMembre.objects.get(membre=membre, annee=2024).statut == "inactif"


def test_ausfuehren_nur_ausgewaehlte_zeilen_werden_ueberschrieben():
    a = MembreFactory(email="a@example.de", cin="10000001")
    b = MembreFactory(email="b@example.de", cin="10000002")
    _eintrag(a, 2024, StatutMembre.INACTIF)
    _eintrag(b, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(["email", "2024"], [["a@example.de", "actif"], ["b@example.de", "actif"]])

    ausfuehren_historique(fichier, {3})

    assert HistoriqueStatutMembre.objects.get(membre=a, annee=2024).statut == "inactif"
    assert HistoriqueStatutMembre.objects.get(membre=b, annee=2024).statut == "actif"


def test_ausfuehren_beruehrt_das_aktuelle_statut_bei_altem_nachtrag_nicht():
    membre = MembreFactory(
        email="rattrapage@example.de", cin="66667777", statut=StatutMembre.INACTIF
    )
    _eintrag(membre, 2026, StatutMembre.INACTIF)
    fichier = _xlsx_file(["email", "cin", "2023"], [["rattrapage@example.de", "66667777", "actif"]])

    ausfuehren_historique(fichier, set())

    membre.refresh_from_db()
    assert membre.statut == StatutMembre.INACTIF
    assert HistoriqueStatutMembre.objects.get(membre=membre, annee=2023).statut == "actif"


def test_ausfuehren_fehlerzeile_wird_nicht_geschrieben():
    MembreFactory(email="x@example.de", cin="12121212")
    fichier = _xlsx_file(["email", "2024"], [["x@example.de", "vielleicht"]])

    ergebnis = ausfuehren_historique(fichier, {2})

    assert _ergebnisse(ergebnis)[2].ergebnis == "fehler"
    assert HistoriqueStatutMembre.objects.count() == 0


# --- Vue API ---


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def rh_user():
    return User.objects.create_user(
        email="rh-hist@example.de", password="Password123!", role=Role.RH, is_active=True
    )


@pytest.fixture
def membre_user():
    return User.objects.create_user(
        email="membre-hist@example.de", password="Password123!", role=Role.MEMBRE, is_active=True
    )


def _auth(api_client, user):
    api_client.force_authenticate(user=user)
    return api_client


def _pruefen(api_client, fichier):
    return api_client.post(
        reverse("membres:membre-import-historique-pruefen"),
        {"fichier": fichier},
        format="multipart",
    )


def _bestaetigen(api_client, fichier, ueberschreiben=None):
    data = {"fichier": fichier}
    if ueberschreiben is not None:
        data["ueberschreiben"] = ueberschreiben
    return api_client.post(
        reverse("membres:membre-import-historique-bestaetigen"), data, format="multipart"
    )


def test_vue_pruefen_ok_und_schreibt_nichts(api_client, rh_user):
    MembreFactory(email="vue@example.de", cin="10101010")
    fichier = _xlsx_file(["email", "cin", "2025"], [["vue@example.de", "10101010", "actif"]])

    resp = _pruefen(_auth(api_client, rh_user), fichier)

    assert resp.status_code == 200, resp.data
    assert resp.data["zaehler"]["neu"] == 1
    assert resp.data["zeilen"][0]["aenderungen"][0]["champ"] == "2025"
    assert HistoriqueStatutMembre.objects.count() == 0


def test_vue_bestaetigen_ueberschreibt_gewaehlte_zeile_und_liefert_bericht(api_client, rh_user):
    a = MembreFactory(email="a@example.de", cin="10000001")
    b = MembreFactory(email="b@example.de", cin="10000002")
    _eintrag(a, 2024, StatutMembre.INACTIF)
    _eintrag(b, 2024, StatutMembre.INACTIF)
    fichier = _xlsx_file(
        ["email", "2024"],
        [["a@example.de", "actif"], ["b@example.de", "actif"], ["x@example.de", "actif"]],
        filename="historie 2024.xlsx",
    )

    resp = _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["2"])

    assert resp.status_code == 200, resp.data
    assert resp.data["zaehler"]["ueberschrieben"] == 1
    assert resp.data["zaehler"]["uebersprungen"] == 1
    assert resp.data["zaehler"]["fehler"] == 1
    assert HistoriqueStatutMembre.objects.get(membre=a, annee=2024).statut == "actif"
    assert HistoriqueStatutMembre.objects.get(membre=b, annee=2024).statut == "inactif"

    assert resp.data["bericht"]["dateiname"] == "historie_2024_bericht.xlsx"
    blatt = openpyxl.load_workbook(
        io.BytesIO(base64.b64decode(resp.data["bericht"]["inhalt_base64"]))
    ).worksheets[0]
    assert [c.value for c in blatt[1]] == ["email", "2024", "Status", "Grund"]
    assert [blatt.cell(row=r, column=3).value for r in (2, 3, 4)] == [
        "Überschrieben",
        "Übersprungen",
        "Fehler",
    ]


def test_vue_bestaetigen_ungueltige_auswahl_400(api_client, rh_user):
    fichier = _xlsx_file(["email", "2025"], [["x@example.de", "actif"]])
    resp = _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["zwei"])
    assert resp.status_code == 400
    assert resp.data["code"] == "ueberschreiben_ungueltig"


@pytest.mark.parametrize("aktion", [_pruefen, _bestaetigen])
def test_vue_comme_membre_refuse_403(api_client, membre_user, aktion):
    fichier = _xlsx_file(["email", "cin", "2025"], [["x@example.de", "1", "actif"]])
    assert aktion(_auth(api_client, membre_user), fichier).status_code == 403


@pytest.mark.parametrize("aktion", [_pruefen, _bestaetigen])
def test_vue_non_authentifie_401(api_client, aktion):
    fichier = _xlsx_file(["email", "cin", "2025"], [["x@example.de", "1", "actif"]])
    assert aktion(api_client, fichier).status_code == 401


@pytest.mark.parametrize("aktion", [_pruefen, _bestaetigen])
def test_vue_colonnes_manquantes_400(api_client, rh_user, aktion):
    fichier = _xlsx_file(["prenom"], [["Riadh"]])
    resp = aktion(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 400
    assert resp.data["code"] == "schema_invalide"


@pytest.mark.parametrize("aktion", [_pruefen, _bestaetigen])
def test_vue_fichier_non_excel_rejete_400(api_client, rh_user, aktion):
    fichier = SimpleUploadedFile("f.txt", b"pas un excel", content_type="text/plain")
    resp = aktion(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 400
    assert resp.data["code"] == "type_fichier_invalide"


def test_vue_telechargement_template_comme_rh_ok(api_client, rh_user):
    _auth(api_client, rh_user)
    resp = api_client.get(reverse("membres:membre-import-historique-template"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME


def test_vue_telechargement_template_comme_membre_refuse_403(api_client, membre_user):
    _auth(api_client, membre_user)
    resp = api_client.get(reverse("membres:membre-import-historique-template"))
    assert resp.status_code == 403
