"""
Tests export Excel du répertoire des membres (demande utilisateur du 2026-09-16, RICEFW
R-002) — voir apps.membres.exports/export_views.
"""

import io
from datetime import date

import openpyxl
import pytest
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.cotisations.models import Cotisation, StatutCotisation, TypeArticle
from apps.membres.exports import CHAMPS_EXPORT
from apps.membres.models import Bundesland, StatutMembre
from apps.membres.tests.factories import MembreFactory

pytestmark = pytest.mark.django_db

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def rh_user():
    return User.objects.create_user(
        email="rh@example.de", password="Password123!", role=Role.RH, is_active=True
    )


@pytest.fixture
def bureau_admin_user():
    return User.objects.create_user(
        email="bureau@example.de", password="Password123!", role=Role.BUREAU_ADMIN, is_active=True
    )


@pytest.fixture
def membre_user():
    return User.objects.create_user(
        email="membre@example.de", password="Password123!", role=Role.MEMBRE, is_active=True
    )


def _auth(api_client, user):
    api_client.force_authenticate(user=user)
    return api_client


def _export(api_client, **params):
    return api_client.get(reverse("membres:membre-export"), params)


def _classeur(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def _lignes(resp):
    feuille = _classeur(resp).active
    lignes = list(feuille.iter_rows(values_only=True))
    return lignes[0], lignes[1:]  # (en-têtes, lignes de données)


# --- Cas nominal ---


def test_export_comme_rh_ok(api_client, rh_user):
    MembreFactory(prenom="Riadh", nom="Bchini", ville_de="Berlin")
    MembreFactory(prenom="Sana", nom="Werfelli", ville_de="Hambourg")

    resp = _export(_auth(api_client, rh_user))

    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert "export_membres_" in resp["Content-Disposition"]
    assert ".xlsx" in resp["Content-Disposition"]

    entetes, lignes = _lignes(resp)
    assert entetes[0] == "N° membre"
    assert "Statut" in entetes
    assert len(lignes) == 2


def test_export_comme_bureau_admin_ok(api_client, bureau_admin_user):
    MembreFactory()
    resp = _export(_auth(api_client, bureau_admin_user))
    assert resp.status_code == 200


def test_export_cin_toujours_masque(api_client, rh_user):
    MembreFactory(cin="99998888")
    resp = _export(_auth(api_client, rh_user))
    entetes, lignes = _lignes(resp)
    idx_cin = entetes.index("CIN")
    assert lignes[0][idx_cin] == "•••••888"
    assert "99998888" not in str(lignes[0])


def test_export_colonnes_identifiants_au_format_texte(api_client, rh_user):
    # Demande utilisateur du 2026-09-16 (2e) : N° membre, Téléphone, CIN, Code postal ne
    # doivent jamais être réinterprétés par Excel comme un nombre à l'ouverture.
    MembreFactory(cin="99998888", telephone="+4915112345678", code_postal_de="01067")
    resp = _export(_auth(api_client, rh_user))
    feuille = _classeur(resp).active
    entetes = list(next(feuille.iter_rows(min_row=1, max_row=1, values_only=True)))
    ligne = list(feuille.iter_rows(min_row=2, max_row=2))[0]

    for libelle in ("N° membre", "Téléphone", "CIN", "Code postal"):
        idx = entetes.index(libelle)
        assert ligne[idx].number_format == "@", libelle

    # Une colonne numérique "normale" (Âge) n'a pas de raison d'être forcée en texte.
    idx_age = entetes.index("Âge")
    assert ligne[idx_age].number_format != "@"


def test_export_colonne_cotisation_annee_en_cours(api_client, rh_user):
    membre_paye = MembreFactory()
    membre_impaye = MembreFactory()
    annee_courante = date.today().year
    Cotisation.objects.create(
        membre=membre_paye,
        type_article=TypeArticle.COTISATION,
        libelle="Cotisation annuelle",
        montant="45.00",
        statut=StatutCotisation.PAYEE,
        annee=annee_courante,
    )

    resp = _export(_auth(api_client, rh_user))
    entetes, lignes = _lignes(resp)
    idx_cotisation = entetes.index(f"Cotisation {annee_courante}")
    idx_email = entetes.index("Email")
    valeurs_par_email = {ligne[idx_email]: ligne[idx_cotisation] for ligne in lignes}
    assert valeurs_par_email[membre_paye.email] == "Payée"
    assert valeurs_par_email[membre_impaye.email] == "En attente"


# --- Sélection des colonnes (demande utilisateur du 2026-09-16, 2e) ---


def test_export_champs_selection_reduit_les_colonnes(api_client, rh_user):
    MembreFactory(prenom="Sana", nom="Werfelli")
    resp = _export(_auth(api_client, rh_user), champs="prenom,nom,email")
    entetes, lignes = _lignes(resp)
    assert list(entetes) == ["Prénom", "Nom", "Email"]
    assert len(lignes[0]) == 3


def test_export_champs_respecte_toujours_l_ordre_canonique(api_client, rh_user):
    MembreFactory(prenom="Sana", nom="Werfelli")
    # Ordre de sélection volontairement inversé par rapport à CHAMPS_EXPORT.
    resp = _export(_auth(api_client, rh_user), champs="email,nom,prenom")
    entetes, _ = _lignes(resp)
    assert list(entetes) == ["Prénom", "Nom", "Email"]


def test_export_champs_ignore_les_cles_inconnues(api_client, rh_user):
    MembreFactory()
    resp = _export(_auth(api_client, rh_user), champs="prenom,nom,mot_de_passe")
    entetes, _ = _lignes(resp)
    assert list(entetes) == ["Prénom", "Nom"]


def test_export_champs_aucune_cle_connue_retombe_sur_toutes_les_colonnes(api_client, rh_user):
    MembreFactory()
    resp = _export(_auth(api_client, rh_user), champs="inconnu1,inconnu2")
    entetes, _ = _lignes(resp)
    assert len(entetes) == len(CHAMPS_EXPORT)


def test_export_sans_parametre_champs_retourne_toutes_les_colonnes(api_client, rh_user):
    MembreFactory()
    resp = _export(_auth(api_client, rh_user))
    entetes, _ = _lignes(resp)
    assert len(entetes) == len(CHAMPS_EXPORT)


def test_export_champs_cotisation_libelle_inclut_l_annee(api_client, rh_user):
    MembreFactory()
    resp = _export(_auth(api_client, rh_user), champs="prenom,cotisation_annee_en_cours")
    entetes, _ = _lignes(resp)
    annee_courante = date.today().year
    assert list(entetes) == ["Prénom", f"Cotisation {annee_courante}"]


# --- Filtrage (même FilterSet que GET /membres/) ---


def test_export_filtre_par_statut(api_client, rh_user):
    MembreFactory(statut=StatutMembre.ACTIF)
    MembreFactory(statut=StatutMembre.INACTIF)

    resp = _export(_auth(api_client, rh_user), statut=StatutMembre.ACTIF)
    _, lignes = _lignes(resp)
    assert len(lignes) == 1


def test_export_filtre_par_ville(api_client, rh_user):
    MembreFactory(ville_de="Berlin")
    MembreFactory(ville_de="Munich")

    resp = _export(_auth(api_client, rh_user), ville="berl")
    _, lignes = _lignes(resp)
    assert len(lignes) == 1


def test_export_filtre_recherche_q(api_client, rh_user):
    MembreFactory(prenom="Ahmed", nom="Zribi")
    MembreFactory(prenom="Sana", nom="Werfelli")

    resp = _export(_auth(api_client, rh_user), q="zribi")
    _, lignes = _lignes(resp)
    assert len(lignes) == 1


# --- Tri ---


def test_export_tri_par_defaut_nom_prenom(api_client, rh_user):
    MembreFactory(nom="Zribi", prenom="Ahmed")
    MembreFactory(nom="Bchini", prenom="Riadh")

    resp = _export(_auth(api_client, rh_user))
    entetes, lignes = _lignes(resp)
    idx_nom = entetes.index("Nom")
    assert [ligne[idx_nom] for ligne in lignes] == ["Bchini", "Zribi"]


def test_export_tri_descendant_par_date_adhesion(api_client, rh_user):
    MembreFactory(nom="Ancien", date_adhesion=date(2018, 1, 1))
    MembreFactory(nom="Recent", date_adhesion=date(2024, 1, 1))

    resp = _export(_auth(api_client, rh_user), ordering="-date_adhesion")
    entetes, lignes = _lignes(resp)
    idx_nom = entetes.index("Nom")
    assert [ligne[idx_nom] for ligne in lignes] == ["Recent", "Ancien"]


def test_export_tri_champ_inconnu_retombe_sur_le_defaut(api_client, rh_user):
    MembreFactory(nom="Zribi")
    MembreFactory(nom="Bchini")

    resp = _export(_auth(api_client, rh_user), ordering="mot_de_passe")
    assert resp.status_code == 200
    entetes, lignes = _lignes(resp)
    idx_nom = entetes.index("Nom")
    assert [ligne[idx_nom] for ligne in lignes] == ["Bchini", "Zribi"]


# --- Permissions (IDOR / RBAC) ---


def test_export_comme_membre_refuse_403(api_client, membre_user):
    resp = _export(_auth(api_client, membre_user))
    assert resp.status_code == 403


def test_export_non_authentifie_401(api_client):
    resp = _export(api_client)
    assert resp.status_code == 401


def test_export_land_de_vide_ne_leve_pas_erreur(api_client, rh_user):
    # land_de blank=True (voir Membre) — get_land_de_display() sur une valeur vide ne doit pas
    # planter (choices ne couvre pas la chaîne vide) : la vue doit gérer ce cas explicitement.
    MembreFactory(land_de="")
    resp = _export(_auth(api_client, rh_user))
    assert resp.status_code == 200


def test_export_bundesland_affiche_le_libelle_complet(api_client, rh_user):
    MembreFactory(land_de=Bundesland.BAYERN)
    resp = _export(_auth(api_client, rh_user))
    entetes, lignes = _lignes(resp)
    idx_land = entetes.index("Land")
    assert lignes[0][idx_land] == "Bayern"
