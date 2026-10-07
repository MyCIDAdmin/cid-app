"""Tests — Mitglieder-Reporting (Mitgliederliste mit Historie, Aktivitäten, Filter, Export)."""

import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.adhesions.models import StatutSouscription
from apps.adhesions.tests.factories import SouscriptionFactory
from apps.boutique.models import StatutCommande
from apps.boutique.tests.factories import CommandeFactory
from apps.cotisations.models import StatutCotisation
from apps.cotisations.tests.factories import CotisationFactory
from apps.evenements.tests.factories import InscriptionFactory
from apps.membres.models import HistoriqueStatutMembre, RaisonChangementStatut, StatutMembre
from apps.membres.tests.factories import MembreFactory
from apps.projets.models import ProjetMitglied
from apps.projets.tests.factories import ProjetFactory

pytestmark = pytest.mark.django_db


def _client(role):
    email = f"{role}@example.com"
    user = User.objects.filter(email=email).first() or User.objects.create_user(
        email=email, password="Password123!", role=role, is_active=True
    )
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def _historie(membre, annee, statut):
    return HistoriqueStatutMembre.objects.create(
        membre=membre,
        annee=annee,
        statut=statut,
        raison=RaisonChangementStatut.MANUEL,
        date_effet=timezone.now(),
    )


@pytest.fixture
def daten():
    aktiv = MembreFactory(nom="Aktiv", prenom="Anna", pays="DE")
    ruhig = MembreFactory(nom="Ruhig", prenom="Rolf", pays="FR", statut=StatutMembre.INACTIF)
    SouscriptionFactory(membre=aktiv, prix_paye=Decimal("50.00"), statut=StatutSouscription.PAYEE)
    SouscriptionFactory(membre=aktiv, prix_paye=Decimal("40.00"), statut=StatutSouscription.ANNULEE)
    CommandeFactory(membre=aktiv, montant_total=Decimal("30.00"), statut=StatutCommande.LIVREE)
    InscriptionFactory(membre=aktiv, montant_paye=Decimal("35.00"))
    CotisationFactory(membre=aktiv, montant=Decimal("20.00"), statut=StatutCotisation.PAYEE)
    projet = ProjetFactory()
    CotisationFactory(
        membre=aktiv, montant=Decimal("15.00"), statut=StatutCotisation.PAYEE, projet=projet
    )
    ProjetMitglied.objects.create(projet=projet, membre=aktiv)
    _historie(aktiv, 2025, StatutMembre.ACTIF)
    _historie(ruhig, 2024, StatutMembre.ACTIF)
    _historie(ruhig, 2025, StatutMembre.INACTIF)
    return aktiv, ruhig


def _liste(client, **params):
    antwort = client.get(reverse("membres:membre-reporting"), params)
    assert antwort.status_code == 200, antwort.content
    return antwort.json()


def _namen(antwort):
    # Weitere Mitglieder entstehen als Nebeneffekt der Factories (Kampagnen-Ersteller usw.).
    return [z["nom"] for z in antwort["ergebnisse"] if z["nom"] in ("Aktiv", "Ruhig")]


def test_mitgliederliste_mit_historie_und_zahlen(daten):
    antwort = _liste(_client(Role.RH), historie_jahr="2025")
    assert antwort["count"] == 2
    aktiv = antwort["ergebnisse"][0]
    assert aktiv["nom"] == "Aktiv"
    assert aktiv["aktivitaeten"]["mitgliedschaft"] == 1  # stornierte zählt nicht
    assert aktiv["aktivitaeten"]["bestellung"] == 1
    assert aktiv["aktivitaeten"]["teilnahme"] == 1
    assert aktiv["aktivitaeten"]["beitrag"] == 1
    assert aktiv["aktivitaeten"]["projektbeitrag"] == 1
    assert aktiv["aktivitaeten"]["projektmitarbeit"] == 1
    assert aktiv["betrag_gesamt"] == "150.00"
    ruhig = antwort["ergebnisse"][1]
    assert [h["annee"] for h in ruhig["historie"]] == [2025, 2024]
    assert ruhig["aktivitaeten_gesamt"] == 0
    assert antwort["summen"]["betrag"] == "150.00"


def test_mitgliederliste_filter(daten):
    client = _client(Role.RH)
    assert _namen(_liste(client, statut="inactif")) == ["Ruhig"]
    assert _namen(_liste(client, pays="DE")) == ["Aktiv"]
    assert _namen(_liste(client, q="rolf")) == ["Ruhig"]
    assert _namen(_liste(client, historie_statut="inactif")) == ["Ruhig"]
    assert _namen(_liste(client, historie_jahr="2024")) == ["Ruhig"]
    assert _namen(_liste(client, typ="bestellung")) == ["Aktiv"]
    assert _namen(_liste(client, ohne_aktivitaet="1")) == ["Ruhig"]
    assert _namen(_liste(client, min_aktivitaeten="3")) == ["Aktiv"]
    assert _namen(_liste(client, min_aktivitaeten="99")) == []
    assert _namen(_liste(client, sortierung="aktivitaeten")) == ["Aktiv", "Ruhig"]


def test_mitgliederliste_zeitraum_begrenzt_aktivitaeten(daten):
    client = _client(Role.RH)
    zukunft = (datetime.date.today() + datetime.timedelta(days=400)).isoformat()
    assert _namen(_liste(client, typ="bestellung", von=zukunft)) == []


def test_mitgliederliste_seiten(daten):
    antwort = _liste(_client(Role.RH), historie_jahr="2025", page_size="1", page="2")
    assert antwort["count"] == 2
    assert _namen(antwort) == ["Ruhig"]


def test_aktivitaeten_liste_und_summen(daten):
    antwort = _client(Role.RH).get(reverse("membres:membre-reporting-aktivitaeten"))
    assert antwort.status_code == 200
    daten_ = antwort.json()
    typen = {z["typ"] for z in daten_["ergebnisse"]}
    assert typen == {
        "mitgliedschaft",
        "bestellung",
        "teilnahme",
        "beitrag",
        "projektbeitrag",
        "projektmitarbeit",
        "statuswechsel",
    }
    assert daten_["summen"]["mitgliedschaft"] == {"anzahl": 2, "betrag": "50.00"}
    assert daten_["summen"]["bestellung"]["betrag"] == "30.00"
    assert daten_["summen"]["projektmitarbeit"]["betrag"] == "0.00"
    daten_listen = [z["datum"] for z in daten_["ergebnisse"]]
    assert daten_listen == sorted(daten_listen, reverse=True)


def test_aktivitaeten_filter(daten):
    aktiv, ruhig = daten
    client = _client(Role.RH)
    url = reverse("membres:membre-reporting-aktivitaeten")

    def lade(**params):
        return client.get(url, params).json()

    assert {z["typ"] for z in lade(typ="bestellung,teilnahme")["ergebnisse"]} == {
        "bestellung",
        "teilnahme",
    }
    nur_ruhig = lade(membre=str(ruhig.pk))
    assert {z["typ"] for z in nur_ruhig["ergebnisse"]} == {"statuswechsel"}
    assert lade(typ="mitgliedschaft", aktivitaet_status="annulee")["count"] == 1
    assert lade(typ="mitgliedschaft", min_betrag="45")["count"] == 1
    assert lade(typ="mitgliedschaft", max_betrag="45")["count"] == 1
    assert lade(statut="inactif")["count"] == 2  # Historie von Rolf
    assert lade(page_size="3")["count"] > 3
    assert len(lade(page_size="3")["ergebnisse"]) == 3
    assert lade(typ="bestellung", bis="2000-01-01")["count"] == 0


def test_ungueltige_parameter(daten):
    client = _client(Role.RH)
    url = reverse("membres:membre-reporting-aktivitaeten")
    assert client.get(url, {"von": "gestern"}).status_code == 400
    assert client.get(url, {"min_betrag": "viel"}).status_code == 400
    assert (
        client.get(reverse("membres:membre-reporting"), {"historie_jahr": "x"}).status_code == 400
    )


@pytest.mark.parametrize(
    "name", ["membre-reporting", "membre-reporting-aktivitaeten", "membre-reporting-export"]
)
def test_zugriff_nur_ab_rh(daten, name):
    url = reverse(f"membres:{name}")
    assert APIClient().get(url).status_code in (401, 403)
    assert _client(Role.MEMBRE).get(url).status_code == 403
    assert _client(Role.RH).get(url).status_code == 200


def test_export_mitglieder_und_aktivitaeten(daten):
    client = _client(Role.RH)
    url = reverse("membres:membre-reporting-export")
    mappe = load_workbook(io.BytesIO(client.get(url, {"statut": "inactif"}).content))
    assert [z[1].value for z in mappe["Mitglieder"].iter_rows(min_row=2)] == ["Ruhig"]
    assert len(list(mappe["Historie"].iter_rows(min_row=2))) == 2
    mappe = load_workbook(
        io.BytesIO(client.get(url, {"ansicht": "aktivitaeten", "typ": "bestellung"}).content)
    )
    zeilen = list(mappe["Aktivitäten"].iter_rows(min_row=2))
    assert len(zeilen) == 1
    assert zeilen[0][1].value == "bestellung"


def test_summe_ohne_betrag_hat_typisiertes_null():
    """PostgreSQL kennt `SUM(NULL)` nicht eindeutig; die Betragsspalte muss ein CAST sein."""
    from apps.membres.reporting import _quelle

    for typ in ("projektmitarbeit", "statuswechsel"):
        sql = str(_quelle(typ).query).upper()
        assert "CAST(NULL AS" in sql.replace("  ", " ") or "NULL::" in sql, sql
