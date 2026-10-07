"""Tests — Schutz vor Formel-Injektion in Excel-/CSV-Exporten (Sicherheitsprüfung 2026-10-07)."""

import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.membres.utils_http import safe_cell, safe_csv, safe_zeile


def test_safe_cell_neutralisiert_formeln_und_laesst_normalen_text_unveraendert():
    assert safe_cell('=HYPERLINK("https://evil","x")').startswith("'=")
    assert safe_cell("Anna") == "Anna"
    assert safe_cell("+49 170 1234567") == "+49 170 1234567"
    assert safe_cell(12) == 12
    assert safe_cell(None) is None


def test_safe_zeile_wirkt_auf_jede_zelle():
    assert safe_zeile(["=1+1", "ok", 3]) == ["'=1+1", "ok", 3]


@pytest.mark.parametrize("wert", ["=A1", "+A1", "-A1", "@SUM(A1)", "\tA1", "\rA1"])
def test_safe_csv_neutralisiert_alle_formelanfaenge(wert):
    assert safe_csv(wert) == "'" + wert


def test_safe_csv_beloest_text_und_zahlen():
    assert safe_csv("Spende") == "Spende"
    assert safe_csv(5) == 5


def test_xlsx_mit_formel_im_text_wird_nicht_zur_formel():
    mappe = Workbook()
    blatt = mappe.active
    blatt.append(safe_zeile(['=HYPERLINK("https://evil","x")']))
    puffer = io.BytesIO()
    mappe.save(puffer)
    puffer.seek(0)
    zelle = load_workbook(puffer).active["A1"]
    assert zelle.data_type != "f"
