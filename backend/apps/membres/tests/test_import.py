"""
Tests import Excel des membres (RICEFW C-001/W-008) — en deux temps depuis le 2026-10-08 :
POST import/pruefen/ (analyse, rien n'est écrit) puis POST import/bestaetigen/ (import, écrasement
des doublons choisis, rapport Excel).
"""

import base64
import datetime
import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.membres.imports import analyser_membres
from apps.membres.models import Membre, StatutMembre
from apps.membres.tests.factories import MembreFactory

pytestmark = pytest.mark.django_db

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DEFAULT_HEADERS = [
    "prenom",
    "nom",
    "date_naissance",
    "sexe",
    "email",
    "telephone",
    "cin",
    "passeport",
    "adresse_de",
    "code_postal_de",
    "ville_de",
    "land_de",
    "ville_origine_tn",
    "gouvernorat_tn",
    "statut",
    "date_adhesion",
]


def _xlsx_file(rows, headers=None, filename="import.xlsx"):
    headers = headers if headers is not None else DEFAULT_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(filename, buf.getvalue(), content_type=XLSX_MIME)


def _ligne(
    prenom="Ahmed",
    nom="Zribi",
    date_naissance="15/03/1985",
    sexe="Homme",
    email="ahmed.zribi@example.de",
    telephone="+49 170 1111111",
    cin="10000001",
    passeport="",
    adresse_de="Musterstr. 1",
    code_postal_de="10115",
    ville_de="Berlin",
    land_de="BE",
    ville_origine_tn="Tunis",
    gouvernorat_tn="Tunis",
    statut="",
    date_adhesion="01/09/2020",
):
    return [
        prenom,
        nom,
        date_naissance,
        sexe,
        email,
        telephone,
        cin,
        passeport,
        adresse_de,
        code_postal_de,
        ville_de,
        land_de,
        ville_origine_tn,
        gouvernorat_tn,
        statut,
        date_adhesion,
    ]


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture
def rh_user():
    return User.objects.create_user(
        email="rh@example.de", password="Password123!", role=Role.RH, is_active=True
    )


@pytest.fixture
def membre_user():
    return User.objects.create_user(
        email="membre@example.de", password="Password123!", role=Role.MEMBRE, is_active=True
    )


def _auth(api_client, user):
    api_client.force_authenticate(user=user)
    return api_client


def _pruefen(api_client, fichier):
    return api_client.post(
        reverse("membres:membre-import-pruefen"), {"fichier": fichier}, format="multipart"
    )


def _bestaetigen(api_client, fichier, ueberschreiben=None):
    data = {"fichier": fichier}
    if ueberschreiben is not None:
        data["ueberschreiben"] = ueberschreiben
    return api_client.post(reverse("membres:membre-import-bestaetigen"), data, format="multipart")


def _zeilen(resp):
    return {z["ligne"]: z for z in resp.data["zeilen"]}


def _bericht(resp):
    inhalt = base64.b64decode(resp.data["bericht"]["inhalt_base64"])
    return openpyxl.load_workbook(io.BytesIO(inhalt)).worksheets[0]


# --- Prüfphase : Auswertung ohne zu schreiben ---


def test_pruefen_schreibt_nichts_und_listet_neue_zeilen(api_client, rh_user):
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001", nom="Zribi"),
            _ligne(email="b@example.de", cin="10000002", nom="Bchini"),
        ]
    )
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 200
    assert resp.data["total"] == 2
    assert resp.data["zaehler"] == {"neu": 2, "dublette": 0, "unveraendert": 0, "fehler": 0}
    assert [z["ligne"] for z in resp.data["zeilen"]] == [2, 3]
    assert resp.data["zeilen"][0]["anzeige"]["email"] == "a@example.de"
    assert Membre.objects.count() == 0


def test_pruefen_fehler_pro_zeile_mit_grund(api_client, rh_user):
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001"),
            _ligne(email="b@example.de", cin="10000002", nom=""),  # nom manquant
            _ligne(email="c@example.de", cin="10000003", land_de="Atlantide"),
            _ligne(email="pas-un-email", cin="10000004"),
        ]
    )
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    zeilen = _zeilen(resp)
    assert zeilen[2]["status"] == "neu"
    assert zeilen[3]["status"] == "fehler" and "nom" in zeilen[3]["grund"]
    assert zeilen[4]["status"] == "fehler" and "Bundesland" in zeilen[4]["grund"]
    assert zeilen[5]["status"] == "fehler" and "E-Mail" in zeilen[5]["grund"]
    assert resp.data["zaehler"]["fehler"] == 3


def test_pruefen_leere_zeilen_werden_ignoriert(api_client, rh_user):
    fichier = _xlsx_file([_ligne(email="a@example.de", cin="10000001"), [None] * 16])
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    assert resp.data["total"] == 1


def test_pruefen_doppelte_zeile_in_der_datei_ist_fehler(api_client, rh_user):
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001"),
            _ligne(email="a2@example.de", cin="10000001"),  # même CIN, email différent
        ]
    )
    zeilen = _zeilen(_pruefen(_auth(api_client, rh_user), fichier))
    assert zeilen[2]["status"] == "neu"
    assert zeilen[3]["status"] == "fehler"
    assert "Zeile 2" in zeilen[3]["grund"]


def test_pruefen_dublette_per_email_zeigt_abweichungen(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="99999999", telephone="+49 1", ville_de="Bonn")
    fichier = _xlsx_file(
        [_ligne(email="EX@example.de", cin="99999999", telephone="+49 2", ville_de="Bonn")]
    )
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    zeile = _zeilen(resp)[2]
    assert zeile["status"] == "dublette"
    assert zeile["ueberschreibbar"] is True
    assert zeile["existant"]["numero_membre"].startswith("CA-")
    felder = {a["champ"]: a for a in zeile["aenderungen"]}
    assert felder["telephone"]["alt"] == "+49 1" and felder["telephone"]["neu"] == "+49 2"
    assert "ville_de" not in felder  # gleicher Wert -> keine Änderung
    assert "email" not in felder  # nur Schreibweise verschieden


def test_pruefen_dublette_ohne_abweichung_ist_unveraendert(api_client, rh_user):
    MembreFactory(
        prenom="Ahmed",
        nom="Zribi",
        date_naissance=datetime.date(1985, 3, 15),
        email="ahmed.zribi@example.de",
        telephone="+49 170 1111111",
        cin="10000001",
        adresse_de="Musterstr. 1",
        code_postal_de="10115",
        ville_de="Berlin",
        land_de="BE",
        ville_origine_tn="Tunis",
        gouvernorat_tn="Tunis",
        date_adhesion=datetime.date(2020, 9, 1),
        sexe="homme",
    )
    resp = _pruefen(_auth(api_client, rh_user), _xlsx_file([_ligne()]))
    zeile = _zeilen(resp)[2]
    assert zeile["status"] == "unveraendert"
    assert zeile["ueberschreibbar"] is False
    assert resp.data["zaehler"]["unveraendert"] == 1


def test_pruefen_leere_zellen_gelten_nicht_als_abweichung(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="99999999", telephone="+49 5", statut="inactif")
    fichier = _xlsx_file(
        [_ligne(email="ex@example.de", cin="", telephone="", sexe="", statut="", passeport="")]
    )
    felder = {
        a["champ"] for a in _zeilen(_pruefen(_auth(api_client, rh_user), fichier))[2]["aenderungen"]
    }
    assert "telephone" not in felder
    assert "cin" not in felder
    assert "sexe" not in felder
    assert "statut" not in felder  # Standardwert ACTIF darf einen Bestand nie überschreiben
    assert "passeport" not in felder


def test_pruefen_statut_nur_wenn_spalte_gefuellt(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="99999999", statut="actif")
    fichier = _xlsx_file([_ligne(email="ex@example.de", cin="99999999", statut="inactif")])
    felder = {
        a["champ"]: a
        for a in _zeilen(_pruefen(_auth(api_client, rh_user), fichier))[2]["aenderungen"]
    }
    assert felder["statut"]["neu"] == "Inactif"


def test_pruefen_cin_wird_maskiert(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="12345678")
    fichier = _xlsx_file([_ligne(email="ex@example.de", cin="12345678")])
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    assert "12345678" not in str(resp.data)

    MembreFactory(email="ex2@example.de", cin="22222222")
    fichier = _xlsx_file([_ligne(email="ex2@example.de", cin="33333333")])
    cin = {a["champ"]: a for a in _zeilen(_pruefen(api_client, fichier))[2]["aenderungen"]}["cin"]
    assert cin["alt"].endswith("22") and cin["neu"].endswith("33")
    assert "22222222" not in cin["alt"] and "33333333" not in cin["neu"]


def test_pruefen_dublette_per_cin_mit_anderer_email(api_client, rh_user):
    MembreFactory(email="alt@example.de", cin="99999999")
    fichier = _xlsx_file([_ligne(email="neu@example.de", cin="99999999")])
    zeile = _zeilen(_pruefen(_auth(api_client, rh_user), fichier))[2]
    assert zeile["status"] == "dublette"
    assert "CIN" in zeile["grund"]
    assert "email" in {a["champ"] for a in zeile["aenderungen"]}  # Fiche ohne Konto


def test_pruefen_email_einer_fiche_mit_konto_bleibt_unveraendert(api_client, rh_user):
    user = User.objects.create_user(email="konto@example.de", password="Password123!")
    MembreFactory(user=user, email="konto@example.de", cin="99999999")
    fichier = _xlsx_file([_ligne(email="andere@example.de", cin="99999999")])
    zeile = _zeilen(_pruefen(_auth(api_client, rh_user), fichier))[2]
    assert zeile["status"] == "dublette"
    assert "email" not in {a["champ"] for a in zeile["aenderungen"]}


def test_pruefen_email_und_cin_verschiedener_mitglieder_ist_fehler(api_client, rh_user):
    MembreFactory(email="a@example.de", cin="11111111")
    MembreFactory(email="b@example.de", cin="22222222")
    fichier = _xlsx_file([_ligne(email="a@example.de", cin="22222222")])
    zeile = _zeilen(_pruefen(_auth(api_client, rh_user), fichier))[2]
    assert zeile["status"] == "fehler"
    assert "mehreren" in zeile["grund"]


# --- CIN und téléphone facultatifs (décision du 2026-10-07) ---


def test_pruefen_ohne_cin_und_telefon(api_client, rh_user):
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="", telephone=""),
            _ligne(email="b@example.de", cin=None, telephone=None),  # 2 sans CIN : pas un doublon
        ]
    )
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    assert resp.data["zaehler"]["neu"] == 2


def test_pruefen_spalten_cin_und_telefon_fehlen_in_der_datei(api_client, rh_user):
    idx = [DEFAULT_HEADERS.index("cin"), DEFAULT_HEADERS.index("telephone")]
    headers = [h for i, h in enumerate(DEFAULT_HEADERS) if i not in idx]
    ligne = [v for i, v in enumerate(_ligne(email="a@example.de")) if i not in idx]
    resp = _pruefen(_auth(api_client, rh_user), _xlsx_file([ligne], headers=headers))
    assert resp.status_code == 200
    assert resp.data["zaehler"]["neu"] == 1


def test_pruefen_membre_en_base_avec_cin_platzhalter_bloque_pas(api_client, rh_user):
    MembreFactory(email="existant@example.de", cin="00000000")
    MembreFactory(email="sans@example.de", cin=None)
    fichier = _xlsx_file([_ligne(email="nouveau@example.de", cin="")])
    resp = _pruefen(_auth(api_client, rh_user), fichier)
    assert resp.data["zaehler"]["neu"] == 1


# --- Bestätigungsphase ---


def test_bestaetigen_importiert_neue_zeilen(api_client, rh_user):
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001", statut=""),
            _ligne(email="b@example.de", cin="10000002", land_de="Bayern"),
        ]
    )
    resp = _bestaetigen(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 200
    assert resp.data["zaehler"]["importiert"] == 2
    membre = Membre.objects.get(email="a@example.de")
    assert membre.numero_membre.startswith("CA-2020-")  # année de date_adhesion
    assert membre.statut == StatutMembre.ACTIF
    assert Membre.objects.get(email="b@example.de").land_de == "BY"


def test_bestaetigen_cin_und_telefon_als_excel_zahlen(api_client, rh_user):
    fichier = _xlsx_file([_ligne(email="a@example.de", cin=12345678.0, telephone=4917011111)])
    _bestaetigen(_auth(api_client, rh_user), fichier)
    membre = Membre.objects.get(email="a@example.de")
    assert membre.cin == "12345678"
    assert membre.telephone == "4917011111"


def test_bestaetigen_leere_cin_wird_platzhalter(api_client, rh_user):
    fichier = _xlsx_file([_ligne(email="a@example.de", cin="", telephone="")])
    _bestaetigen(_auth(api_client, rh_user), fichier)
    membre = Membre.objects.get(email="a@example.de")
    assert membre.cin == "00000000"
    assert membre.telephone == ""


def test_bestaetigen_dublette_ohne_auswahl_bleibt_unveraendert(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="99999999", telephone="+49 1")
    fichier = _xlsx_file([_ligne(email="ex@example.de", cin="99999999", telephone="+49 2")])
    resp = _bestaetigen(_auth(api_client, rh_user), fichier)
    assert resp.data["zaehler"]["uebersprungen"] == 1
    assert Membre.objects.get(email="ex@example.de").telephone == "+49 1"
    assert "nicht zum Überschreiben" in _zeilen(resp)[2]["grund"]


def test_bestaetigen_ueberschreibt_nur_ausgewaehlte_dubletten_und_nur_gefuellte_zellen(
    api_client, rh_user
):
    a = MembreFactory(
        email="a@example.de", cin="10000001", telephone="+49 1", ville_de="Bonn", passeport="P1"
    )
    b = MembreFactory(email="b@example.de", cin="10000002", telephone="+49 1")
    nummer_a = a.numero_membre
    fichier = _xlsx_file(
        [
            _ligne(
                email="a@example.de",
                cin="10000001",
                telephone="+49 2",
                ville_de="Köln",
                passeport="",
                statut="",
            ),
            _ligne(email="b@example.de", cin="10000002", telephone="+49 3"),
        ]
    )
    resp = _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["2"])
    assert resp.data["zaehler"]["ueberschrieben"] == 1
    assert resp.data["zaehler"]["uebersprungen"] == 1

    a.refresh_from_db()
    b.refresh_from_db()
    assert a.telephone == "+49 2" and a.ville_de == "Köln"
    assert a.passeport == "P1"  # Zelle leer -> bleibt
    assert a.numero_membre == nummer_a
    assert b.telephone == "+49 1"  # nicht ausgewählt
    assert Membre.objects.count() == 2  # keine Dublette angelegt


def test_bestaetigen_email_einer_fiche_mit_konto_bleibt(api_client, rh_user):
    user = User.objects.create_user(email="konto@example.de", password="Password123!")
    membre = MembreFactory(user=user, email="konto@example.de", cin="99999999", telephone="+49 1")
    fichier = _xlsx_file([_ligne(email="andere@example.de", cin="99999999", telephone="+49 2")])
    _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["2"])
    membre.refresh_from_db()
    assert membre.email == "konto@example.de"
    assert membre.telephone == "+49 2"


def test_bestaetigen_statut_wird_nur_bei_gefuellter_spalte_ueberschrieben(api_client, rh_user):
    a = MembreFactory(email="a@example.de", cin="10000001", statut="inactif", telephone="+49 1")
    b = MembreFactory(email="b@example.de", cin="10000002", statut="actif", telephone="+49 1")
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001", telephone="+49 2", statut=""),
            _ligne(email="b@example.de", cin="10000002", telephone="+49 2", statut="inactif"),
        ]
    )
    _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["2", "3"])
    a.refresh_from_db()
    b.refresh_from_db()
    assert a.statut == StatutMembre.INACTIF  # Standardwert ACTIF hat ihn nicht überschrieben
    assert b.statut == StatutMembre.INACTIF


def test_bestaetigen_auswahl_einer_nicht_dublette_wird_ignoriert(api_client, rh_user):
    fichier = _xlsx_file([_ligne(email="a@example.de", cin="10000001")])
    resp = _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["2", "99"])
    assert resp.status_code == 200
    assert resp.data["zaehler"]["importiert"] == 1


def test_bestaetigen_ungueltige_auswahl_400(api_client, rh_user):
    fichier = _xlsx_file([_ligne()])
    resp = _bestaetigen(_auth(api_client, rh_user), fichier, ueberschreiben=["abc"])
    assert resp.status_code == 400
    assert resp.data["code"] == "ueberschreiben_ungueltig"
    assert Membre.objects.count() == 0


def test_bestaetigen_bericht_hat_status_und_grund_spalten(api_client, rh_user):
    MembreFactory(email="ex@example.de", cin="99999999", telephone="+49 1")
    fichier = _xlsx_file(
        [
            _ligne(email="a@example.de", cin="10000001"),
            _ligne(email="ex@example.de", cin="99999999", telephone="+49 2"),
            _ligne(email="x@example.de", cin="10000003", nom=""),
        ],
        filename="meine liste.xlsx",
    )
    resp = _bestaetigen(_auth(api_client, rh_user), fichier)
    assert resp.data["bericht"]["dateiname"] == "meine_liste_bericht.xlsx"

    blatt = _bericht(resp)
    kopf = [c.value for c in blatt[1]]
    assert kopf[: len(DEFAULT_HEADERS)] == DEFAULT_HEADERS  # Originalspalten bleiben
    assert kopf[-2:] == ["Status", "Grund"]
    assert blatt.cell(row=2, column=len(kopf) - 1).value == "Importiert"
    assert "angelegt" in blatt.cell(row=2, column=len(kopf)).value
    assert blatt.cell(row=3, column=len(kopf) - 1).value == "Übersprungen"
    assert blatt.cell(row=4, column=len(kopf) - 1).value == "Fehler"
    assert blatt.cell(row=2, column=1).value == "Ahmed"  # Originaldaten unverändert


def test_bestaetigen_bericht_neutralisiert_formeln(api_client, rh_user):
    fichier = _xlsx_file([_ligne(email="=1+1", cin="10000001")])
    resp = _bestaetigen(_auth(api_client, rh_user), fichier)
    blatt = _bericht(resp)
    grund = blatt.cell(row=2, column=len(DEFAULT_HEADERS) + 2)
    assert grund.data_type != "f"


# --- Validation fichier / permissions ---


@pytest.mark.parametrize("aktion", ["pruefen", "bestaetigen"])
def test_spalte_fehlt_400(api_client, rh_user, aktion):
    headers_incomplets = [h for h in DEFAULT_HEADERS if h != "email"]
    fichier = _xlsx_file([], headers=headers_incomplets)
    funktion = _pruefen if aktion == "pruefen" else _bestaetigen
    resp = funktion(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 400
    assert resp.data["code"] == "schema_invalide"


@pytest.mark.parametrize("aktion", ["pruefen", "bestaetigen"])
def test_fichier_non_excel_rejete_400(api_client, rh_user, aktion):
    fichier = SimpleUploadedFile(
        "import.xlsx", b"ceci n'est pas un fichier excel", content_type=XLSX_MIME
    )
    funktion = _pruefen if aktion == "pruefen" else _bestaetigen
    resp = funktion(_auth(api_client, rh_user), fichier)
    assert resp.status_code == 400
    assert resp.data["code"] == "type_fichier_invalide"


def test_sans_fichier_400(api_client, rh_user):
    _auth(api_client, rh_user)
    resp = api_client.post(reverse("membres:membre-import-pruefen"), {}, format="multipart")
    assert resp.status_code == 400
    assert resp.data["code"] == "fichier_requis"


@pytest.mark.parametrize("aktion", ["pruefen", "bestaetigen"])
def test_comme_membre_refuse_403(api_client, membre_user, aktion):
    funktion = _pruefen if aktion == "pruefen" else _bestaetigen
    resp = funktion(_auth(api_client, membre_user), _xlsx_file([_ligne()]))
    assert resp.status_code == 403


@pytest.mark.parametrize("aktion", ["pruefen", "bestaetigen"])
def test_non_authentifie_401(api_client, aktion):
    funktion = _pruefen if aktion == "pruefen" else _bestaetigen
    resp = funktion(api_client, _xlsx_file([_ligne()]))
    assert resp.status_code == 401


def test_analyser_membres_direkt_ohne_schreiben():
    analyse = analyser_membres(_xlsx_file([_ligne()]), "x.xlsx")
    assert analyse.zaehler()["neu"] == 1
    assert Membre.objects.count() == 0


# --- Téléchargement du template (RICEFW W-008/F-019) ---


def test_telechargement_template_comme_rh_ok(api_client, rh_user):
    _auth(api_client, rh_user)
    resp = api_client.get(reverse("membres:membre-import-template"))

    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert "template_import_membres.xlsx" in resp["Content-Disposition"]

    classeur = openpyxl.load_workbook(io.BytesIO(resp.content))
    en_tetes = next(classeur.active.iter_rows(values_only=True))
    assert set(DEFAULT_HEADERS) <= set(en_tetes)


def test_telechargement_template_comme_membre_refuse_403(api_client, membre_user):
    _auth(api_client, membre_user)
    resp = api_client.get(reverse("membres:membre-import-template"))
    assert resp.status_code == 403


def test_telechargement_template_non_authentifie_401(api_client):
    resp = api_client.get(reverse("membres:membre-import-template"))
    assert resp.status_code == 401
