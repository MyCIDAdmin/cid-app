"""Reporting über Business Partner & Lieferanten (Nutzerwunsch 2026-10-07) : Liste aller Partner
mit verknüpften Elementen (Projekte, Veranstaltungen, Shop-Artikel), Umsatz und Details sowie
Excel-Export. Lesen ab Rolle RH wie die übrige Partnerverwaltung (PartnerPermission).

Umsatz je Partner = Einnahmen (`PartnerEinnahme`, z. B. Sponsoring) + genehmigte Ausgaben
(`finances.Depense` mit `partner`) ; Saldo = Einnahmen - Ausgaben. Der Zeitraum (`von`/`bis`)
begrenzt beide Summen."""

from datetime import date
from decimal import Decimal

from django.db.models import Count, Q, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.finances.models import Depense, StatutDepense
from apps.membres.utils_http import xlsx_response

from .models import Angebot, AngebotStatus, PartnerEinnahme, VerknuepfungRolle
from .permissions import PartnerPermission
from .views import filtere_partner, partner_basis_queryset

NULL = Decimal("0.00")
ZIEL_TYPEN = ("projet", "evenement", "produit")
ORDERINGS = {
    "nom": lambda z: z["nom"].lower(),
    "umsatz": lambda z: -z["umsatz"],
    "einnahmen": lambda z: -z["einnahmen_summe"],
    "ausgaben": lambda z: -z["ausgaben_summe"],
    "saldo": lambda z: -z["saldo"],
    "verknuepfungen": lambda z: -len(z["verknuepfungen"]),
}


def _datum(wert, name):
    if not wert:
        return None
    try:
        return date.fromisoformat(wert)
    except ValueError as exc:
        raise ValidationError({name: "Datum im Format JJJJ-MM-TT erwartet."}) from exc


def _dezimal(wert, name):
    if wert in (None, ""):
        return None
    try:
        return Decimal(str(wert))
    except Exception as exc:  # decimal.InvalidOperation
        raise ValidationError({name: "Zahl erwartet."}) from exc


def _summen(qs, feld_datum, feld_betrag, partner_feld, von, bis):
    if von:
        qs = qs.filter(**{f"{feld_datum}__gte": von})
    if bis:
        qs = qs.filter(**{f"{feld_datum}__lte": bis})
    zeilen = (
        qs.filter(**{f"{partner_feld}__isnull": False})
        .values(partner_feld)
        .annotate(summe=Sum(feld_betrag), anzahl=Count("id"))
    )
    return {z[partner_feld]: (z["summe"] or NULL, z["anzahl"]) for z in zeilen}


def baue_reporting(params):
    """Liefert (zeilen, summen) für die Query-Parameter `params`."""
    von, bis = _datum(params.get("von"), "von"), _datum(params.get("bis"), "bis")
    min_umsatz = _dezimal(params.get("min_umsatz"), "min_umsatz")
    max_umsatz = _dezimal(params.get("max_umsatz"), "max_umsatz")

    partner = filtere_partner(partner_basis_queryset(), params)
    if params.get("rolle") in VerknuepfungRolle.values:
        partner = partner.filter(verknuepfungen__rolle=params["rolle"]).distinct()
    if params.get("ziel_typ") in ZIEL_TYPEN:
        partner = partner.filter(
            **{f"verknuepfungen__{params['ziel_typ']}__isnull": False}
        ).distinct()
    partner = list(
        partner.prefetch_related(
            "verknuepfungen__projet", "verknuepfungen__evenement", "verknuepfungen__produit"
        )
    )

    ausgaben = _summen(
        Depense.objects.filter(statut=StatutDepense.APPROUVEE),
        "date_depense",
        "montant",
        "partner",
        von,
        bis,
    )
    einnahmen = _summen(PartnerEinnahme.objects.all(), "datum", "betrag", "partner", von, bis)
    angebote = {
        a["partner"]: a
        for a in Angebot.objects.values("partner").annotate(
            anzahl=Count("id"),
            zuschlag_anzahl=Count("id", filter=Q(status=AngebotStatus.ZUSCHLAG)),
            zuschlag_summe=Sum("betrag", filter=Q(status=AngebotStatus.ZUSCHLAG)),
        )
    }

    zeilen = []
    for p in partner:
        aus_summe, aus_anzahl = ausgaben.get(p.pk, (NULL, 0))
        ein_summe, ein_anzahl = einnahmen.get(p.pk, (NULL, 0))
        ang = angebote.get(p.pk, {})
        verknuepfungen = [
            {
                "id": str(v.pk),
                "ziel_typ": v.ziel_typ,
                "ziel_id": str(v.ziel.pk) if v.ziel else None,
                "ziel_label": v.ziel_label,
                "rolle": v.rolle,
            }
            for v in p.verknuepfungen.all()
        ]
        kontakte = list(p.kontakte.all())
        zeile = {
            "id": str(p.pk),
            "nom": p.nom,
            "typ": p.typ,
            "statut": p.statut,
            "bevorzugt": p.bevorzugt,
            "auf_startseite": p.auf_startseite,
            "kategorien_namen": [k.nom for k in p.kategorien.all()],
            "website": p.website,
            "email": p.email,
            "telefon": p.telefon,
            "ville": p.ville,
            "pays": p.pays,
            "hauptkontakt_name": kontakte[0].name if kontakte else "",
            "bewertung_schnitt": (
                round(p.bewertung_schnitt, 2) if p.bewertung_schnitt is not None else None
            ),
            "bewertung_anzahl": p.bewertung_anzahl,
            "ausgaben_summe": aus_summe,
            "ausgaben_anzahl": aus_anzahl,
            "einnahmen_summe": ein_summe,
            "einnahmen_anzahl": ein_anzahl,
            "umsatz": aus_summe + ein_summe,
            "saldo": ein_summe - aus_summe,
            "angebote_anzahl": ang.get("anzahl", 0),
            "zuschlaege_anzahl": ang.get("zuschlag_anzahl", 0),
            "zuschlaege_summe": ang.get("zuschlag_summe") or NULL,
            "verknuepfungen": verknuepfungen,
        }
        if min_umsatz is not None and zeile["umsatz"] < min_umsatz:
            continue
        if max_umsatz is not None and zeile["umsatz"] > max_umsatz:
            continue
        zeilen.append(zeile)

    zeilen.sort(key=ORDERINGS.get(params.get("sortierung"), ORDERINGS["nom"]))
    summen = {
        "partner": len(zeilen),
        "ausgaben": sum((z["ausgaben_summe"] for z in zeilen), NULL),
        "einnahmen": sum((z["einnahmen_summe"] for z in zeilen), NULL),
        "umsatz": sum((z["umsatz"] for z in zeilen), NULL),
        "saldo": sum((z["saldo"] for z in zeilen), NULL),
        "verknuepfungen": sum(len(z["verknuepfungen"]) for z in zeilen),
    }
    return zeilen, summen


def _json(zeile):
    return {k: (f"{v:.2f}" if isinstance(v, Decimal) else v) for k, v in zeile.items()}


class PartnerReportingView(APIView):
    """GET /partenaires/reporting/ — gefilterte Partnerliste mit Verknüpfungen und Umsatz.
    Filter: typ, kategorie, statut/alle, bevorzugt, auf_startseite, q, rolle, ziel_typ, projet/
    evenement/produit (Ziel-ID), min_note, von/bis (Zeitraum der Summen), min_umsatz/max_umsatz,
    sortierung (nom, umsatz, einnahmen, ausgaben, saldo, verknuepfungen)."""

    permission_classes = [PartnerPermission]

    def get(self, request):
        zeilen, summen = baue_reporting(request.query_params)
        return Response(
            {
                "summen": {
                    k: (f"{v:.2f}" if isinstance(v, Decimal) else v) for k, v in summen.items()
                },
                "ergebnisse": [_json(z) for z in zeilen],
            }
        )


KOPF = [
    "Name",
    "Typ",
    "Status",
    "Kategorien",
    "Hauptkontakt",
    "E-Mail",
    "Telefon",
    "Ort",
    "Website",
    "Bewertung",
    "Einnahmen",
    "Ausgaben",
    "Umsatz",
    "Saldo",
    "Angebote",
    "Zuschläge",
    "Verknüpfungen",
]


class PartnerReportingExportView(APIView):
    """GET /partenaires/reporting/export/ — gleiche Filter wie das Reporting, als Excel."""

    permission_classes = [PartnerPermission]

    def get(self, request):
        zeilen, _summen_unused = baue_reporting(request.query_params)
        wb = Workbook()
        blatt = wb.active
        blatt.title = "Partner"
        blatt.append(KOPF)
        for z in zeilen:
            blatt.append(
                [
                    z["nom"],
                    z["typ"],
                    z["statut"],
                    ", ".join(z["kategorien_namen"]),
                    z["hauptkontakt_name"],
                    z["email"],
                    z["telefon"],
                    z["ville"],
                    z["website"],
                    z["bewertung_schnitt"],
                    float(z["einnahmen_summe"]),
                    float(z["ausgaben_summe"]),
                    float(z["umsatz"]),
                    float(z["saldo"]),
                    z["angebote_anzahl"],
                    z["zuschlaege_anzahl"],
                    len(z["verknuepfungen"]),
                ]
            )
        verknuepfungen = wb.create_sheet("Verknüpfungen")
        verknuepfungen.append(["Partner", "Art", "Element", "Rolle"])
        for z in zeilen:
            for v in z["verknuepfungen"]:
                verknuepfungen.append([z["nom"], v["ziel_typ"], v["ziel_label"], v["rolle"]])
        for blatt_ in (blatt, verknuepfungen):
            for zelle in blatt_[1]:
                zelle.font = Font(bold=True, color="FFFFFF")
                zelle.fill = PatternFill("solid", fgColor="C8102E")
            blatt_.freeze_panes = "A2"
            for spalte in blatt_.columns:
                blatt_.column_dimensions[spalte[0].column_letter].width = 22
        return xlsx_response(wb, f"reporting_partner_{timezone.localtime():%Y%m%d_%H%M}.xlsx")
