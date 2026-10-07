"""Tests — Partner-Reporting, Partner-Einnahmen und Startseiten-Banner (2026-10-07)."""

import datetime
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.finances.models import CategorieDepense, Depense, StatutDepense
from apps.partenaires.models import (
    Partner,
    PartnerEinnahme,
    PartnerKategorie,
    PartnerStatus,
    PartnerVerknuepfung,
    VerknuepfungRolle,
)
from apps.projets.tests.factories import ProjetFactory

pytestmark = pytest.mark.django_db


def _client(role):
    email = f"{role}@example.com"
    user = User.objects.filter(email=email).first() or User.objects.create_user(
        email=email, password="Password123!", role=role, is_active=True
    )
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def _png():
    buf = io.BytesIO()
    Image.new("RGB", (20, 20), "blue").save(buf, format="PNG")
    return SimpleUploadedFile("l.png", buf.getvalue(), content_type="image/png")


def _ausgabe(partner, montant, datum=datetime.date(2026, 4, 2), statut=StatutDepense.APPROUVEE):
    return Depense.objects.create(
        date_depense=datum,
        montant=Decimal(montant),
        categorie=CategorieDepense.objects.filter(actif=True).first(),
        fournisseur=partner.nom,
        partner=partner,
        statut=statut,
    )


def _einnahme(partner, betrag, datum=datetime.date(2026, 5, 1), **kw):
    return PartnerEinnahme.objects.create(
        partner=partner, datum=datum, betrag=Decimal(betrag), **kw
    )


def _zeile(antwort, nom):
    return next(z for z in antwort.json()["ergebnisse"] if z["nom"] == nom)


@pytest.fixture
def daten():
    sponsor = Partner.objects.create(nom="Sponsor AG", typ="partner")
    lieferant = Partner.objects.create(nom="Druck GmbH", typ="lieferant")
    projet = ProjetFactory()
    PartnerVerknuepfung.objects.create(
        partner=sponsor, projet=projet, rolle=VerknuepfungRolle.SPONSOR
    )
    _einnahme(sponsor, "500.00")
    _ausgabe(lieferant, "120.00")
    _ausgabe(lieferant, "80.00")
    _ausgabe(lieferant, "999.00", statut=StatutDepense.EN_ATTENTE)  # zählt nicht
    return sponsor, lieferant, projet


def test_reporting_summen_und_verknuepfungen(daten):
    antwort = _client(Role.RH).get(reverse("partenaires:reporting"))
    assert antwort.status_code == 200
    sponsor = _zeile(antwort, "Sponsor AG")
    assert sponsor["einnahmen_summe"] == "500.00"
    assert sponsor["ausgaben_summe"] == "0.00"
    assert sponsor["saldo"] == "500.00"
    assert sponsor["verknuepfungen"][0]["ziel_typ"] == "projet"
    lieferant = _zeile(antwort, "Druck GmbH")
    assert lieferant["ausgaben_summe"] == "200.00"
    assert lieferant["ausgaben_anzahl"] == 2
    assert lieferant["umsatz"] == "200.00"
    assert lieferant["saldo"] == "-200.00"
    summen = antwort.json()["summen"]
    assert summen["partner"] == 2
    assert summen["umsatz"] == "700.00"
    assert summen["einnahmen"] == "500.00"


def test_reporting_zeitraum_begrenzt_summen(daten):
    antwort = _client(Role.RH).get(
        reverse("partenaires:reporting"), {"von": "2026-05-01", "bis": "2026-12-31"}
    )
    assert _zeile(antwort, "Druck GmbH")["ausgaben_summe"] == "0.00"
    assert _zeile(antwort, "Sponsor AG")["einnahmen_summe"] == "500.00"


def test_reporting_filter(daten):
    sponsor, lieferant, projet = daten
    client = _client(Role.RH)
    url = reverse("partenaires:reporting")

    def namen(**params):
        return [z["nom"] for z in client.get(url, params).json()["ergebnisse"]]

    assert namen(typ="lieferant") == ["Druck GmbH"]
    assert namen(rolle="sponsor") == ["Sponsor AG"]
    assert namen(ziel_typ="projet") == ["Sponsor AG"]
    assert namen(projet=str(projet.pk)) == ["Sponsor AG"]
    assert namen(min_umsatz="300") == ["Sponsor AG"]
    assert namen(max_umsatz="300") == ["Druck GmbH"]
    assert namen(q="druck") == ["Druck GmbH"]
    kategorie = PartnerKategorie.objects.create(nom="Druck")
    lieferant.kategorien.add(kategorie)
    assert namen(kategorie=str(kategorie.pk)) == ["Druck GmbH"]
    assert namen(sortierung="umsatz") == ["Sponsor AG", "Druck GmbH"]


def test_reporting_archivierte_nur_mit_filter(daten):
    sponsor, _, _ = daten
    sponsor.statut = PartnerStatus.ARCHIVIERT
    sponsor.save()
    client = _client(Role.RH)
    url = reverse("partenaires:reporting")
    assert [z["nom"] for z in client.get(url).json()["ergebnisse"]] == ["Druck GmbH"]
    assert len(client.get(url, {"alle": "1"}).json()["ergebnisse"]) == 2


def test_reporting_ungueltige_parameter(daten):
    client = _client(Role.RH)
    url = reverse("partenaires:reporting")
    assert client.get(url, {"von": "gestern"}).status_code == 400
    assert client.get(url, {"min_umsatz": "viel"}).status_code == 400


def test_reporting_zugriff(daten):
    url = reverse("partenaires:reporting")
    assert APIClient().get(url).status_code in (401, 403)
    assert _client(Role.MEMBRE).get(url).status_code == 403
    assert _client(Role.RH).get(url).status_code == 200


def test_reporting_export_xlsx(daten):
    antwort = _client(Role.RH).get(reverse("partenaires:reporting-export"), {"typ": "lieferant"})
    assert antwort.status_code == 200
    blatt = load_workbook(io.BytesIO(antwort.content))["Partner"]
    assert [z[0].value for z in blatt.iter_rows(min_row=2)] == ["Druck GmbH"]


def test_einnahmen_crud_rechte(daten):
    sponsor, _, _ = daten
    url = reverse("partenaires:einnahme-list")
    daten_neu = {"partner": str(sponsor.pk), "datum": "2026-06-01", "betrag": "250.00"}
    assert _client(Role.RH).post(url, daten_neu, format="json").status_code == 403
    bureau = _client(Role.BUREAU_ADMIN)
    antwort = bureau.post(url, daten_neu, format="json")
    assert antwort.status_code == 201
    liste = _client(Role.RH).get(url, {"partner": str(sponsor.pk)}).json()
    assert len(liste) == 2
    detail = reverse("partenaires:einnahme-detail", args=[antwort.json()["id"]])
    assert bureau.patch(detail, {"betrag": "300.00"}, format="json").status_code == 200
    assert bureau.delete(detail).status_code == 204


def test_einnahme_betrag_muss_positiv_sein(daten):
    sponsor, _, _ = daten
    antwort = _client(Role.BUREAU_ADMIN).post(
        reverse("partenaires:einnahme-list"),
        {"partner": str(sponsor.pk), "datum": "2026-06-01", "betrag": "0"},
        format="json",
    )
    assert antwort.status_code == 400


def test_banner_ist_oeffentlich_und_filtert(tmp_path):
    url = reverse("partenaires:banner")
    sichtbar = Partner.objects.create(nom="A Sponsor", auf_startseite=True, website="https://a.de")
    sichtbar.logo.save("a.png", _png())
    ohne_logo = Partner.objects.create(nom="Ohne Logo", auf_startseite=True)
    aus = Partner.objects.create(nom="Aus", auf_startseite=False)
    aus.logo.save("b.png", _png())
    archiviert = Partner.objects.create(
        nom="Archiv", auf_startseite=True, statut=PartnerStatus.ARCHIVIERT
    )
    archiviert.logo.save("c.png", _png())
    inaktiv = Partner.objects.create(
        nom="Inaktiv", auf_startseite=True, statut=PartnerStatus.INAKTIV
    )
    inaktiv.logo.save("d.png", _png())
    antwort = APIClient().get(url)  # ohne Anmeldung
    assert antwort.status_code == 200
    assert [p["nom"] for p in antwort.json()] == ["A Sponsor"]
    eintrag = antwort.json()[0]
    assert eintrag["website"] == "https://a.de"
    assert set(eintrag) == {"id", "nom", "logo_url", "website"}
    assert ohne_logo.pk and aus.pk


def test_partner_schalter_auf_startseite_setzen():
    partner = Partner.objects.create(nom="X")
    bureau = _client(Role.BUREAU_ADMIN)
    antwort = bureau.patch(
        reverse("partenaires:partner-detail", args=[partner.pk]),
        {"auf_startseite": True},
        format="json",
    )
    assert antwort.status_code == 200
    assert antwort.json()["auf_startseite"] is True
