"""Einlass (2026-10-23, T2) : Prüfen, Teil-Einlass, Rücknahme, Rechte, Helfer."""

import datetime
from unittest import mock

import pytest
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.eintritt import einlass, tasks
from apps.eintritt.models import Einlass, EinlassHelfer, Eintrittsticket
from apps.evenements.tests.factories import EvenementFactory, InscriptionFactory
from apps.membres.models import StatutMembre
from apps.membres.tests.factories import MembreFactory
from apps.notifications.models import Notification
from apps.regionen.models import (
    Region,
    Teilnahme,
    Veranstaltung,
    VeranstaltungStatus,
    Verantwortlicher,
)

pytestmark = pytest.mark.django_db
API = "/api/v1/eintritt/einlass/"


@pytest.fixture(autouse=True)
def _kein_versand():
    with mock.patch.object(tasks.ticket_senden, "delay"):
        yield


def _user(email, role=Role.MEMBRE, prenom="Sami", nom="Ben Ali"):
    user = User.objects.create_user(email=email, password="Password123!", role=role, is_active=True)
    membre = MembreFactory(user=user, statut=StatutMembre.ACTIF, prenom=prenom, nom=nom)
    return user, membre


def _client(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def _buero():
    return _user("buero@example.de", Role.BUREAU_ADMIN, "Bea", "Büro")[0]


def _ticket(places=2, adults=0, kids=1, tage=0, email="gast@example.de", **ev):
    _, membre = _user(email)
    evenement = EvenementFactory(
        ticket_aktiv=True,
        date_evenement=timezone.localdate() + datetime.timedelta(days=tage),
        **ev,
    )
    ins = InscriptionFactory(
        evenement=evenement,
        membre=membre,
        places=places,
        nombre_accompagnants_adultes=adults,
        nombre_accompagnants_enfants=kids,
    )
    return evenement, Eintrittsticket.objects.get(inscription=ins)


def _url(ziel, aktion=""):
    typ = "veranstaltung" if hasattr(ziel, "titre") else "treffen"
    return f"{API}{typ}/{ziel.pk}/{aktion}"


class TestPruefen:
    def test_gueltiger_code_zeigt_stand(self):
        ev, t = _ticket()
        r = _client(_buero()).post(_url(ev, "pruefen/"), {"code": f"CID-T:{t.code}"}, format="json")
        assert r.status_code == 200
        assert r.data["ergebnis"] == "gueltig"
        assert r.data["ticket"]["offen_erwachsene"] == 2
        assert r.data["ticket"]["offen_kinder"] == 1
        assert r.data["ticket"]["name"] == "Sami Ben Ali"

    def test_code_ohne_praefix_wird_akzeptiert(self):
        ev, t = _ticket()
        r = _client(_buero()).post(_url(ev, "pruefen/"), {"code": f" {t.code} "}, format="json")
        assert r.data["ergebnis"] == "gueltig"

    def test_unbekannter_code(self):
        ev, _ = _ticket()
        r = _client(_buero()).post(
            _url(ev, "pruefen/"), {"code": "CID-T:gibtsnicht"}, format="json"
        )
        assert r.data["ergebnis"] == "unbekannt"

    def test_erneuertes_ticket_alter_code_unbekannt(self):
        from apps.eintritt import services

        ev, t = _ticket()
        alt = t.code
        services.erneuern(t)
        r = _client(_buero()).post(_url(ev, "pruefen/"), {"code": alt}, format="json")
        assert r.data["ergebnis"] == "unbekannt"

    def test_ticket_einer_anderen_veranstaltung(self):
        ev, _ = _ticket()
        _, anderes = _ticket(email="x@example.de")
        r = _client(_buero()).post(_url(ev, "pruefen/"), {"code": anderes.code}, format="json")
        assert r.data["ergebnis"] == "anderes_ereignis"
        assert "ticket" not in r.data

    def test_ungueltiges_ticket(self):
        ev, t = _ticket()
        t.gueltig = False
        t.save()
        r = _client(_buero()).post(_url(ev, "pruefen/"), {"code": t.code}, format="json")
        assert r.data["ergebnis"] == "ungueltig"

    def test_ohne_code_400(self):
        ev, _ = _ticket()
        r = _client(_buero()).post(_url(ev, "pruefen/"), {}, format="json")
        assert r.status_code == 400


class TestEinlassen:
    def _buchen(self, client, ev, t, erw, kin):
        return client.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": erw, "kinder": kin},
            format="json",
        )

    def test_teil_einlass_mit_demselben_code(self):
        ev, t = _ticket(places=2, adults=0, kids=0)
        c = _client(_buero())
        r = self._buchen(c, ev, t, 1, 0)
        assert r.status_code == 200
        assert r.data["stand"]["status"] == "teilweise"
        assert r.data["stand"]["offen_erwachsene"] == 1
        pruef = c.post(_url(ev, "pruefen/"), {"code": t.code}, format="json")
        assert pruef.data["ergebnis"] == "teilweise"
        r2 = self._buchen(c, ev, t, 1, 0)
        assert r2.data["stand"]["status"] == "vollstaendig"
        pruef = c.post(_url(ev, "pruefen/"), {"code": t.code}, format="json")
        assert pruef.data["ergebnis"] == "vollstaendig"

    def test_nie_mehr_als_gebucht(self):
        ev, t = _ticket(places=2, adults=0, kids=1)
        c = _client(_buero())
        assert self._buchen(c, ev, t, 3, 0).status_code == 400
        assert self._buchen(c, ev, t, 0, 2).status_code == 400
        assert self._buchen(c, ev, t, 2, 1).status_code == 200
        assert self._buchen(c, ev, t, 1, 0).status_code == 400
        assert Einlass.objects.count() == 1

    def test_null_personen_und_schlechte_werte(self):
        ev, t = _ticket()
        c = _client(_buero())
        assert self._buchen(c, ev, t, 0, 0).status_code == 400
        assert self._buchen(c, ev, t, -1, 0).status_code == 400
        assert self._buchen(c, ev, t, "2", 0).status_code == 400
        assert self._buchen(c, ev, t, True, 0).status_code == 400

    def test_ticket_einer_anderen_veranstaltung_wird_abgelehnt(self):
        ev, _ = _ticket()
        _, fremd = _ticket(email="x@example.de")
        r = self._buchen(_client(_buero()), ev, fremd, 1, 0)
        assert r.status_code == 400
        assert not Einlass.objects.exists()

    def test_ungueltiges_ticket_blockiert_nur_den_scan_nicht_den_eintrag_von_hand(self):
        ev, t = _ticket()
        t.gueltig = False
        t.save()
        c = _client(_buero())
        assert c.post(_url(ev, "pruefen/"), {"code": t.code}, format="json").data["ergebnis"] == (
            "ungueltig"
        )
        assert self._buchen(c, ev, t, 1, 0).status_code == 200

    def test_person_ohne_bestaetigte_anmeldung_wird_abgelehnt(self):
        from apps.evenements.models import StatutInscription

        ev, t = _ticket()
        ins = t.inscription
        ins.statut = StatutInscription.LISTE_ATTENTE
        ins.save()
        assert self._buchen(_client(_buero()), ev, t, 1, 0).status_code == 400

    def test_ungueltige_personen_id(self):
        ev, _ = _ticket()
        r = _client(_buero()).post(
            _url(ev, "einlassen/"), {"person": "kein-uuid", "erwachsene": 1}, format="json"
        )
        assert r.status_code == 400

    def test_treffen_setzt_anwesend(self):
        v = Veranstaltung.objects.create(
            region=Region.objects.get(slug="nord"),
            titel="Public Viewing",
            datum=timezone.localdate(),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
            ticket_aktiv=True,
        )
        user, _ = _user("tn@example.de")
        tn = Teilnahme.objects.create(veranstaltung=v, user=user)
        t = Eintrittsticket.objects.get(teilnahme=tn)
        r = _client(_buero()).post(
            _url(v, "einlassen/"),
            {"person": str(t.teilnahme_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        assert r.status_code == 200
        tn.refresh_from_db()
        assert tn.anwesend is True


class TestRuecknahme:
    def test_rueckgaengig_gibt_plaetze_wieder_frei(self):
        ev, t = _ticket(places=1, adults=0, kids=0)
        c = _client(_buero())
        r = c.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        eintrag = r.data["einlass"]
        r2 = c.post(f"{API}ruecknahme/{eintrag}/")
        assert r2.status_code == 200
        assert r2.data["stand"]["status"] == "offen"
        assert c.post(f"{API}ruecknahme/{eintrag}/").status_code == 400

    def test_nach_der_frist_nicht_mehr(self):
        ev, t = _ticket(places=1, adults=0, kids=0)
        c = _client(_buero())
        r = c.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        Einlass.objects.update(zeit=timezone.now() - einlass.RUECKNAHME_FRIST * 2)
        assert c.post(f"{API}ruecknahme/{r.data['einlass']}/").status_code == 400

    def test_fremder_ohne_recht_darf_nicht(self):
        ev, t = _ticket(places=1, adults=0, kids=0)
        r = _client(_buero()).post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        zufaellig, _ = _user("zufall@example.de")
        assert _client(zufaellig).post(f"{API}ruecknahme/{r.data['einlass']}/").status_code == 403

    def test_treffen_anwesend_wird_zurueckgesetzt(self):
        v = Veranstaltung.objects.create(
            region=Region.objects.get(slug="nord"),
            titel="Public Viewing",
            datum=timezone.localdate(),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
            ticket_aktiv=True,
        )
        user, _ = _user("tn@example.de")
        tn = Teilnahme.objects.create(veranstaltung=v, user=user)
        t = Eintrittsticket.objects.get(teilnahme=tn)
        c = _client(_buero())
        r = c.post(
            _url(v, "einlassen/"),
            {"person": str(t.teilnahme_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        c.post(f"{API}ruecknahme/{r.data['einlass']}/")
        tn.refresh_from_db()
        assert tn.anwesend is False


class TestRechte:
    def test_normales_mitglied_darf_nicht(self):
        ev, t = _ticket()
        user, _ = _user("normal@example.de")
        c = _client(user)
        assert c.get(_url(ev)).status_code == 403
        assert c.post(_url(ev, "pruefen/"), {"code": t.code}, format="json").status_code == 403
        assert c.post(_url(ev, "einlassen/"), {}, format="json").status_code == 403

    def test_anonym_401(self):
        ev, _ = _ticket()
        assert APIClient().get(_url(ev)).status_code == 401

    def test_helfer_nur_am_veranstaltungstag(self):
        helfer, _ = _user("helfer@example.de")
        heute_ev, t = _ticket()
        spaeter_ev, _t2 = _ticket(tage=3, email="g2@example.de")
        EinlassHelfer.objects.create(evenement=heute_ev, user=helfer)
        EinlassHelfer.objects.create(evenement=spaeter_ev, user=helfer)
        c = _client(helfer)
        assert c.get(_url(heute_ev)).status_code == 200
        r = c.get(_url(spaeter_ev))
        assert r.status_code == 403
        assert (
            c.post(_url(heute_ev, "pruefen/"), {"code": t.code}, format="json").status_code == 200
        )

    def test_helfer_gilt_nur_fuer_seine_veranstaltung(self):
        helfer, _ = _user("helfer@example.de")
        ev, _ = _ticket()
        anderes, _ = _ticket(email="g2@example.de")
        EinlassHelfer.objects.create(evenement=ev, user=helfer)
        assert _client(helfer).get(_url(anderes)).status_code == 403

    def test_helfer_darf_keine_helfer_verwalten(self):
        helfer, _ = _user("helfer@example.de")
        ev, _ = _ticket()
        EinlassHelfer.objects.create(evenement=ev, user=helfer)
        c = _client(helfer)
        assert c.get(_url(ev, "helfer/")).status_code == 403
        assert c.post(_url(ev, "helfer/"), {"email": "x@y.de"}, format="json").status_code == 403

    def test_verantwortlicher_der_region_darf_bei_treffen(self):
        region = Region.objects.get(slug="nord")
        v = Veranstaltung.objects.create(
            region=region,
            titel="Public Viewing",
            datum=timezone.localdate() + datetime.timedelta(days=2),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
            ticket_aktiv=True,
        )
        boss, _ = _user("boss@example.de")
        Verantwortlicher.objects.create(region=region, user=boss, anzeige_name="Boss")
        fremder, _ = _user("fremd@example.de")
        other = Region.objects.exclude(slug="nord").first()
        Verantwortlicher.objects.create(region=other, user=fremder, anzeige_name="Fremd")
        assert _client(boss).get(_url(v)).status_code == 200
        assert _client(fremder).get(_url(v)).status_code == 403

    def test_unbekannter_typ_404(self):
        ev, _ = _ticket()
        r = _client(_buero()).get(f"{API}quatsch/{ev.pk}/")
        assert r.status_code == 404


class TestUebersicht:
    def test_zaehler_und_personenliste(self):
        ev, t = _ticket(places=2, adults=0, kids=1)
        _, m2 = _user("zwei@example.de", prenom="Aya", nom="Zed")
        ins2 = InscriptionFactory(evenement=ev, membre=m2, places=1)
        c = _client(_buero())
        c.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        r = c.get(_url(ev))
        assert r.status_code == 200
        assert r.data["zaehler"] == {
            "gebucht_erwachsene": 3,
            "gebucht_kinder": 1,
            "drin_erwachsene": 1,
            "drin_kinder": 0,
        }
        assert [p["name"] for p in r.data["personen"]] == ["Aya Zed", "Sami Ben Ali"]
        assert {p["status"] for p in r.data["personen"]} == {"offen", "teilweise"}
        assert r.data["darf_helfer_verwalten"] is True
        assert ins2.pk

    def test_stornierte_anmeldung_nicht_in_der_liste(self):
        from apps.evenements.models import StatutInscription

        ev, t = _ticket()
        ins = t.inscription
        ins.statut = StatutInscription.ANNULEE
        ins.save()
        r = _client(_buero()).get(_url(ev))
        assert r.data["personen"] == []


class TestHelferVerwaltung:
    def test_benennen_und_entfernen(self):
        ev, _ = _ticket()
        _, _m = _user("helfer@example.de", prenom="Hela", nom="Fer")
        c = _client(_buero())
        r = c.post(_url(ev, "helfer/"), {"email": "HELFER@example.de"}, format="json")
        assert r.status_code == 201
        assert [h["name"] for h in r.data["ergebnisse"]] == ["Hela Fer"]
        assert Notification.objects.filter(destinataire__email="helfer@example.de").exists()
        # doppelt benennen ändert nichts
        c.post(_url(ev, "helfer/"), {"email": "helfer@example.de"}, format="json")
        assert EinlassHelfer.objects.count() == 1
        helfer_id = r.data["ergebnisse"][0]["id"]
        r2 = c.delete(_url(ev, f"helfer/{helfer_id}/"))
        assert r2.status_code == 200
        assert r2.data["ergebnisse"] == []

    def test_unbekannte_email(self):
        ev, _ = _ticket()
        r = _client(_buero()).post(_url(ev, "helfer/"), {"email": "nix@example.de"}, format="json")
        assert r.status_code == 400

    def test_leere_email(self):
        ev, _ = _ticket()
        r = _client(_buero()).post(_url(ev, "helfer/"), {"email": " "}, format="json")
        assert r.status_code == 400

    def test_helfer_fuer_treffen(self):
        v = Veranstaltung.objects.create(
            region=Region.objects.get(slug="nord"),
            titel="Public Viewing",
            datum=timezone.localdate(),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
            ticket_aktiv=True,
        )
        _user("helfer@example.de")
        r = _client(_buero()).post(
            _url(v, "helfer/"), {"email": "helfer@example.de"}, format="json"
        )
        assert r.status_code == 201
        assert EinlassHelfer.objects.get().veranstaltung_id == v.pk


class TestMeine:
    def test_buero_sieht_alle_veranstaltungen_helfer_nur_ihre(self):
        ev, _ = _ticket()
        ohne_ticket = EvenementFactory(
            ticket_aktiv=False, date_evenement=timezone.localdate() + datetime.timedelta(days=1)
        )
        helfer, _ = _user("helfer@example.de")
        EinlassHelfer.objects.create(evenement=ev, user=helfer)
        buero = _client(_buero()).get(f"{API}meine/")
        assert {e["id"]: e["ticket_aktiv"] for e in buero.data["ergebnisse"]} == {
            str(ev.pk): True,
            str(ohne_ticket.pk): False,
        }
        assert {e["rolle"] for e in buero.data["ergebnisse"]} == {"verantwortlich"}
        r = _client(helfer).get(f"{API}meine/")
        assert [e["id"] for e in r.data["ergebnisse"]] == [str(ev.pk)]
        assert r.data["ergebnisse"][0]["rolle"] == "helfer"
        assert r.data["ergebnisse"][0]["heute"] is True

    def test_normales_mitglied_sieht_nichts(self):
        _ticket()
        user, _ = _user("normal@example.de")
        assert _client(user).get(f"{API}meine/").data["ergebnisse"] == []


class TestAnwesenheitOhneTickets:
    """Die Liste gilt für jede Veranstaltung, auch wenn keine Tickets ausgestellt werden."""

    def _ohne_ticket(self):
        _, membre = _user("gast@example.de")
        ev = EvenementFactory(ticket_aktiv=False)
        ins = InscriptionFactory(evenement=ev, membre=membre, places=2)
        return ev, ins

    def test_liste_zeigt_bestaetigte_anmeldungen_ohne_ticket(self):
        ev, ins = self._ohne_ticket()
        r = _client(_buero()).get(_url(ev))
        assert r.status_code == 200
        assert r.data["ticket_aktiv"] is False
        assert [(p["person"], p["ticket"], p["status"]) for p in r.data["personen"]] == [
            (str(ins.pk), False, "offen")
        ]

    def test_von_hand_eintragen_und_zaehler(self):
        ev, ins = self._ohne_ticket()
        c = _client(_buero())
        r = c.post(
            _url(ev, "einlassen/"),
            {"person": str(ins.pk), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        assert r.status_code == 200
        liste = c.get(_url(ev)).data
        assert liste["personen"][0]["status"] == "teilweise"
        assert liste["personen"][0]["anwesend_seit"] is not None
        assert liste["zaehler"]["drin_erwachsene"] == 1
        assert Einlass.objects.get().ticket is None

    def test_scan_und_liste_zeigen_denselben_stand(self):
        ev, t = _ticket(places=1, adults=0, kids=0)
        c = _client(_buero())
        c.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        person = c.get(_url(ev)).data["personen"][0]
        assert person["status"] == "vollstaendig"
        assert person["ticket"] is True
        assert Einlass.objects.get().ticket_id == t.pk

    def test_treffen_haekchen_von_hand_zaehlt_als_anwesend(self):
        v = Veranstaltung.objects.create(
            region=Region.objects.get(slug="nord"),
            titel="Public Viewing",
            datum=timezone.localdate(),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
            ticket_aktiv=False,
        )
        user, _ = _user("tn@example.de")
        Teilnahme.objects.create(veranstaltung=v, user=user, anwesend=True)
        liste = _client(_buero()).get(_url(v)).data
        assert liste["personen"][0]["status"] == "vollstaendig"
        assert liste["zaehler"]["drin_erwachsene"] == 1

    def test_unbestaetigte_anmeldung_steht_nicht_in_der_liste(self):
        from apps.evenements.models import StatutInscription

        ev, ins = self._ohne_ticket()
        ins.statut = StatutInscription.EN_ATTENTE_PAIEMENT
        ins.save()
        assert _client(_buero()).get(_url(ev)).data["personen"] == []


class TestExport:
    def test_excel_mit_namen_status_und_summe(self):
        import io

        from openpyxl import load_workbook

        ev, t = _ticket(places=2, adults=0, kids=1)
        c = _client(_buero())
        c.post(
            _url(ev, "einlassen/"),
            {"person": str(t.inscription_id), "erwachsene": 1, "kinder": 0},
            format="json",
        )
        r = c.get(_url(ev, "export/"))
        assert r.status_code == 200
        assert "spreadsheetml" in r["Content-Type"]
        assert "anwesenheit_" in r["Content-Disposition"]
        blatt = load_workbook(io.BytesIO(r.content)).active
        zeilen = [[z.value for z in zeile] for zeile in blatt.iter_rows()]
        assert zeilen[3][:2] == ["Name", "Status"]
        assert zeilen[4][:6] == ["Sami Ben Ali", "Teilweise anwesend", 2, 1, 1, 0]
        assert zeilen[4][6] == "ja"
        assert zeilen[-1][:6] == ["Summe", None, 2, 1, 1, 0]

    def test_excel_fuer_treffen(self):
        import io

        from openpyxl import load_workbook

        v = Veranstaltung.objects.create(
            region=Region.objects.get(slug="nord"),
            titel="Public Viewing",
            datum=timezone.localdate(),
            ort="Hamburg",
            status=VeranstaltungStatus.VEROEFFENTLICHT,
        )
        user, _ = _user("tn@example.de")
        Teilnahme.objects.create(veranstaltung=v, user=user)
        r = _client(_buero()).get(_url(v, "export/"))
        assert r.status_code == 200
        blatt = load_workbook(io.BytesIO(r.content)).active
        assert blatt["A5"].value == "Sami Ben Ali"

    def test_formel_im_namen_wird_neutralisiert(self):
        import io

        from openpyxl import load_workbook

        _, membre = _user("evil@example.de", prenom="=HYPERLINK(1)", nom="X")
        ev = EvenementFactory(ticket_aktiv=False)
        InscriptionFactory(evenement=ev, membre=membre)
        r = _client(_buero()).get(_url(ev, "export/"))
        blatt = load_workbook(io.BytesIO(r.content)).active
        assert blatt["A5"].value.startswith("'=")

    def test_helfer_und_fremde_duerfen_nicht_exportieren(self):
        helfer, _ = _user("helfer@example.de")
        ev, _ = _ticket()
        EinlassHelfer.objects.create(evenement=ev, user=helfer)
        assert _client(helfer).get(_url(ev, "export/")).status_code == 403
        zufaellig, _ = _user("zufall@example.de")
        assert _client(zufaellig).get(_url(ev, "export/")).status_code == 403
        assert APIClient().get(_url(ev, "export/")).status_code == 401

    def test_export_steht_im_aktivitaetsprotokoll(self):
        from apps.accounts.models import AuditLogEntry

        ev, _ = _ticket()
        _client(_buero()).get(_url(ev, "export/"))
        assert AuditLogEntry.objects.filter(action="eintritt_export").exists()
