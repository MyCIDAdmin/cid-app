"""Tests Welle 2 : Teilgruppen, gemeinsame Ankunft, Checkliste, Exporte, Saison, Steckbriefe,
Bereinigung."""

from datetime import timedelta
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import AuditLogEntry, Role
from apps.communaute.models import RencontreCalendrier
from apps.notifications.models import Notification, TypeNotification
from apps.reisen.bereinigung import bereinige_persoenliche_daten
from apps.reisen.models import Etappe, Reise, ReiseTeilnehmer

from .test_reisen import BASE, api, dabei, etappe, neue_reise, offen, teilnehmer_sicht, user

pytestmark = pytest.mark.django_db


@pytest.fixture
def orga():
    return user("orga@example.de")


@pytest.fixture
def anna():
    return user("anna@example.de")


@pytest.fixture
def ben():
    return user("ben@example.de")


@pytest.fixture
def clara():
    return user("clara@example.de")


def tn_id(rid, u):
    return str(ReiseTeilnehmer.objects.get(reise_id=rid, user=u).pk)


def gruppe_anlegen(orga_user, rid, name, *mitglieder):
    r = api(orga_user).post(
        f"{BASE}{rid}/teilgruppen/",
        {"name": name, "teilnehmer": [tn_id(rid, m) for m in mitglieder]},
        format="json",
    )
    assert r.status_code == 201, r.data
    return r.data["id"]


# --- Teilgruppen ------------------------------------------------------------------------------


def test_teilgruppe_filtert_info_entscheidung_und_aufgabe(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    gid = gruppe_anlegen(orga, rid, "Abflug Hamburg", anna)
    cl = api(orga)
    info = cl.post(
        f"{BASE}{rid}/infos/",
        {"titel": "Treffpunkt HH", "text": "<p>Terminal 1</p>", "wichtig": True, "teilgruppe": gid},
        format="json",
    )
    assert info.status_code == 201, info.data
    ent = cl.post(
        f"{BASE}{rid}/entscheidungen/",
        {
            "typ": "auswahl",
            "baustein": "transport",
            "titel": "Zug oder Auto?",
            "teilgruppe": gid,
            "optionen": [{"titel": "Zug"}, {"titel": "Auto"}],
        },
        format="json",
    )
    assert ent.status_code == 201, ent.data
    auf = cl.post(
        f"{BASE}{rid}/aufgaben/", {"titel": "Zug buchen", "teilgruppe": gid}, format="json"
    )
    assert auf.status_code == 201

    # Anna (Mitglied) und Orga sehen alles, Ben nichts davon
    for u in (anna, orga):
        assert len(api(u).get(f"{BASE}{rid}/infos/").data) == 1
        assert len(api(u).get(f"{BASE}{rid}/entscheidungen/").data) == 1
        assert len(api(u).get(f"{BASE}{rid}/aufgaben/").data) == 1
    assert api(ben).get(f"{BASE}{rid}/infos/").data == []
    assert api(ben).get(f"{BASE}{rid}/entscheidungen/").data == []
    assert api(ben).get(f"{BASE}{rid}/aufgaben/").data == []
    detail = api(ben).get(f"{BASE}{rid}/").data
    assert detail["offene_entscheidungen"] == [] and detail["infos"] == []
    # Ben kann auch nicht darauf zugreifen
    eid = ent.data["id"]
    opt = api(anna).get(f"{BASE}{rid}/entscheidungen/").data[0]["optionen"][0]["id"]
    assert (
        api(ben)
        .post(
            f"{BASE}{rid}/entscheidungen/{eid}/stimmen/",
            {"stimmen": [{"option": opt}]},
            format="json",
        )
        .status_code
        == 404
    )
    assert api(ben).post(f"{BASE}{rid}/infos/{info.data['id']}/gelesen/").status_code == 404
    # Benachrichtigungen gehen nur an die Gruppe
    assert Notification.objects.filter(
        destinataire=anna, type_notification=TypeNotification.REISE_ENTSCHEIDUNG
    ).exists()
    assert not Notification.objects.filter(
        destinataire=ben, type_notification=TypeNotification.REISE_ENTSCHEIDUNG
    ).exists()
    assert not Notification.objects.filter(
        destinataire=ben, type_notification=TypeNotification.REISE_INFO
    ).exists()


def test_teilgruppe_einer_anderen_reise_wird_abgelehnt(orga):
    r1 = neue_reise(orga, **offen())
    r2 = neue_reise(orga, titel="Andere", **offen())
    gid = gruppe_anlegen(orga, r2, "Fremd")
    r = api(orga).post(f"{BASE}{r1}/infos/", {"titel": "X", "teilgruppe": gid}, format="json")
    assert r.status_code == 400


def test_teilgruppen_namen_nur_fuer_mitglieder_und_orga(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    gruppe_anlegen(orga, rid, "Hotel Nil", anna)
    g_ben = api(ben).get(f"{BASE}{rid}/teilgruppen/").data["gruppen"][0]
    assert g_ben["anzahl"] == 1 and "mitglieder" not in g_ben and not g_ben["ich_bin_mitglied"]
    assert "vorschlaege" not in api(ben).get(f"{BASE}{rid}/teilgruppen/").data
    g_anna = api(anna).get(f"{BASE}{rid}/teilgruppen/").data["gruppen"][0]
    assert g_anna["mitglieder"][0]["name"]
    # Nur die Organisation legt an
    assert (
        api(ben).post(f"{BASE}{rid}/teilgruppen/", {"name": "X"}, format="json").status_code == 403
    )


def test_freigabe_teilgruppe(orga, anna, ben, clara):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben, clara)
    gruppe_anlegen(orga, rid, "HH", anna, clara)
    etappe(anna, rid)
    assert teilnehmer_sicht(clara, rid, anna)["etappen"] == []
    api(anna).patch(f"{BASE}{rid}/meine/", {"freigabe_hin": "teilgruppe"}, format="json")
    assert len(teilnehmer_sicht(clara, rid, anna)["etappen"]) == 1
    assert teilnehmer_sicht(ben, rid, anna)["etappen"] == []


def test_vorschlaege_nur_aus_sichtbaren_angaben(orga, anna, ben, clara):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben, clara)
    etappe(anna, rid)
    etappe(ben, rid, von_ort="hamburg")
    etappe(clara, rid, von_ort="Hamburg")
    api(clara).patch(f"{BASE}{rid}/meine/", {"verbergen_vor_orga": True}, format="json")
    v = api(orga).get(f"{BASE}{rid}/teilgruppen/").data["vorschlaege"]
    assert len(v) == 1 and v[0]["name"].startswith("Abflug") and len(v[0]["teilnehmer"]) == 2


# --- Gemeinsame Ankunft -----------------------------------------------------------------------


def test_gemeinsame_ankunft_namen_erst_bei_beidseitiger_zustimmung(orga, anna, ben, clara):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben, clara)
    etappe(anna, rid, ankunft="2026-11-26T13:00:00Z")
    etappe(ben, rid, ankunft="2026-11-26T14:30:00Z")
    etappe(clara, rid, ankunft="2026-11-26T20:00:00Z", von_ort="Berlin")
    for u in (anna, ben, clara):
        api(u).patch(f"{BASE}{rid}/meine/", {"offen_gemeinsame_anreise": True}, format="json")
    g = api(anna).get(f"{BASE}{rid}/gemeinsame-ankunft/").data
    ank = next(x for x in g["gruppen"] if x["typ"] == "ankunft")
    assert ank["weitere"] == 1 and ank["kontakte"] == [] and not ank["ich_zugestimmt"]
    key = {"typ": "ankunft", "ort": ank["ort"], "tag": ank["tag"]}
    # Nur Anna stimmt zu : noch keine Namen
    g = (
        api(anna)
        .post(f"{BASE}{rid}/gemeinsame-ankunft/", {**key, "zustimmen": True}, format="json")
        .data
    )
    assert next(x for x in g["gruppen"] if x["typ"] == "ankunft")["kontakte"] == []
    # Ben stimmt zu : beide sehen sich
    g = (
        api(ben)
        .post(f"{BASE}{rid}/gemeinsame-ankunft/", {**key, "zustimmen": True}, format="json")
        .data
    )
    assert len(next(x for x in g["gruppen"] if x["typ"] == "ankunft")["kontakte"]) == 1
    g = api(anna).get(f"{BASE}{rid}/gemeinsame-ankunft/").data
    assert len(next(x for x in g["gruppen"] if x["typ"] == "ankunft")["kontakte"]) == 1
    # Clara (anderes Zeitfenster) ist nicht dabei
    g = api(clara).get(f"{BASE}{rid}/gemeinsame-ankunft/").data
    assert all(x["typ"] != "ankunft" or x["weitere"] == 0 for x in g["gruppen"])
    # Ausschalten nimmt die Zustimmung zurück
    api(ben).patch(f"{BASE}{rid}/meine/", {"offen_gemeinsame_anreise": False}, format="json")
    assert api(ben).get(f"{BASE}{rid}/gemeinsame-ankunft/").data == {"offen": False, "gruppen": []}
    g = api(anna).get(f"{BASE}{rid}/gemeinsame-ankunft/").data
    assert g["gruppen"] == [] or all(x["weitere"] == 0 for x in g["gruppen"])


def test_ohne_einwilligung_wird_niemand_verglichen(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    etappe(anna, rid)
    etappe(ben, rid)
    api(anna).patch(f"{BASE}{rid}/meine/", {"offen_gemeinsame_anreise": True}, format="json")
    g = api(anna).get(f"{BASE}{rid}/gemeinsame-ankunft/").data
    assert g["gruppen"] == []
    # Zustimmung zu einer Gruppe, die es nicht gibt
    r = api(anna).post(
        f"{BASE}{rid}/gemeinsame-ankunft/",
        {"typ": "ankunft", "ort": "Kairo", "tag": "2026-11-26", "zustimmen": True},
        format="json",
    )
    assert r.status_code == 400


# --- Checkliste -------------------------------------------------------------------------------


def test_checkliste_gemeinsam_persoenlich_und_erledigt(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    r = api(orga).post(
        f"{BASE}{rid}/checkliste/",
        {"titel": "Reisepass prüfen", "frist": "2020-01-01"},
        format="json",
    )
    assert r.status_code == 201
    api(anna).post(
        f"{BASE}{rid}/checkliste/", {"titel": "Adapter kaufen", "persoenlich": True}, format="json"
    )
    # gemeinsame Punkte legt nur die Organisation an
    assert (
        api(anna).post(f"{BASE}{rid}/checkliste/", {"titel": "X"}, format="json").status_code == 403
    )
    liste_anna = api(anna).get(f"{BASE}{rid}/checkliste/").data
    assert {p["titel"] for p in liste_anna} == {"Reisepass prüfen", "Adapter kaufen"}
    assert {p["titel"] for p in api(ben).get(f"{BASE}{rid}/checkliste/").data} == {
        "Reisepass prüfen"
    }
    punkt = next(p for p in liste_anna if not p["persoenlich"])
    assert punkt["ueberfaellig"] is True
    api(anna).post(
        f"{BASE}{rid}/checkliste/{punkt['id']}/erledigt/", {"erledigt": True}, format="json"
    )
    assert next(
        p for p in api(anna).get(f"{BASE}{rid}/checkliste/").data if p["id"] == punkt["id"]
    )["erledigt"]
    assert not next(
        p for p in api(ben).get(f"{BASE}{rid}/checkliste/").data if p["id"] == punkt["id"]
    )["erledigt"]
    # Organisation sieht den Stand (1 von 3 : Orga, Anna, Ben)
    o = next(p for p in api(orga).get(f"{BASE}{rid}/checkliste/").data if p["id"] == punkt["id"])
    assert o["erledigt_zahl"] == 1 and o["gesamt"] == 3
    # Fremden persönlichen Punkt kann niemand löschen
    pers = next(p for p in liste_anna if p["persoenlich"])
    assert api(ben).delete(f"{BASE}{rid}/checkliste/{pers['id']}/").status_code == 404
    assert api(anna).delete(f"{BASE}{rid}/checkliste/{pers['id']}/").status_code == 204


def test_vorlage_fuellt_checkliste(orga):
    rid = neue_reise(orga, vorlage="auswaerts_ausland")
    assert len(api(orga).get(f"{BASE}{rid}/checkliste/").data) >= 4


# --- Exporte ----------------------------------------------------------------------------------


def blatt(antwort, name):
    wb = load_workbook(BytesIO(antwort.content))
    return [[c.value for c in zeile] for zeile in wb[name].iter_rows()]


def test_export_nur_organisation_mit_freigaben_und_protokoll(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    etappe(anna, rid, anbieter="=HYPERLINK(1)")
    etappe(ben, rid, von_ort="Berlin")
    api(ben).patch(
        f"{BASE}{rid}/meine/", {"verbergen_vor_orga": True, "notiz": "Allergie"}, format="json"
    )
    assert api(anna).get(f"{BASE}{rid}/export-excel/").status_code == 403
    r = api(orga).get(f"{BASE}{rid}/export-excel/")
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    plan = blatt(r, "Anreiseplan")
    text = str(plan)
    assert "ABC123" not in text  # nie die Buchungsreferenz
    assert "Berlin" not in text  # Ben verbirgt Details vor der Organisation
    assert "'=HYPERLINK(1)" in text  # Formeln im Freitext werden entschärft
    assert "Allergie" not in str(blatt(r, "Teilnehmende"))
    assert AuditLogEntry.objects.filter(action="reisen_export", user=orga).exists()


def test_namensliste_nur_notwendiges_und_filter(orga, anna, ben):
    rid = neue_reise(orga, **offen())
    dabei(rid, anna, ben)
    api(anna).patch(
        f"{BASE}{rid}/meine/",
        {"geburtsdatum": "1990-05-17", "unterkunft_eigen": "Hotel Nil"},
        format="json",
    )
    gid = gruppe_anlegen(orga, rid, "HH", anna)
    r = api(orga).get(f"{BASE}{rid}/export-namensliste/")
    zeilen = blatt(r, "Namensliste")
    assert zeilen[0] == ["Name", "Personen", "Geburtsdatum"] and len(zeilen) == 4
    assert any("17.05.1990" in str(z) for z in zeilen)
    gefiltert = blatt(
        api(orga).get(f"{BASE}{rid}/export-namensliste/?teilgruppe={gid}"), "Namensliste"
    )
    assert len(gefiltert) == 2
    eigene = blatt(
        api(orga).get(f"{BASE}{rid}/export-namensliste/?unterkunft=eigene"), "Namensliste"
    )
    assert len(eigene) == 2
    assert api(anna).get(f"{BASE}{rid}/export-namensliste/").status_code == 403
    assert api(orga).get(f"{BASE}{rid}/export-namensliste/?teilgruppe={rid}").status_code == 404


# --- Saison -----------------------------------------------------------------------------------


def test_saison_zeigt_auswaertsspiele_mit_sichtbaren_reisen(orga, anna):
    heim = RencontreCalendrier.objects.create(
        evenement_externe_id="h1",
        competition="Ligue 1",
        equipe_domicile="Club Africain",
        equipe_exterieur="ES Tunis",
        date_heure=timezone.now() + timedelta(days=10),
    )
    ausw = RencontreCalendrier.objects.create(
        evenement_externe_id="a1",
        competition="CAF",
        equipe_domicile="Zamalek SC",
        equipe_exterieur="Club Africain",
        date_heure=timezone.now() + timedelta(days=20),
    )
    sichtbar = (
        api(orga).post(BASE, {"rencontre": str(ausw.id), **offen()}, format="json").data["id"]
    )
    api(orga).post(BASE, {"rencontre": str(ausw.id), "titel": "Privat"}, format="json")
    daten = api(anna).get(f"{BASE}saison/").data
    assert [s["rencontre"] for s in daten] == [str(ausw.id)] and str(heim.id) not in str(daten)
    assert [r["id"] for r in daten[0]["reisen"]] == [sichtbar]


# --- Steckbriefe ------------------------------------------------------------------------------


def test_laenderprofil_pflege_nur_bureau_und_in_reise_sichtbar(orga, anna):
    admin = user("admin@example.de", role=Role.BUREAU_ADMIN)
    daten = {
        "land": "Ägypten",
        "waehrung": "EGP",
        "steckdose": "Typ C/F",
        "einreise": "<p>Visum</p>",
    }
    assert api(anna).post(f"{BASE}laenderprofile/", daten, format="json").status_code == 403
    r = api(admin).post(f"{BASE}laenderprofile/", daten, format="json")
    assert r.status_code == 201 and r.data["stand"] == timezone.localdate().isoformat()
    assert api(admin).post(f"{BASE}laenderprofile/", daten, format="json").status_code == 400
    rid = neue_reise(orga, **offen(land="ägypten"))
    dabei(rid, anna)
    assert api(anna).get(f"{BASE}{rid}/").data["laenderprofil"]["waehrung"] == "EGP"
    assert api(anna).get(f"{BASE}laenderprofile/").status_code == 200
    pid = r.data["id"]
    assert (
        api(admin)
        .patch(f"{BASE}laenderprofile/{pid}/", {"waehrung": "EGP (£E)"}, format="json")
        .status_code
        == 200
    )
    assert api(admin).delete(f"{BASE}laenderprofile/{pid}/").status_code == 204


# --- Bereinigung nach 90 Tagen ----------------------------------------------------------------


def test_bereinigung_loescht_persoenliche_daten_nach_90_tagen(orga, anna):
    alt = neue_reise(orga, **offen())
    neu = neue_reise(orga, titel="Neu", **offen())
    dabei(alt, anna)
    dabei(neu, anna)
    for rid in (alt, neu):
        etappe(anna, rid)
        api(anna).patch(
            f"{BASE}{rid}/meine/",
            {
                "notiz": "privat",
                "geburtsdatum": "1990-01-01",
                "unterkunft_eigen": "Hotel",
                "freigabe_hin": "alle",
            },
            format="json",
        )
        api(anna).post(
            f"{BASE}{rid}/checkliste/", {"titel": "Persönlich", "persoenlich": True}, format="json"
        )
    heute = timezone.localdate()
    Reise.objects.filter(pk=alt).update(ende=heute - timedelta(days=91))
    Reise.objects.filter(pk=neu).update(ende=heute - timedelta(days=89))
    assert bereinige_persoenliche_daten() == {"reisen": 1}
    tn = ReiseTeilnehmer.objects.get(reise_id=alt, user=anna)
    assert tn.notiz == "" and tn.geburtsdatum is None and tn.unterkunft_eigen == ""
    assert tn.freigabe_hin == "privat"
    assert not Etappe.objects.filter(teilnehmer=tn).exists()
    assert not api(anna).get(f"{BASE}{alt}/checkliste/").data
    assert Reise.objects.get(pk=alt).daten_bereinigt_am is not None
    # Die jüngere Reise bleibt unberührt, ein zweiter Lauf ändert nichts
    assert ReiseTeilnehmer.objects.get(reise_id=neu, user=anna).notiz == "privat"
    assert bereinige_persoenliche_daten() == {"reisen": 0}
    # Die Reise selbst bleibt als Rückblick (Teilnehmerliste bleibt)
    assert ReiseTeilnehmer.objects.filter(reise_id=alt).count() == 2
    assert AuditLogEntry.objects.filter(action="reisen_bereinigt").exists()


def test_aenderungen_an_reisen_stehen_im_aktivitaetsprotokoll(orga):
    neue_reise(orga)
    assert AuditLogEntry.objects.filter(action="reisen_erstellt", user=orga).exists()
