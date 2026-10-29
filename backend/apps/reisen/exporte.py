"""Excel-Exporte für die Organisation. Sie enthalten personenbezogene Daten ; jeder Abruf steht
im Aktivitätsprotokoll (die Middleware erfasst Pfade mit `export-…` als „Reisen: Export“).

Die Exporte folgen denselben Freigaben wie die Oberfläche : was eine Person vor der Organisation
verbirgt, erscheint nur grob (Zielort, Tag, gebucht ja/nein). Buchungsreferenzen gehören nie
hinein."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import services as sv
from .models import Rolle, TeilgruppeMitglied, TStatus

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GEFAEHRLICH = ("=", "+", "-", "@", "\t", "\r")


def _zelle(wert):
    """Schützt vor Formeln in Freitext (Excel führt Zellen mit = + - @ sonst aus)."""
    if isinstance(wert, str) and wert.startswith(GEFAEHRLICH):
        return "'" + wert
    return wert


def _blatt(wb, titel, kopf, zeilen, erstes=False):
    ws = wb.active if erstes else wb.create_sheet()
    ws.title = titel[:31]
    ws.append(kopf)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)
    for zeile in zeilen:
        ws.append([_zelle(w) for w in zeile])
    for i, spalte in enumerate(kopf, start=1):
        laenge = max([len(str(spalte))] + [len(str(z[i - 1])) for z in zeilen if i - 1 < len(z)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(laenge + 2, 10), 50)
    ws.freeze_panes = "A2"
    return ws


def _bytes(wb) -> bytes:
    puffer = BytesIO()
    wb.save(puffer)
    return puffer.getvalue()


def _ja(wert) -> str:
    return "ja" if wert else "nein"


def reise_excel(ctx) -> bytes:
    reise = ctx.reise
    dabei = list(
        reise.teilnehmer.filter(status=TStatus.DABEI)
        .select_related("user__membre", "unterkunft")
        .prefetch_related("etappen")
        .order_by("created_at")
    )
    wb = Workbook()
    teilnehmer, plan, belegung = [], [], []
    for tn in dabei:
        name = sv.name_von(tn.user)
        verborgen = tn.verbergen_vor_orga
        teilnehmer.append(
            [
                name,
                dict(Rolle.choices).get(tn.rolle, tn.rolle),
                1 + tn.begleitpersonen,
                _ja(sv.hat_gebucht(tn, "hin")),
                _ja(sv.hat_gebucht(tn, "rueck")),
                _ja(tn.unterkunft_gebucht),
                "" if verborgen else _ja(tn.stadionkarte),
                "" if verborgen else _ja(tn.dokument_gueltig),
                "" if verborgen else _ja(tn.regeln_bestaetigt_am),
                "" if verborgen else tn.notiz,
            ]
        )
        for e in sorted(tn.etappen.all(), key=lambda x: (x.abfahrt is None, x.abfahrt)):
            if verborgen:
                g = sv.etappe_grob(e)
                plan.append(
                    [
                        name,
                        g["richtung"],
                        "",
                        "",
                        g["nach_ort"],
                        g["tag"] or "",
                        "",
                        "",
                        _ja(g["gebucht"]),
                    ]
                )
            else:
                plan.append(
                    [
                        name,
                        e.richtung,
                        e.art,
                        e.von_ort,
                        e.nach_ort,
                        _lokal(e.abfahrt),
                        _lokal(e.ankunft),
                        e.anbieter,
                        e.status,
                    ]
                )
        if verborgen:
            ort = ""
        elif tn.unterkunft_id:
            ort = tn.unterkunft.name
        else:
            ort = tn.unterkunft_eigen
        if ort:
            belegung.append([ort, name, 1 + tn.begleitpersonen, _ja(tn.unterkunft_gebucht)])
    belegung.sort(key=lambda z: (z[0].lower(), z[1].lower()))
    _blatt(
        wb,
        "Teilnehmende",
        [
            "Name",
            "Rolle",
            "Personen",
            "Hinreise gebucht",
            "Rückreise gebucht",
            "Unterkunft gebucht",
            "Stadionkarte",
            "Reisedokument gültig",
            "Regeln bestätigt",
            "Notiz",
        ],
        teilnehmer,
        erstes=True,
    )
    _blatt(
        wb,
        "Anreiseplan",
        ["Name", "Richtung", "Art", "Von", "Nach", "Abfahrt", "Ankunft", "Anbieter", "Status"],
        plan,
    )
    _blatt(wb, "Belegung", ["Unterkunft", "Name", "Personen", "Gebucht"], belegung)
    return _bytes(wb)


def _lokal(zeit) -> str:
    from django.utils import timezone

    return timezone.localtime(zeit).strftime("%d.%m.%Y %H:%M") if zeit else ""


def namensliste_excel(ctx, unterkunft: str | None = None, teilgruppe: str | None = None) -> bytes:
    """Nur das Nötigste für Dritte (Hotel, Stadion) : Name, Personen, Geburtsdatum (nur wenn die
    Person es für diese Reise angegeben hat). Personen, die Angaben vor der Organisation
    verbergen, fehlen in gefilterten Listen, weil die Zuordnung nicht sichtbar ist."""
    qs = ctx.reise.teilnehmer.filter(status=TStatus.DABEI).select_related(
        "user__membre", "unterkunft"
    )
    gefiltert = bool(unterkunft or teilgruppe)
    if gefiltert:
        qs = qs.filter(verbergen_vor_orga=False)
    if unterkunft == "eigene":
        qs = qs.filter(unterkunft__isnull=True).exclude(unterkunft_eigen="")
    elif unterkunft:
        qs = qs.filter(unterkunft_id=unterkunft)
    if teilgruppe:
        ids = TeilgruppeMitglied.objects.filter(
            teilgruppe_id=teilgruppe, teilgruppe__reise=ctx.reise
        ).values_list("teilnehmer_id", flat=True)
        qs = qs.filter(pk__in=list(ids))
    zeilen = []
    for tn in qs.order_by("user__membre__nom", "user__membre__prenom"):
        zeilen.append(
            [
                sv.name_von(tn.user),
                1 + tn.begleitpersonen,
                tn.geburtsdatum.strftime("%d.%m.%Y") if tn.geburtsdatum else "",
            ]
        )
    wb = Workbook()
    _blatt(wb, "Namensliste", ["Name", "Personen", "Geburtsdatum"], zeilen, erstes=True)
    return _bytes(wb)
