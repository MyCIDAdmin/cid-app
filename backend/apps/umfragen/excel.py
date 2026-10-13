"""Excel-Export der Umfrage-Ergebnisse (Release 14) : Blatt „Auswertung“ (je Frage und
Antwortmöglichkeit Anzahl und Anteil), Blatt „Freitext“ (Freitext-Antworten und Bemerkungen) und –
nur für die Verwaltung bei namentlichen Umfragen – Blatt „Einzelantworten“. Sprache DE/FR wie beim
PDF-Bericht (Übersetzungen aus apps.uebersetzung, sonst Originaltext)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pdf import _text

TEXTE = {
    "de": {
        "auswertung": "Auswertung",
        "freitext": "Freitext",
        "einzel": "Einzelantworten",
        "thema": "Thema",
        "frage": "Frage",
        "antwort": "Antwort",
        "anzahl": "Anzahl",
        "anteil": "Anteil (%)",
        "text": "Text",
        "art": "Art",
        "bemerkung": "Bemerkung",
        "freitext_art": "Freitext",
        "name": "Name",
        "email": "E-Mail",
        "datum": "Datum",
        "teilnehmende": "Teilnahmen",
        "ohne_thema": "Allgemein",
    },
    "fr": {
        "auswertung": "Résultats",
        "freitext": "Texte libre",
        "einzel": "Réponses individuelles",
        "thema": "Thème",
        "frage": "Question",
        "antwort": "Réponse",
        "anzahl": "Nombre",
        "anteil": "Part (%)",
        "text": "Texte",
        "art": "Type",
        "bemerkung": "Remarque",
        "freitext_art": "Texte libre",
        "name": "Nom",
        "email": "E-mail",
        "datum": "Date",
        "teilnehmende": "Participations",
        "ohne_thema": "Général",
    },
}

_KOPF = PatternFill("solid", start_color="CC0000")


def _kopfzeile(blatt, spalten, breiten):
    blatt.append(spalten)
    for zelle in blatt[1]:
        zelle.font = Font(bold=True, color="FFFFFF")
        zelle.fill = _KOPF
        zelle.alignment = Alignment(vertical="center", wrap_text=True)
    for i, breite in enumerate(breiten, start=1):
        blatt.column_dimensions[get_column_letter(i)].width = breite
    blatt.freeze_panes = "A2"


def erzeuge_bericht_excel(*, umfrage, daten: dict, sprache: str) -> bytes:
    sprache = sprache if sprache in TEXTE else "de"
    t = TEXTE[sprache]
    arbeitsmappe = Workbook()

    blatt = arbeitsmappe.active
    blatt.title = t["auswertung"]
    _kopfzeile(
        blatt,
        [t["thema"], t["frage"], t["antwort"], t["anzahl"], t["anteil"]],
        [24, 48, 36, 10, 12],
    )
    fragen_titel = []
    for frage in daten["fragen"]:
        thema = _text(frage, "thema", frage.get("thema") or "", sprache) or t["ohne_thema"]
        titel = _text(frage, "text", frage["text"], sprache)
        fragen_titel.append(titel)
        for option in frage.get("optionen") or []:
            blatt.append(
                [
                    thema,
                    titel,
                    _text(option, "text", option["text"], sprache),
                    option["anzahl"],
                    option["anteil"],
                ]
            )
    if blatt.max_row > 1:
        blatt.conditional_formatting.add(
            f"E2:E{blatt.max_row}",
            DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=100, color="CC0000"
            ),
        )
    blatt.append([])
    blatt.append([t["teilnehmende"], daten["teilnehmende"]])

    texte = arbeitsmappe.create_sheet(t["freitext"])
    _kopfzeile(texte, [t["thema"], t["frage"], t["art"], t["text"]], [24, 48, 14, 80])
    for frage in daten["fragen"]:
        thema = _text(frage, "thema", frage.get("thema") or "", sprache) or t["ohne_thema"]
        titel = _text(frage, "text", frage["text"], sprache)
        art = t["freitext_art"] if frage["typ"] == "freitext" else t["bemerkung"]
        for eintrag in frage.get("freitexte") or []:
            texte.append([thema, titel, art, eintrag])
            texte.cell(row=texte.max_row, column=4).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

    if daten.get("einreichungen") is not None:
        einzel = arbeitsmappe.create_sheet(t["einzel"])
        originale = [f["text"] for f in daten["fragen"]]
        _kopfzeile(
            einzel,
            [t["name"], t["email"], t["datum"], *fragen_titel],
            [26, 30, 18] + [32] * len(fragen_titel),
        )
        for e in daten["einreichungen"]:
            nach_frage = {a["frage"]: ", ".join(a["werte"]) for a in e["antworten"]}
            datum = e["datum"]
            einzel.append(
                [
                    e["name"],
                    e["email"],
                    datum.strftime("%d.%m.%Y %H:%M") if hasattr(datum, "strftime") else str(datum),
                    *[nach_frage.get(text, "") for text in originale],
                ]
            )

    puffer = BytesIO()
    arbeitsmappe.save(puffer)
    return puffer.getvalue()
