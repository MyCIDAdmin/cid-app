"""Tests — Release 6 (2026-10-12): Preis pro Variante, Zahlungsfrist mit automatischem Verfall,
Abholung ohne Lieferadresse, Excel-Export pro Artikel mit Status/Filtern."""

import io
from datetime import timedelta
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.boutique.models import (
    Commande,
    LigneCommande,
    Reservierung,
    ReservierungsModus,
    StatutCommande,
)
from apps.boutique.tasks import verfallene_zahlungen_stornieren
from apps.boutique.tests.factories import CommandeFactory, ProduitFactory, VarianteProduitFactory
from apps.membres.models import StatutMembre
from apps.membres.tests.factories import MembreFactory

pytestmark = pytest.mark.django_db

PASSER = "boutique:commande-passer"


def _user(role, email):
    user = User.objects.create_user(email=email, password="Password123!", role=role, is_active=True)
    return user, MembreFactory(user=user, statut=StatutMembre.ACTIF)


def _client(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


# ── Preis pro Variante ───────────────────────────────────────────────────────────────────────


def test_variante_ohne_eigenen_preis_nutzt_produktpreis():
    produkt = ProduitFactory(prix=Decimal("30.00"), pourcentage_reduction=10)
    variante = VarianteProduitFactory(produit=produkt, taille="M")
    assert produkt.prix_pour_membre(None, variante) == (Decimal("27.00"), False)


def test_variante_mit_eigenem_preis_und_rabatt():
    produkt = ProduitFactory(prix=Decimal("30.00"), pourcentage_reduction=10)
    variante = VarianteProduitFactory(produit=produkt, taille="L", prix=Decimal("40.00"))
    assert produkt.prix_pour_membre(None, variante) == (Decimal("36.00"), False)


def test_mitgliederpreis_regeln_der_variante():
    _, membre = _user(Role.MEMBRE, "preis@example.de")
    produkt = ProduitFactory(prix=Decimal("30.00"), prix_membre=Decimal("25.00"))
    normal = VarianteProduitFactory(produit=produkt, taille="M")
    eigener_preis = VarianteProduitFactory(produit=produkt, taille="L", prix=Decimal("40.00"))
    eigener_mitglied = VarianteProduitFactory(
        produit=produkt, taille="XL", prix=Decimal("40.00"), prix_membre=Decimal("33.00")
    )
    assert produkt.prix_pour_membre(membre, normal) == (Decimal("25.00"), True)
    # eigener Normalpreis ohne eigenen Mitgliederpreis: der feste Produkt-Mitgliederpreis gilt nicht
    assert produkt.prix_pour_membre(membre, eigener_preis) == (Decimal("40.00"), False)
    assert produkt.prix_pour_membre(membre, eigener_mitglied) == (Decimal("33.00"), True)


def test_api_zeigt_variantenpreise_und_ab_preis():
    produkt = ProduitFactory(prix=Decimal("30.00"))
    VarianteProduitFactory(produit=produkt, taille="S", prix=Decimal("25.00"))
    VarianteProduitFactory(produit=produkt, taille="L", prix=Decimal("35.00"))
    user, _ = _user(Role.MEMBRE, "api-preis@example.de")
    resp = _client(user).get(reverse("boutique:produit-detail", args=[produkt.id]))
    assert resp.status_code == 200
    assert resp.data["prix_ab"] == "25.00"
    assert resp.data["preise_unterschiedlich"] is True
    preise = {v["taille"]: v["prix_affiche"] for v in resp.data["variantes"]}
    assert preise == {"S": "25.00", "L": "35.00"}


def test_bestellung_friert_den_variantenpreis_ein():
    produkt = ProduitFactory(prix=Decimal("30.00"))
    variante = VarianteProduitFactory(produit=produkt, taille="L", prix=Decimal("45.00"), stock=5)
    user, _ = _user(Role.MEMBRE, "bestell-preis@example.de")
    resp = _client(user).post(
        reverse(PASSER),
        {
            "lignes": [{"variante": str(variante.id), "quantite": 2}],
            "nom_destinataire": "T",
            "adresse_livraison": "Str. 1",
            "code_postal_livraison": "10115",
            "ville_livraison": "Berlin",
        },
        format="json",
    )
    assert resp.status_code == 201, resp.data
    assert Decimal(resp.data["montant_total"]) == Decimal("90.00")
    assert LigneCommande.objects.get().prix_unitaire == Decimal("45.00")


def test_verwaltung_kann_variantenpreis_setzen():
    admin, _ = _user(Role.BUREAU_ADMIN, "var-admin@example.de")
    produkt = ProduitFactory()
    variante = VarianteProduitFactory(produit=produkt)
    resp = _client(admin).patch(
        reverse("boutique:variante-detail", args=[variante.id]),
        {"prix": "12.50", "prix_membre": "10.00"},
        format="json",
    )
    assert resp.status_code == 200, resp.data
    variante.refresh_from_db()
    assert variante.prix == Decimal("12.50") and variante.prix_membre == Decimal("10.00")


# ── Reservierung mit Zahlung: Frist, Abholung, Verfall ───────────────────────────────────────


def _reservierungsprodukt(tage=7, stock=10):
    produkt = ProduitFactory(
        reservierung_modus=ReservierungsModus.RESERVIEREN_UND_BEZAHLEN,
        reservierung_zahlungsfrist_tage=tage,
    )
    return produkt, VarianteProduitFactory(produit=produkt, taille="", couleur="", stock=stock)


def _ohne_adresse(client, variante, menge=1):
    return client.post(
        reverse(PASSER),
        {"lignes": [{"variante": str(variante.id), "quantite": menge}]},
        format="json",
    )


def test_reservierungsartikel_ohne_adresse_mit_zahlungsfrist():
    _, variante = _reservierungsprodukt(tage=7)
    user, _ = _user(Role.MEMBRE, "res-pay@example.de")
    resp = _ohne_adresse(_client(user), variante, 2)
    assert resp.status_code == 201, resp.data
    commande = Commande.objects.get()
    assert commande.adresse_livraison.startswith("Abholung")
    assert commande.zahlung_bis is not None
    delta = commande.zahlung_bis - timezone.now()
    assert timedelta(days=6, hours=23) < delta <= timedelta(days=7)
    variante.refresh_from_db()
    assert variante.stock == 8


def test_normaler_artikel_braucht_weiterhin_eine_adresse():
    variante = VarianteProduitFactory(stock=5)
    user, _ = _user(Role.MEMBRE, "ohne-adr@example.de")
    resp = _ohne_adresse(_client(user), variante)
    assert resp.status_code == 400
    assert "adresse_livraison" in resp.data or "adresse_livraison" in str(resp.data)


def test_gemischter_warenkorb_braucht_adresse():
    _, res_variante = _reservierungsprodukt()
    normal = VarianteProduitFactory(stock=5)
    user, _ = _user(Role.MEMBRE, "mix@example.de")
    resp = _client(user).post(
        reverse(PASSER),
        {
            "lignes": [
                {"variante": str(res_variante.id), "quantite": 1},
                {"variante": str(normal.id), "quantite": 1},
            ]
        },
        format="json",
    )
    assert resp.status_code == 400


def test_ohne_frist_kein_zahlungsziel():
    _, variante = _reservierungsprodukt(tage=None)
    user, _ = _user(Role.MEMBRE, "keine-frist@example.de")
    assert _ohne_adresse(_client(user), variante).status_code == 201
    assert Commande.objects.get().zahlung_bis is None


def test_verfall_storniert_unbezahlte_und_gibt_menge_frei():
    _, variante = _reservierungsprodukt(tage=1, stock=10)
    user, _ = _user(Role.MEMBRE, "verfall@example.de")
    assert _ohne_adresse(_client(user), variante, 3).status_code == 201
    variante.refresh_from_db()
    assert variante.stock == 7

    assert verfallene_zahlungen_stornieren(jetzt=timezone.now()) == 0  # Frist läuft noch
    assert verfallene_zahlungen_stornieren(jetzt=timezone.now() + timedelta(days=2)) == 1

    commande = Commande.objects.get()
    assert commande.statut == StatutCommande.ANNULEE
    variante.refresh_from_db()
    assert variante.stock == 10
    # idempotent
    assert verfallene_zahlungen_stornieren(jetzt=timezone.now() + timedelta(days=3)) == 0


def test_verfall_laesst_bezahlte_bestellungen_in_ruhe():
    commande = CommandeFactory(statut=StatutCommande.CONFIRMEE)
    Commande.objects.filter(pk=commande.pk).update(zahlung_bis=timezone.now() - timedelta(days=1))
    assert verfallene_zahlungen_stornieren() == 0
    commande.refresh_from_db()
    assert commande.statut == StatutCommande.CONFIRMEE


def test_verfallene_menge_zaehlt_nicht_mehr_gegen_das_limit():
    produkt, variante = _reservierungsprodukt(tage=1, stock=10)
    produkt.reservierung_max_gesamt = 2
    produkt.save()
    user, _ = _user(Role.MEMBRE, "limit@example.de")
    client = _client(user)
    assert _ohne_adresse(client, variante, 2).status_code == 201
    assert _ohne_adresse(client, variante, 1).status_code == 400  # Limit erreicht
    verfallene_zahlungen_stornieren(jetzt=timezone.now() + timedelta(days=2))
    assert _ohne_adresse(client, variante, 2).status_code == 201


# ── Excel-Export pro Artikel ─────────────────────────────────────────────────────────────────


def _bestellung(membre, variante, menge, statut):
    commande = CommandeFactory(membre=membre, statut=statut)
    LigneCommande.objects.create(
        commande=commande,
        variante=variante,
        quantite=menge,
        prix_unitaire=Decimal("20.00"),
    )
    return commande


def test_export_nach_artikeln_und_status_mit_summen():
    admin, _ = _user(Role.BUREAU_ADMIN, "exp-admin@example.de")
    _, m1 = _user(Role.MEMBRE, "exp-m1@example.de")
    _, m2 = _user(Role.MEMBRE, "exp-m2@example.de")
    trikot = ProduitFactory(nom="Trikot")
    ticket = ProduitFactory(nom="Ticket")
    trikot_l = VarianteProduitFactory(produit=trikot, taille="L", couleur="")
    ticket_v = VarianteProduitFactory(produit=ticket, taille="Block A", couleur="")
    _bestellung(m1, trikot_l, 2, StatutCommande.CONFIRMEE)
    _bestellung(m2, trikot_l, 1, StatutCommande.EN_ATTENTE)
    _bestellung(m1, ticket_v, 4, StatutCommande.CONFIRMEE)
    Reservierung.objects.create(produit=trikot, variante=trikot_l, membre=m1, quantite=1)

    resp = _client(admin).get(reverse("boutique:commande-export"), {"produit": str(trikot.id)})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Commandes", "Positionen", "Summen", "Reservierungen"]
    assert wb["Commandes"].max_row == 3  # nur Bestellungen mit Trikot
    positionen = [[c.value for c in row] for row in wb["Positionen"].iter_rows(min_row=2)]
    assert {p[6] for p in positionen} == {"Trikot"}
    summe = [c.value for c in wb["Summen"][2]]
    assert summe[:3] == ["Trikot", "L", ""] or summe[0] == "Trikot"
    assert summe[3] == 3 and summe[4] == 2 and summe[5] == 1
    assert wb["Reservierungen"].max_row == 2

    resp = _client(admin).get(
        reverse("boutique:commande-export"),
        {"produit": f"{trikot.id},{ticket.id}", "statuts": StatutCommande.CONFIRMEE},
    )
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb["Commandes"].max_row == 3  # zwei bestätigte Bestellungen
    assert {r[6].value for r in wb["Positionen"].iter_rows(min_row=2)} == {"Trikot", "Ticket"}


def test_mitglied_exportiert_nur_eigene_ohne_reservierungsblatt():
    _, m1 = _user(Role.MEMBRE, "own-m1@example.de")
    user2, m2 = _user(Role.MEMBRE, "own-m2@example.de")
    variante = VarianteProduitFactory()
    _bestellung(m1, variante, 1, StatutCommande.CONFIRMEE)
    _bestellung(m2, variante, 1, StatutCommande.CONFIRMEE)
    resp = _client(user2).get(reverse("boutique:commande-export"))
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Reservierungen" not in wb.sheetnames
    assert wb["Commandes"].max_row == 2
