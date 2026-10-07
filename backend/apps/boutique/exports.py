"""
Export Excel des commandes — app boutique (demande utilisateur du 2026-09-25, module
"Shop-Verwaltung" : "Es soll möglich sein die Bestellungen als Excel zu exportieren").

Le queryset est filtré/scopé par l'appelant (CommandeViewSet.export — mêmes filtres que la
liste, voir CommandeFilter, et même scope IDOR que get_queryset) : ce module ne fait que
construire le classeur, même séparation des responsabilités que
apps.membres.exports/export_views. Même style visuel (en-tête rouge #CC0000) que
apps.membres.exports.construire_classeur_export, pour une seule identité visuelle d'export dans
toute l'application.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.membres.utils_http import safe_cell

_EN_TETE_STYLE = Font(bold=True, color="FFFFFF")
_REMPLISSAGE = PatternFill(start_color="CC0000", end_color="CC0000", fill_type="solid")

# Colonnes disponibles à l'export, dans leur ordre d'affichage — pas de sélection de colonnes
# ici (contrairement à apps.membres.exports) : la demande utilisateur ne porte que sur l'export
# en lui-même, sans mention d'un besoin de personnalisation des colonnes.
_COLONNES = [
    "N° commande",
    "Destinataire",
    "Email membre",
    "Montant total (€)",
    "Bon d'achat déduit (€)",
    "Montant dû (€)",
    "Statut",
    "Mode de paiement",
    "Date de commande",
    "Date de paiement",
    "N° de suivi",
    "Transporteur",
    "Date d'expédition",
]

# Colonnes forcées au format Texte Excel (number_format '@') — mêmes raisons que
# apps.membres.exports.CHAMPS_TEXTE : un identifiant/numéro de suivi ne doit jamais être
# réinterprété comme un nombre à l'ouverture du fichier.
_COLONNES_TEXTE = {1, 11}  # N° commande, N° de suivi


def _formater_date(valeur) -> str:
    return valeur.strftime("%d/%m/%Y %H:%M") if valeur else ""


def construire_classeur_commandes(queryset) -> Workbook:
    """
    Construit (sans l'enregistrer) le classeur .xlsx d'export des commandes — `queryset` doit
    déjà être filtré/scopé par l'appelant (voir CommandeViewSet.export). Un export n'est pas
    paginé : il contient toute la sélection filtrée, comme pour
    apps.membres.exports.construire_classeur_export.
    """
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Commandes"

    for col_idx, libelle in enumerate(_COLONNES, start=1):
        cellule = feuille.cell(row=1, column=col_idx, value=libelle)
        cellule.font = _EN_TETE_STYLE
        cellule.fill = _REMPLISSAGE

    ligne_idx = 2
    for commande in queryset.select_related("membre"):
        valeurs = [
            commande.numero_commande,
            commande.nom_destinataire,
            commande.membre.email if commande.membre else "",
            float(commande.montant_total),
            float(commande.montant_bon_achat),
            float(commande.montant_du),
            commande.get_statut_display(),
            commande.get_mode_paiement_display() if commande.mode_paiement else "",
            _formater_date(commande.created_at),
            _formater_date(commande.date_paiement_confirme),
            commande.numero_suivi,
            commande.transporteur,
            _formater_date(commande.date_expedition),
        ]
        for col_idx, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne_idx, column=col_idx, value=safe_cell(valeur))
            if col_idx in _COLONNES_TEXTE:
                cellule.number_format = "@"
        ligne_idx += 1

    for col_idx, libelle in enumerate(_COLONNES, start=1):
        feuille.column_dimensions[feuille.cell(row=1, column=col_idx).column_letter].width = max(
            len(libelle) + 2, 12
        )
    feuille.freeze_panes = "A2"

    return classeur
